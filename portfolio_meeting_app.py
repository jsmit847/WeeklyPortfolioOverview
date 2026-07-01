from __future__ import annotations

import datetime as dt
import html
import io
import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook


APP_TITLE = "Weekly Portfolio Meeting Deck"
APP_DIR = Path(__file__).resolve().parent

DEFAULT_SAMPLE_FILES = [
    APP_DIR / "Portfolio Overview 3.31.26.xlsx",
    APP_DIR / "portfolio_overview.xlsx",
    Path("Portfolio Overview 3.31.26.xlsx"),
    Path("portfolio_overview.xlsx"),
]

# Optional local file. Do not commit real deal-level history to a public repo.
STATUS_HISTORY_PATH = APP_DIR / "status_history.csv"
STATUS_HISTORY_SHEET_NAME = "_Status History"
STATUS_STALE_MEETING_THRESHOLD = 3
STATUS_HISTORY_COLUMNS = [
    "meeting_date",
    "sheet",
    "deal_number",
    "deal_key",
    "status",
    "owner",
    "commentary",
    "workbook_name",
    "saved_at",
]

DATE_HDR = re.compile(r"^\s*(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s+(.*\S)\s*$")

COMMON_STATUS_SUGGESTIONS = [
    "Surveillance",
    "Upcoming Maturity",
    "Modification",
    "Foreclosure",
    "Litigation",
    "Receivership",
    "Probate Court",
    "Needs Follow-Up",
    "Resolved",
]

# -----------------------------------------------------------------------------
# Loan Modification sheet settings
# -----------------------------------------------------------------------------
LOAN_MOD_SHEET_NAME = "Loan Modifications"

LOAN_MOD_DATE_COLUMNS = [
    "Cancellation Date",
    "Forbearance Termination Date",
    "Last Activity Date",
    "System Task Completed Date",
    "Mod Effective Date",
    "Modification Finalized Date",
    "Updated Maturity Date",
    "Previous Expiration Date",
    "Previous Maturity Date",
    "Prior Forbearance Termination Date",
    "Updated Expiration Date",
    "Created Date",
    "Last Modified Date",
    "System Modstamp",
    "Opportunity Close Date",
]

LOAN_MOD_MONEY_COLUMNS = [
    "Previous Loan Commitment",
    "Updated Loan Commitment",
]

LOAN_MOD_PERCENT_COLUMNS = [
    "Previous Exit Fee",
    "Previous Floor",
    "Previous Index Margin",
    "Previous Interest Rate",
    "Previous Max LTV %",
    "Previous Total ARV LTV",
    "Updated Exit Fee",
    "Updated Floor",
    "Updated Index Margin",
    "Updated Interest Rate",
    "Updated Max LTV %",
    "Updated Total ARV LTV",
]


# These are the columns the app needs to match loan mods back to the agenda.
# They stay lightweight and avoid loading the full Salesforce history table.
LOAN_MOD_REQUIRED_COLUMNS = [
    "Portfolio Overview Match",
    "Visible Overview Sheet(s)",
    "Deal Number",
    "Deal Name",
    "Opportunity Deal Loan Number",
    "Opportunity Deal Name",
]

# These are the highlighted fields in the Loan Modifications sheet.
# They get priority in Overview and Presentation.
LOAN_MOD_SPOTLIGHT_COLUMNS = [
    "Loan Mod Order Number",
    "Modification Type",
    "System Task Completed Date",
    "Mod Effective Date",
    "Previous Maturity Date",
    "Updated Maturity Date",
    "Previous Loan Commitment",
    "Updated Loan Commitment",
    "Comments",
]

# Small set of helper fields used only for sorting and labels.
LOAN_MOD_SUPPORT_COLUMNS = [
    "Loan Modification Name",
    "Status",
    "Loan Mod Type",
    "Mod Reporting Type",
    "Modification Finalized Date",
    "Deal Type",
    "Last Modified Date",
    "Created Date",
]

LOAN_MOD_CORE_COLUMNS = list(
    dict.fromkeys(
        LOAN_MOD_REQUIRED_COLUMNS
        + LOAN_MOD_SPOTLIGHT_COLUMNS
        + LOAN_MOD_SUPPORT_COLUMNS
    )
)



# -----------------------------------------------------------------------------
# Formatting / normalization helpers
# -----------------------------------------------------------------------------
def norm_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and np.isnan(value):
        return ""
    text = str(value).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def canon_header(value: object) -> str:
    return norm_text(value).upper()


def display_text(value: object, blank: str = "-") -> str:
    text = norm_text(value)
    return text if text else blank


def get_override_key(sheet: str, deal_number: str) -> str:
    return f"{norm_text(sheet)}::{norm_text(deal_number)}"


def first_existing_file(paths: Iterable[Path]) -> Optional[Path]:
    for path in paths:
        if path.exists():
            return path
    return None


def value_is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and np.isnan(value):
        return True
    return norm_text(value) == ""


def first_nonblank_value(*values: object) -> object:
    for value in values:
        if not value_is_blank(value):
            return value
    return None


def fmt_money(value: object, decimals: int = 0, blank: str = "-") -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return blank
    try:
        number = float(value)
    except Exception:
        return blank
    if np.isnan(number):
        return blank
    if abs(number) >= 1_000_000:
        return f"${number / 1_000_000:,.1f}MM"
    if decimals == 0:
        return f"${number:,.0f}"
    return f"${number:,.{decimals}f}"


def fmt_int(value: object, blank: str = "-") -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return blank
    try:
        return f"{int(round(float(value))):,}"
    except Exception:
        return blank


def coerce_datetime_value(value: object):
    if value is None:
        return pd.NaT
    if isinstance(value, float) and np.isnan(value):
        return pd.NaT
    if isinstance(value, (pd.Timestamp, dt.datetime, dt.date)):
        return pd.to_datetime(value, errors="coerce")

    try:
        number = float(value)
        if 20000 <= number <= 80000:
            return pd.Timestamp("1899-12-30") + pd.to_timedelta(number, unit="D")
    except Exception:
        pass

    return pd.to_datetime(value, errors="coerce")


def coerce_datetime_series(values: object) -> pd.Series:
    series = pd.Series(values)
    parsed = pd.to_datetime(series, errors="coerce")
    numeric = pd.to_numeric(series, errors="coerce")
    serial_mask = numeric.notna() & numeric.between(20000, 80000)

    if serial_mask.any():
        parsed.loc[serial_mask] = pd.to_datetime(
            numeric.loc[serial_mask],
            unit="D",
            origin="1899-12-30",
            errors="coerce",
        )

    return parsed


def fmt_date(value: object, blank: str = "-") -> str:
    timestamp = coerce_datetime_value(value)
    if pd.isna(timestamp):
        return blank
    return pd.Timestamp(timestamp).strftime("%m/%d/%Y")


def fmt_percent(value: object, decimals: int = 2, blank: str = "-") -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return blank
    try:
        number = float(value)
    except Exception:
        return blank
    if np.isnan(number):
        return blank
    if abs(number) <= 1.5:
        number *= 100
    return f"{number:,.{decimals}f}%"


def fmt_day_delta(days: object, blank: str = "-") -> str:
    if days is None or (isinstance(days, float) and np.isnan(days)):
        return blank
    try:
        days_i = int(round(float(days)))
    except Exception:
        return blank
    if days_i == 0:
        return "Today"
    if days_i > 0:
        return f"In {days_i:,}d"
    return f"{abs(days_i):,}d overdue"


def safe_series(df: pd.DataFrame, column_name: Optional[str], default: object = "") -> pd.Series:
    if column_name and column_name in df.columns:
        return df[column_name]
    return pd.Series([default] * len(df), index=df.index)


def safe_numeric_series(df: pd.DataFrame, column_name: Optional[str]) -> pd.Series:
    if column_name and column_name in df.columns:
        return pd.to_numeric(df[column_name], errors="coerce")
    return pd.Series([np.nan] * len(df), index=df.index, dtype=float)


def safe_datetime_series(df: pd.DataFrame, column_name: Optional[str]) -> pd.Series:
    if column_name and column_name in df.columns:
        return coerce_datetime_series(df[column_name])
    return pd.Series([pd.NaT] * len(df), index=df.index)


def resolve_column(columns: Iterable[object], *candidates: Optional[str]) -> Optional[str]:
    cols = [str(c) for c in columns if c is not None]
    for candidate in candidates:
        if not candidate:
            continue
        for col in cols:
            if canon_header(col) == canon_header(candidate):
                return col
    return None


def matching_columns(columns: Iterable[object], candidate: str) -> List[str]:
    return [
        str(c)
        for c in columns
        if c is not None and canon_header(c) == canon_header(candidate)
    ]


def find_header_row(ws, search_rows: int = 30, key_header: str = "Deal Number") -> Tuple[int, int]:
    max_row = min(search_rows, ws.max_row)
    for r in range(1, max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(canon_header(v) == canon_header(key_header) for v in values if v is not None):
            return r, ws.max_column
    raise ValueError(f"Could not find header row in sheet {ws.title!r}.")


def parse_header_date(
    month: int,
    day: int,
    year_text: Optional[str],
    as_of_date: dt.date,
) -> Optional[dt.date]:
    try:
        if year_text:
            year = int(year_text)
            if year < 100:
                year += 2000
            return dt.date(year, month, day)

        candidate = dt.date(as_of_date.year, month, day)
        if candidate > as_of_date:
            candidate = dt.date(as_of_date.year - 1, month, day)
        return candidate
    except ValueError:
        return None


def latest_dated_by_suffix(
    headers: Iterable[object],
    as_of_date: Optional[dt.date] = None,
) -> Dict[str, str]:
    as_of_date = as_of_date or dt.date.today()
    best: Dict[str, Tuple[dt.date, str]] = {}

    for header in headers:
        if header is None:
            continue
        text = norm_text(header)
        match = DATE_HDR.match(text)
        if not match:
            continue

        header_date = parse_header_date(
            month=int(match.group(1)),
            day=int(match.group(2)),
            year_text=match.group(3),
            as_of_date=as_of_date,
        )
        if header_date is None:
            continue

        suffix = canon_header(match.group(4))
        if suffix not in best or header_date > best[suffix][0]:
            best[suffix] = (header_date, text)

    return {suffix: full for suffix, (_date_value, full) in best.items()}


def fallback_name(
    primary: pd.Series,
    secondary: pd.Series,
    tertiary: pd.Series,
    final_fallback: pd.Series,
) -> pd.Series:
    result = primary.fillna("").astype(str)
    for candidate in [secondary, tertiary, final_fallback]:
        mask = result.map(norm_text) == ""
        result = result.where(~mask, candidate.fillna("").astype(str))
    return result.map(norm_text)


# -----------------------------------------------------------------------------
# Status history handling
# -----------------------------------------------------------------------------
def empty_status_history() -> pd.DataFrame:
    return pd.DataFrame(columns=STATUS_HISTORY_COLUMNS)


def normalize_status_history(history_df: Optional[pd.DataFrame]) -> pd.DataFrame:
    if history_df is None or history_df.empty:
        return empty_status_history()

    out = history_df.copy()
    for col in STATUS_HISTORY_COLUMNS:
        if col not in out.columns:
            out[col] = ""

    out = out[STATUS_HISTORY_COLUMNS].copy()
    out["meeting_date"] = pd.to_datetime(out["meeting_date"], errors="coerce")
    out["sheet"] = out["sheet"].map(norm_text)
    out["deal_number"] = out["deal_number"].map(norm_text)
    out["deal_key"] = out["deal_key"].map(norm_text)
    out["status"] = out["status"].map(norm_text)
    out["owner"] = out["owner"].map(norm_text)
    out["commentary"] = out["commentary"].fillna("").astype(str)
    out["workbook_name"] = out["workbook_name"].map(norm_text)
    out["saved_at"] = out["saved_at"].map(norm_text)

    out = out[
        out["meeting_date"].notna()
        & (out["deal_key"] != "")
        & (out["status"] != "")
    ].copy()

    if out.empty:
        return empty_status_history()

    out["meeting_date"] = out["meeting_date"].dt.date.astype(str)
    out = out.drop_duplicates(subset=["meeting_date", "deal_key"], keep="last")
    out = out.sort_values(["meeting_date", "sheet", "deal_number"], kind="stable")
    return out.reset_index(drop=True)


def read_status_history_csv(file_obj) -> pd.DataFrame:
    try:
        df = pd.read_csv(file_obj)
    except UnicodeDecodeError:
        file_obj.seek(0)
        df = pd.read_csv(file_obj, encoding="latin-1")
    return normalize_status_history(df)


def load_default_status_history_once() -> None:
    if st.session_state.get("status_history_default_checked", False):
        return

    st.session_state.status_history_default_checked = True
    if not STATUS_HISTORY_PATH.exists():
        return

    try:
        history = pd.read_csv(STATUS_HISTORY_PATH)
        st.session_state.status_history = normalize_status_history(history)
        st.session_state.uploaded_status_history_name = STATUS_HISTORY_PATH.name
    except Exception as exc:
        st.session_state.status_history_load_error = str(exc)


def current_status_history_snapshot(
    deck: pd.DataFrame,
    meeting_date: dt.date,
    workbook_name: str,
) -> pd.DataFrame:
    if deck.empty:
        return empty_status_history()

    snapshot = deck[["sheet", "deal_number", "deal_key", "status", "owner", "commentary"]].copy()
    snapshot.insert(0, "meeting_date", meeting_date.isoformat())
    snapshot["workbook_name"] = workbook_name
    snapshot["saved_at"] = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    snapshot["status"] = snapshot["status"].map(norm_text)
    snapshot = snapshot[snapshot["status"] != ""].copy()
    return normalize_status_history(snapshot)


def merge_status_history(
    existing_history: pd.DataFrame,
    deck: pd.DataFrame,
    meeting_date: dt.date,
    workbook_name: str,
) -> pd.DataFrame:
    existing = normalize_status_history(existing_history)
    current = current_status_history_snapshot(deck, meeting_date, workbook_name)
    merged = pd.concat([existing, current], ignore_index=True)
    return normalize_status_history(merged)


def status_history_to_csv_bytes(history_df: pd.DataFrame) -> bytes:
    history = normalize_status_history(history_df)
    return history.to_csv(index=False).encode("utf-8")


def read_embedded_status_history_from_workbook_bytes(file_bytes: bytes) -> pd.DataFrame:
    """Read hidden/embedded status history from the uploaded workbook, if present."""
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return empty_status_history()

    try:
        if STATUS_HISTORY_SHEET_NAME not in workbook.sheetnames:
            return empty_status_history()

        ws = workbook[STATUS_HISTORY_SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return empty_status_history()

        headers = [norm_text(value) for value in rows[0]]
        records = []
        for row in rows[1:]:
            if row is None or all(norm_text(value) == "" for value in row):
                continue
            record = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                record[header] = row[idx] if idx < len(row) else ""
            records.append(record)

        if not records:
            return empty_status_history()

        return normalize_status_history(pd.DataFrame(records))

    finally:
        workbook.close()


def load_embedded_status_history_if_available(file_bytes: bytes, workbook_name: str) -> None:
    """
    Auto-load status history from the workbook itself.

    Manual CSV upload still wins. Embedded workbook history is used only when the user
    has not uploaded a separate CSV during this session.
    """
    if st.session_state.get("uploaded_status_history") is not None:
        return

    signature = f"embedded:{workbook_name}:{len(file_bytes)}"
    if signature == st.session_state.get("embedded_status_history_signature", ""):
        return

    embedded = read_embedded_status_history_from_workbook_bytes(file_bytes)
    st.session_state.embedded_status_history_signature = signature

    if embedded.empty:
        return

    st.session_state.status_history = embedded
    st.session_state.uploaded_status_history_name = f"Embedded in {workbook_name}"
    st.session_state.status_history_load_error = ""


def write_status_history_sheet_to_workbook_bytes(file_bytes: bytes, history_df: pd.DataFrame) -> bytes:
    """Embed status history directly inside the workbook as a hidden sheet."""
    history = normalize_status_history(history_df)
    workbook = load_workbook(io.BytesIO(file_bytes))

    if STATUS_HISTORY_SHEET_NAME in workbook.sheetnames:
        del workbook[STATUS_HISTORY_SHEET_NAME]

    ws = workbook.create_sheet(STATUS_HISTORY_SHEET_NAME)
    ws.sheet_state = "hidden"

    for col_idx, header in enumerate(STATUS_HISTORY_COLUMNS, start=1):
        ws.cell(1, col_idx).value = header

    for row_idx, record in enumerate(history.to_dict("records"), start=2):
        for col_idx, header in enumerate(STATUS_HISTORY_COLUMNS, start=1):
            ws.cell(row_idx, col_idx).value = record.get(header, "")

    out = io.BytesIO()
    workbook.save(out)
    workbook.close()
    out.seek(0)
    return out.getvalue()


def status_streak_for_deal(
    history: pd.DataFrame,
    deal_key: str,
    current_status: str,
) -> Tuple[int, str]:
    deal_key = norm_text(deal_key)
    current_status = norm_text(current_status)
    if history.empty or not deal_key or not current_status:
        return 0, ""

    rows = history[history["deal_key"].map(norm_text) == deal_key].copy()
    if rows.empty:
        return 0, ""

    rows["meeting_date_dt"] = pd.to_datetime(rows["meeting_date"], errors="coerce")
    rows = rows.dropna(subset=["meeting_date_dt"])
    rows = rows.sort_values("meeting_date_dt", ascending=False, kind="stable")

    same_count = 0
    oldest_same_date = ""
    for _, row in rows.iterrows():
        if norm_text(row.get("status")) != current_status:
            break
        same_count += 1
        oldest_same_date = row["meeting_date_dt"].strftime("%m/%d/%Y")

    return same_count, oldest_same_date


def build_status_prompt_message(row: pd.Series) -> str:
    if not bool(row.get("status_needs_update_prompt", False)):
        return ""

    current_update = display_text(row.get("status"), blank="blank")
    count = int(row.get("status_same_meeting_count") or 0)
    since = display_text(row.get("status_same_since"), blank="unknown")
    return (
        f"This deal has carried the same meeting update for {count} meetings since {since}. "
        f"Confirm whether '{current_update}' still reflects the latest story, or refresh the update/commentary."
    )


def add_status_staleness_columns(
    deck: pd.DataFrame,
    status_history: pd.DataFrame,
    meeting_date: dt.date,
    workbook_name: str,
    threshold: int = STATUS_STALE_MEETING_THRESHOLD,
) -> pd.DataFrame:
    if deck.empty:
        return deck.copy()

    combined_history = merge_status_history(status_history, deck, meeting_date, workbook_name)
    out = deck.copy()
    same_counts: List[int] = []
    same_since: List[str] = []

    for _, row in out.iterrows():
        same_count, oldest_date = status_streak_for_deal(
            combined_history,
            deal_key=str(row.get("deal_key", "")),
            current_status=str(row.get("status", "")),
        )
        same_counts.append(same_count)
        same_since.append(oldest_date)

    out["status_same_meeting_count"] = same_counts
    out["status_same_since"] = same_since
    out["status_needs_update_prompt"] = out["status_same_meeting_count"] >= threshold
    out["status_prompt_message"] = out.apply(build_status_prompt_message, axis=1)
    return out


# -----------------------------------------------------------------------------
# Workbook loading: Bridge / Term
# -----------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def load_portfolio_workbook(
    file_bytes: bytes,
    as_of_date_iso: str,
    include_hidden: bool,
) -> Tuple[pd.DataFrame, Dict[str, Dict[str, str]]]:
    as_of_date = pd.Timestamp(as_of_date_iso).normalize()
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    frames: List[pd.DataFrame] = []
    metadata: Dict[str, Dict[str, str]] = {}

    for sheet_name in ["Bridge", "Term"]:
        if sheet_name not in workbook.sheetnames:
            continue

        ws = workbook[sheet_name]
        header_row, last_col = find_header_row(ws)
        headers = [ws.cell(header_row, c).value for c in range(1, last_col + 1)]
        rows: List[Dict[str, object]] = []

        for r in range(header_row + 1, ws.max_row + 1):
            if not include_hidden and ws.row_dimensions[r].hidden:
                continue

            row_dict: Dict[str, object] = {}
            for c in range(1, last_col + 1):
                header = headers[c - 1]
                if header is None:
                    continue
                row_dict[str(header)] = ws.cell(r, c).value

            deal_col_in_row = resolve_column(row_dict.keys(), "Deal Number")
            deal_value = row_dict.get(deal_col_in_row) if deal_col_in_row else None
            if norm_text(deal_value) == "":
                continue

            row_dict["_excel_row"] = r
            rows.append(row_dict)

        if not rows:
            continue

        df = pd.DataFrame(rows)
        latest_columns = latest_dated_by_suffix(headers, as_of_date=as_of_date.date())

        deal_col = resolve_column(df.columns, "Deal Number")
        deal_name_col = resolve_column(df.columns, "Deal Name", "Loan Name", "Property Name")
        status_col = resolve_column(df.columns, "Status")
        commentary_col = resolve_column(df.columns, "AM Commentary")
        owner_col = resolve_column(df.columns, "Point Person", "Asset Manager", "Active RM")
        borrower_col = resolve_column(df.columns, "Borrower Name", "Borrower Entity", "Account Name")
        account_col = resolve_column(df.columns, "Account", "Account Name")
        servicer_col = resolve_column(df.columns, "Servicer")
        portfolio_col = resolve_column(df.columns, "Portfolio")
        segment_col = resolve_column(df.columns, "Segment")
        financing_col = resolve_column(df.columns, "Financing")
        loan_buyer_col = resolve_column(df.columns, "Loan Buyer")
        maturity_col = resolve_column(
            df.columns,
            "Current Maturity Date",
            "Maturity Date",
            "Original Maturity Date",
            "Next Advance Maturity Date",
        )
        next_payment_col = resolve_column(df.columns, "Next Payment Date")

        if sheet_name == "Bridge":
            upb_col = resolve_column(
                df.columns,
                latest_columns.get("UPB"),
                "UPB",
                "Active Funded Amount",
                "Loan Amount",
            )
        else:
            upb_col = resolve_column(df.columns, latest_columns.get("UPB"), "UPB", "Loan Amount")

        npl_col = resolve_column(
            df.columns,
            latest_columns.get("NPL"),
            "NPL",
            "Loan Level Delinquency",
            "DQ Status",
        )
        dpd_cols = matching_columns(df.columns, "Days Past Due")

        deal_series = safe_series(df, deal_col, "").map(norm_text)
        borrower_series = safe_series(df, borrower_col, "")
        account_series = safe_series(df, account_col, "")
        deal_name_series = fallback_name(
            safe_series(df, deal_name_col, ""),
            borrower_series,
            account_series,
            deal_series,
        )

        df["sheet"] = sheet_name
        df["sheet_order"] = 0 if sheet_name == "Bridge" else 1
        df["original_order"] = pd.to_numeric(df["_excel_row"], errors="coerce").fillna(999999).astype(int)
        df["deal_number"] = deal_series
        df["deal_name"] = deal_name_series
        df["borrower"] = borrower_series.map(norm_text)
        df["account_display"] = account_series.map(norm_text)
        df["servicer"] = safe_series(df, servicer_col, "").map(norm_text)
        df["portfolio"] = safe_series(df, portfolio_col, sheet_name).map(norm_text)
        df["segment"] = safe_series(df, segment_col, "").map(norm_text)
        df["financing"] = safe_series(df, financing_col, "").map(norm_text)
        df["loan_buyer"] = safe_series(df, loan_buyer_col, "").map(norm_text)
        df["owner"] = safe_series(df, owner_col, "").map(norm_text)
        df["status"] = safe_series(df, status_col, "").map(norm_text)
        df["commentary"] = safe_series(df, commentary_col, "").map(norm_text)
        df["saved_at"] = ""
        df["deal_key"] = df.apply(
            lambda row: get_override_key(str(row["sheet"]), str(row["deal_number"])),
            axis=1,
        )
        df["upb"] = safe_numeric_series(df, upb_col)
        df["maturity_date"] = safe_datetime_series(df, maturity_col)
        df["next_payment_date"] = safe_datetime_series(df, next_payment_col)
        df["npl_raw"] = safe_series(df, npl_col, "").fillna("").astype(str)

        if dpd_cols:
            dpd_matrix = np.column_stack(
                [
                    pd.to_numeric(df[col], errors="coerce").fillna(0).to_numpy(dtype=float)
                    for col in dpd_cols
                ]
            )
            df["days_past_due"] = np.maximum(dpd_matrix.max(axis=1), 0)
        else:
            df["days_past_due"] = 0

        if sheet_name == "Bridge":
            funded_amount_col = resolve_column(df.columns, "Active Funded Amount")
            commitment_col = resolve_column(df.columns, "Loan Commitment")
            remaining_commitment_col = resolve_column(df.columns, "Remaining Commitment")
            df["funded_amount"] = safe_numeric_series(df, funded_amount_col)
            df["commitment"] = safe_numeric_series(df, commitment_col)
            df["remaining_commitment"] = safe_numeric_series(df, remaining_commitment_col)
            df["loan_amount"] = np.nan
        else:
            loan_amount_col = resolve_column(df.columns, "Loan Amount")
            df["loan_amount"] = safe_numeric_series(df, loan_amount_col)
            df["funded_amount"] = np.nan
            df["commitment"] = np.nan
            df["remaining_commitment"] = np.nan

        metadata[sheet_name] = {
            "upb_header": upb_col or "",
            "maturity_header": maturity_col or "",
            "next_payment_header": next_payment_col or "",
            "owner_header": owner_col or "",
            "status_header": status_col or "",
            "commentary_header": commentary_col or "",
        }
        frames.append(df)

    workbook.close()

    if not frames:
        return pd.DataFrame(), metadata

    deck = pd.concat(frames, ignore_index=True)
    deck["days_to_maturity"] = (deck["maturity_date"] - as_of_date).dt.days
    deck["days_to_next_payment"] = (deck["next_payment_date"] - as_of_date).dt.days
    deck = deck.sort_values(
        ["sheet_order", "original_order"],
        ascending=[True, True],
        kind="stable",
    ).reset_index(drop=True)
    return deck, metadata


# -----------------------------------------------------------------------------
# Loan Modification loading and presentation helpers
# -----------------------------------------------------------------------------
def empty_loan_modifications() -> pd.DataFrame:
    return pd.DataFrame(
        columns=[
            "deal_number_key",
            "Deal Number",
            "Deal Name",
            "Loan Modification Name",
            "Status",
            "Loan Mod Type",
            "Modification Type",
            "Mod Effective Date",
            "Modification Finalized Date",
            "Updated Maturity Date",
            "Comments",
            "Pay Pik Summary",
        ]
    )


def load_workbook_sheet_as_dataframe(
    workbook,
    sheet_name: str,
    key_header: str = "Deal Number",
    desired_columns: Optional[Iterable[str]] = None,
) -> pd.DataFrame:
    if sheet_name not in workbook.sheetnames:
        return pd.DataFrame()

    ws = workbook[sheet_name]
    wanted = {canon_header(col) for col in desired_columns or [] if norm_text(col)}

    header_values: Optional[Tuple[object, ...]] = None
    data_rows: List[Tuple[object, ...]] = []

    for row_idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
        values_tuple = tuple(values or ())

        if header_values is None:
            if any(canon_header(value) == canon_header(key_header) for value in values_tuple if value is not None):
                header_values = values_tuple
            continue

        data_rows.append(values_tuple)

    if header_values is None:
        return pd.DataFrame(columns=list(desired_columns or []))

    headers = [norm_text(value) for value in header_values]
    selected_cols: List[Tuple[int, str]] = []
    for idx, header in enumerate(headers):
        if not header:
            continue
        if wanted and canon_header(header) not in wanted:
            continue
        selected_cols.append((idx, header))

    if not selected_cols:
        return pd.DataFrame(columns=list(desired_columns or []))

    rows: List[Dict[str, object]] = []
    for values in data_rows:
        row_dict: Dict[str, object] = {}
        has_value = False

        for idx, header in selected_cols:
            value = values[idx] if idx < len(values) else None
            row_dict[header] = value
            if not value_is_blank(value):
                has_value = True

        if has_value:
            rows.append(row_dict)

    if not rows:
        return pd.DataFrame(columns=[header for _, header in selected_cols])

    return pd.DataFrame(rows)


def normalize_loan_modifications(raw_df: pd.DataFrame) -> pd.DataFrame:
    if raw_df is None or raw_df.empty:
        return empty_loan_modifications()

    df = raw_df.copy()
    df.columns = [norm_text(col) for col in df.columns]

    deal_col = resolve_column(df.columns, "Deal Number", "Opportunity Deal Loan Number")
    if deal_col is None:
        return empty_loan_modifications()

    if deal_col != "Deal Number":
        df["Deal Number"] = df[deal_col]

    # If the generator already marked which loan mods match the visible overview,
    # keep only those rows. This avoids loading/rendering irrelevant Salesforce history.
    if "Portfolio Overview Match" in df.columns:
        match_series = df["Portfolio Overview Match"].map(norm_text).str.upper()
        if (match_series == "YES").any():
            df = df[match_series == "YES"].copy()

    for col in LOAN_MOD_DATE_COLUMNS:
        if col in df.columns:
            df[col] = coerce_datetime_series(df[col])

    for col in LOAN_MOD_MONEY_COLUMNS + LOAN_MOD_PERCENT_COLUMNS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    text_cols = [
        "Deal Number",
        "Deal Name",
        "Opportunity Deal Name",
        "Loan Modification Name",
        "Status",
        "Loan Mod Type",
        "Modification Type",
        "Mod Reporting Type",
        "Loan Mod Order Number",
        "Cancellation Reason",
        "Comments",
    ]
    for col in text_cols:
        if col in df.columns:
            df[col] = df[col].map(norm_text)

    if "Deal Name" not in df.columns:
        df["Deal Name"] = ""

    if "Opportunity Deal Name" in df.columns:
        df["Deal Name"] = df["Deal Name"].where(
            df["Deal Name"].map(norm_text) != "",
            df["Opportunity Deal Name"].map(norm_text),
        )

    df["deal_number_key"] = df["Deal Number"].map(norm_text)

    # Numeric helper for order-number sorting (falls back to date sorting below).
    if "Loan Mod Order Number" in df.columns:
        df["_mod_order_sort"] = pd.to_numeric(df["Loan Mod Order Number"], errors="coerce")
    else:
        df["_mod_order_sort"] = np.nan

    sort_cols = [
        col
        for col in [
            "_mod_order_sort",
            "System Task Completed Date",
            "Modification Finalized Date",
            "Mod Effective Date",
            "Last Modified Date",
            "Created Date",
        ]
        if col in df.columns
    ]

    if sort_cols:
        # Latest modification first (highest order number / most recent dates).
        df = df.sort_values(
            ["deal_number_key"] + sort_cols,
            ascending=[True] + [False] * len(sort_cols),
            kind="stable",
        )

    return df.reset_index(drop=True)


@st.cache_data(show_spinner=False)
def load_loan_modifications_from_workbook(file_bytes: bytes) -> pd.DataFrame:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return empty_loan_modifications()

    try:
        raw_df = load_workbook_sheet_as_dataframe(
            workbook=workbook,
            sheet_name=LOAN_MOD_SHEET_NAME,
            key_header="Deal Number",
            desired_columns=LOAN_MOD_CORE_COLUMNS,
        )
        return normalize_loan_modifications(raw_df)
    finally:
        workbook.close()


def build_loan_mod_summary_from_record(record: Dict[str, object]) -> str:
    pieces = []

    mod_type = first_nonblank_value(record.get("Loan Mod Type"), record.get("Modification Type"))
    status = record.get("Status")
    effective = fmt_date(record.get("Mod Effective Date"))
    completed = fmt_date(record.get("System Task Completed Date"))
    updated_maturity = fmt_date(record.get("Updated Maturity Date"))

    if display_text(mod_type, blank=""):
        pieces.append(display_text(mod_type))

    if display_text(status, blank=""):
        pieces.append(f"Status: {display_text(status)}")

    if effective != "-":
        pieces.append(f"Effective {effective}")

    if completed != "-":
        pieces.append(f"Completed {completed}")

    if updated_maturity != "-":
        pieces.append(f"New maturity {updated_maturity}")

    return " | ".join(pieces)


def build_loan_mod_commentary_from_record(record: Dict[str, object]) -> str:
    sections = []

    for label in [
        "Comments",
    ]:
        value = display_text(record.get(label), blank="")
        if value:
            sections.append(f"{label}: {value}")

    return "\n\n".join(sections)


def loan_mods_for_deal(loan_mods: pd.DataFrame, deal_number: object) -> pd.DataFrame:
    """Return every loan modification record for a single deal, latest first.

    This is what powers the presentation toggle, so the caller can step through
    each modification on a loan rather than seeing only the most recent one.
    """
    if loan_mods is None or loan_mods.empty or "deal_number_key" not in loan_mods.columns:
        return pd.DataFrame()

    key = norm_text(deal_number)
    if not key:
        return pd.DataFrame()

    subset = loan_mods[loan_mods["deal_number_key"] == key].copy()
    return subset.reset_index(drop=True)


def build_loan_mod_toggle_labels(mods: pd.DataFrame) -> List[str]:
    """Short, unique labels for toggling between each modification of one deal."""
    labels: List[str] = []
    seen: Dict[str, int] = {}

    for position, (_, record) in enumerate(mods.iterrows(), start=1):
        order = display_text(record.get("Loan Mod Order Number"), blank="")
        base = f"Mod #{order}" if order else f"Mod {position}"

        # Guarantee uniqueness so the toggle control always has distinct options.
        if base in seen:
            seen[base] += 1
            base = f"{base} ({seen[base]})"
        else:
            seen[base] = 1

        labels.append(base)

    return labels


def add_loan_modification_columns(deck: pd.DataFrame, loan_mods: pd.DataFrame) -> pd.DataFrame:
    if deck.empty:
        return deck.copy()

    out = deck.copy()

    defaults = {
        "loan_mod_count": 0,
        "loan_mod_latest_summary": "",
        "loan_mod_latest_status": "",
        "loan_mod_latest_type": "",
        "loan_mod_latest_name": "",
        "loan_mod_latest_order_number": "",
        "loan_mod_latest_effective_date": pd.NaT,
        "loan_mod_latest_completed_date": pd.NaT,
        "loan_mod_latest_finalized_date": pd.NaT,
        "loan_mod_latest_previous_maturity_date": pd.NaT,
        "loan_mod_latest_updated_maturity_date": pd.NaT,
        "loan_mod_latest_previous_commitment": np.nan,
        "loan_mod_latest_updated_commitment": np.nan,
        "loan_mod_latest_comments": "",
    }

    for col, default_value in defaults.items():
        out[col] = default_value

    if loan_mods is None or loan_mods.empty or "deal_number_key" not in loan_mods.columns:
        return out

    count_by_deal = loan_mods["deal_number_key"].value_counts().to_dict()
    latest_by_deal = loan_mods.drop_duplicates("deal_number_key", keep="first").set_index("deal_number_key")

    for idx, row in out.iterrows():
        deal_number = norm_text(row.get("deal_number"))
        if not deal_number or deal_number not in latest_by_deal.index:
            continue

        latest = latest_by_deal.loc[deal_number]
        latest_dict = latest.to_dict()
        latest_type = first_nonblank_value(
            latest_dict.get("Loan Mod Type"),
            latest_dict.get("Modification Type"),
            latest_dict.get("Mod Reporting Type"),
        )

        out.at[idx, "loan_mod_count"] = int(count_by_deal.get(deal_number, 0))
        out.at[idx, "loan_mod_latest_summary"] = build_loan_mod_summary_from_record(latest_dict)
        out.at[idx, "loan_mod_latest_status"] = display_text(latest_dict.get("Status"), blank="")
        out.at[idx, "loan_mod_latest_type"] = display_text(latest_type, blank="")
        out.at[idx, "loan_mod_latest_name"] = display_text(latest_dict.get("Loan Modification Name"), blank="")
        out.at[idx, "loan_mod_latest_order_number"] = display_text(latest_dict.get("Loan Mod Order Number"), blank="")
        out.at[idx, "loan_mod_latest_effective_date"] = coerce_datetime_value(latest_dict.get("Mod Effective Date"))
        out.at[idx, "loan_mod_latest_completed_date"] = coerce_datetime_value(latest_dict.get("System Task Completed Date"))
        out.at[idx, "loan_mod_latest_finalized_date"] = coerce_datetime_value(latest_dict.get("Modification Finalized Date"))
        out.at[idx, "loan_mod_latest_previous_maturity_date"] = coerce_datetime_value(latest_dict.get("Previous Maturity Date"))
        out.at[idx, "loan_mod_latest_updated_maturity_date"] = coerce_datetime_value(latest_dict.get("Updated Maturity Date"))
        out.at[idx, "loan_mod_latest_previous_commitment"] = latest_dict.get("Previous Loan Commitment")
        out.at[idx, "loan_mod_latest_updated_commitment"] = latest_dict.get("Updated Loan Commitment")
        out.at[idx, "loan_mod_latest_comments"] = display_text(latest_dict.get("Comments"), blank="")

    return out


def has_loan_mods(row: pd.Series) -> bool:
    try:
        return int(row.get("loan_mod_count") or 0) > 0
    except Exception:
        return False


def compact_note(value: object, max_chars: int = 360) -> str:
    text = display_text(value, blank="")
    if not text:
        return ""
    text = re.sub(r"^Comments\s*:\s*", "", text, flags=re.IGNORECASE).strip()
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip() + "…"


def build_loan_mod_spotlight_dataframe(deck: pd.DataFrame) -> pd.DataFrame:
    if deck.empty or "loan_mod_count" not in deck.columns:
        return pd.DataFrame()

    source = deck[pd.to_numeric(deck["loan_mod_count"], errors="coerce").fillna(0).astype(int) > 0].copy()
    if source.empty:
        return pd.DataFrame()

    rows = []
    for _, row in source.iterrows():
        rows.append(
            {
                "Type": display_text(row.get("sheet")),
                "Deal Name": display_text(row.get("deal_name")),
                "Deal #": display_text(row.get("deal_number")),
                "Mod Count": fmt_int(row.get("loan_mod_count")),
                "Mod Order #": display_text(row.get("loan_mod_latest_order_number")),
                "Modification Type": display_text(row.get("loan_mod_latest_type")),
                "Status": display_text(row.get("loan_mod_latest_status")),
                "Completed": fmt_date(row.get("loan_mod_latest_completed_date")),
                "Effective": fmt_date(row.get("loan_mod_latest_effective_date")),
                "Previous Maturity": fmt_date(row.get("loan_mod_latest_previous_maturity_date")),
                "Updated Maturity": fmt_date(row.get("loan_mod_latest_updated_maturity_date")),
                "Previous Commitment": fmt_money(row.get("loan_mod_latest_previous_commitment"), decimals=0),
                "Updated Commitment": fmt_money(row.get("loan_mod_latest_updated_commitment"), decimals=0),
                "Comments": compact_note(row.get("loan_mod_latest_comments")),
            }
        )

    return pd.DataFrame(rows)


# -----------------------------------------------------------------------------
# App state / overrides
# -----------------------------------------------------------------------------
def available_status_suggestions(deck: pd.DataFrame) -> List[str]:
    if "status" not in deck.columns:
        return COMMON_STATUS_SUGGESTIONS
    existing = sorted({norm_text(v) for v in deck["status"].dropna().tolist() if norm_text(v)})
    return sorted(set(COMMON_STATUS_SUGGESTIONS + existing))


def initialize_state() -> None:
    defaults = {
        "view_mode": "Presentation",
        "sheet_filter": "All",
        "search_query": "",
        "include_hidden": False,
        "stale_only": False,
        "as_of_date": dt.date.today(),
        "selected_deal_key": None,
        "meeting_overrides": {},
        "review_flags": {},
        "dialog_target": None,
        "status_history": empty_status_history(),
        "status_history_default_checked": False,
        "uploaded_status_history_name": "",
        "uploaded_status_history_signature": "",
        "status_history_load_error": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def ensure_override_store() -> Dict[str, Dict[str, str]]:
    if "meeting_overrides" not in st.session_state:
        st.session_state.meeting_overrides = {}
    return st.session_state.meeting_overrides


def ensure_review_flags() -> Dict[str, bool]:
    if "review_flags" not in st.session_state:
        st.session_state.review_flags = {}
    return st.session_state.review_flags


def mark_current_deal_reviewed(deal_key: str) -> None:
    flags = ensure_review_flags()
    flags[deal_key] = True


def upsert_override(
    base_row: pd.Series,
    status: str,
    owner: str,
    commentary: str,
) -> None:
    overrides = ensure_override_store()
    key = get_override_key(str(base_row.get("sheet")), str(base_row.get("deal_number")))
    payload = {
        "sheet": str(base_row.get("sheet")),
        "deal_number": str(base_row.get("deal_number")),
        "status": norm_text(status),
        "owner": norm_text(owner),
        "commentary": str(commentary or "").strip(),
        "saved_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    original_status = norm_text(base_row.get("status"))
    original_owner = norm_text(base_row.get("owner"))
    original_commentary = str(base_row.get("commentary") or "")

    if (
        payload["status"] == original_status
        and payload["owner"] == original_owner
        and payload["commentary"] == original_commentary
    ):
        overrides.pop(key, None)
    else:
        overrides[key] = payload


def apply_overrides(deck: pd.DataFrame, overrides: Dict[str, Dict[str, str]]) -> pd.DataFrame:
    if deck.empty:
        return deck.copy()

    out = deck.copy()
    if "saved_at" not in out.columns:
        out["saved_at"] = ""

    if not overrides:
        return out

    override_df = pd.DataFrame(overrides.values())
    if override_df.empty:
        return out

    for field in ["sheet", "deal_number", "status", "owner", "commentary", "saved_at"]:
        if field not in override_df.columns:
            override_df[field] = ""

    merged = out.merge(
        override_df[["sheet", "deal_number", "status", "owner", "commentary", "saved_at"]],
        on=["sheet", "deal_number"],
        how="left",
        suffixes=("", "_override"),
    )

    for field in ["status", "owner", "commentary", "saved_at"]:
        override_field = f"{field}_override"
        merged[field] = merged[override_field].where(
            merged[override_field].notna() & (merged[override_field].astype(str) != ""),
            merged[field],
        )

    drop_cols = [col for col in merged.columns if col.endswith("_override")]
    return merged.drop(columns=drop_cols)


def apply_filters(deck: pd.DataFrame) -> pd.DataFrame:
    filtered = deck.copy()

    if st.session_state.sheet_filter == "Bridge":
        filtered = filtered[filtered["sheet"] == "Bridge"].copy()
    elif st.session_state.sheet_filter == "Term":
        filtered = filtered[filtered["sheet"] == "Term"].copy()

    query = st.session_state.search_query.strip().upper()
    if query:
        mask = (
            filtered["deal_number"].astype(str).str.upper().str.contains(query, na=False)
            | filtered["deal_name"].astype(str).str.upper().str.contains(query, na=False)
            | filtered["borrower"].astype(str).str.upper().str.contains(query, na=False)
        )

        if "loan_mod_latest_summary" in filtered.columns:
            mask = mask | filtered["loan_mod_latest_summary"].astype(str).str.upper().str.contains(query, na=False)

        if "loan_mod_latest_comments" in filtered.columns:
            mask = mask | filtered["loan_mod_latest_comments"].astype(str).str.upper().str.contains(query, na=False)

        filtered = filtered[mask].copy()

    if st.session_state.stale_only and "status_needs_update_prompt" in filtered.columns:
        filtered = filtered[filtered["status_needs_update_prompt"] == True].copy()  # noqa: E712

    return filtered.sort_values(
        ["sheet_order", "original_order"],
        ascending=[True, True],
        kind="stable",
    ).reset_index(drop=True)


def sync_selected_deal(deck: pd.DataFrame) -> None:
    if deck.empty:
        st.session_state.selected_deal_key = None
        return
    valid_keys = deck["deal_key"].tolist()
    if st.session_state.selected_deal_key not in valid_keys:
        st.session_state.selected_deal_key = valid_keys[0]


def get_selected_row(deck: pd.DataFrame) -> pd.Series:
    sync_selected_deal(deck)
    selected_key = st.session_state.selected_deal_key
    match = deck[deck["deal_key"] == selected_key]
    if match.empty:
        return deck.iloc[0]
    return match.iloc[0]


def get_row_by_key(deck: pd.DataFrame, deal_key: str) -> Optional[pd.Series]:
    match = deck[deck["deal_key"] == deal_key]
    if match.empty:
        return None
    return match.iloc[0]


def move_selection(deck: pd.DataFrame, step: int) -> None:
    sync_selected_deal(deck)
    keys = deck["deal_key"].tolist()
    current_key = st.session_state.selected_deal_key
    idx = keys.index(current_key)
    new_idx = min(max(idx + step, 0), len(keys) - 1)
    st.session_state.selected_deal_key = keys[new_idx]


def build_presenter_prompts(row: pd.Series) -> List[str]:
    prompts: List[str] = []

    if bool(row.get("status_needs_update_prompt", False)):
        prompts.append(str(row.get("status_prompt_message") or "Review whether the meeting update should be refreshed."))

    if has_loan_mods(row):
        prompts.append(
            f"Review loan modification details: {display_text(row.get('loan_mod_latest_summary'), blank='latest modification on file')}."
        )

    maturity_text = fmt_date(row.get("maturity_date"))
    payment_text = fmt_date(row.get("next_payment_date"))

    if maturity_text != "-":
        prompts.append(f"Confirm maturity timing: {maturity_text} ({fmt_day_delta(row.get('days_to_maturity'))}).")

    if payment_text != "-":
        prompts.append(f"Confirm next payment timing: {payment_text} ({fmt_day_delta(row.get('days_to_next_payment'))}).")

    commentary_text = display_text(row.get("commentary"), blank="")
    if commentary_text:
        prompts.append("Decide whether AM commentary needs a live refresh before the meeting ends.")

    if not prompts:
        prompts.append("Confirm the latest story, next milestone, and any follow-up owner.")

    return prompts[:4]


# -----------------------------------------------------------------------------
# Dataframes for tables and exports
# -----------------------------------------------------------------------------
def build_agenda_dataframe(deck: pd.DataFrame) -> pd.DataFrame:
    queue = deck.copy()
    queue.insert(0, "Agenda #", range(1, len(queue) + 1))
    queue["UPB"] = queue["upb"].map(lambda x: fmt_money(x, decimals=0))
    queue["Maturity"] = queue["maturity_date"].map(fmt_date)
    queue["Next Payment"] = queue["next_payment_date"].map(fmt_date)
    queue["Status"] = queue["status"].map(display_text)
    queue["Owner"] = queue["owner"].map(display_text)
    queue["Update Reminder"] = queue["status_needs_update_prompt"].map(lambda value: "Yes" if bool(value) else "")

    if "loan_mod_count" not in queue.columns:
        queue["loan_mod_count"] = 0
    if "loan_mod_latest_summary" not in queue.columns:
        queue["loan_mod_latest_summary"] = ""

    queue["Loan Mods"] = queue["loan_mod_count"].map(lambda value: fmt_int(value) if int(value or 0) else "")
    queue["Latest Loan Mod"] = queue["loan_mod_latest_summary"].map(lambda value: display_text(value, blank=""))

    return queue[
        [
            "Agenda #",
            "sheet",
            "deal_name",
            "deal_number",
            "Update Reminder",
            "Loan Mods",
            "Latest Loan Mod",
            "Status",
            "Owner",
            "Maturity",
            "Next Payment",
            "UPB",
        ]
    ].rename(
        columns={
            "sheet": "Type",
            "deal_name": "Deal Name",
            "deal_number": "Deal #",
        }
    )


def build_review_dataframe(deck: pd.DataFrame) -> pd.DataFrame:
    review = deck.copy().set_index("deal_key")
    review.insert(0, "Agenda #", range(1, len(review) + 1))
    review["Type"] = review["sheet"]
    review["Deal Name"] = review["deal_name"].map(display_text)
    review["Deal #"] = review["deal_number"].map(display_text)
    review["Maturity"] = review["maturity_date"].map(fmt_date)
    review["Next Payment"] = review["next_payment_date"].map(fmt_date)
    review["UPB"] = review["upb"].map(lambda x: fmt_money(x, decimals=0))
    review["Update Reminder"] = review["status_needs_update_prompt"].map(lambda value: "Yes" if bool(value) else "")
    review["Status"] = review["status"].fillna("").astype(str)
    review["Owner"] = review["owner"].fillna("").astype(str)
    review["AM Commentary"] = review["commentary"].fillna("").astype(str)

    return review[
        [
            "Agenda #",
            "Type",
            "Deal Name",
            "Deal #",
            "Maturity",
            "Next Payment",
            "UPB",
            "Update Reminder",
            "Status",
            "Owner",
            "AM Commentary",
        ]
    ]


def apply_review_edits(
    edited_df: pd.DataFrame,
    current_deck: pd.DataFrame,
    raw_deck: pd.DataFrame,
) -> None:
    current_lookup = current_deck.set_index("deal_key")
    raw_lookup = raw_deck.set_index("deal_key")

    for deal_key, edited_row in edited_df.iterrows():
        if deal_key not in current_lookup.index or deal_key not in raw_lookup.index:
            continue

        current_row = current_lookup.loc[deal_key]
        raw_row = raw_lookup.loc[deal_key]
        new_status = norm_text(edited_row.get("Status", ""))
        new_owner = norm_text(edited_row.get("Owner", ""))
        new_commentary = str(edited_row.get("AM Commentary", "") or "").strip()

        if (
            new_status != norm_text(current_row.get("status"))
            or new_owner != norm_text(current_row.get("owner"))
            or new_commentary != str(current_row.get("commentary") or "")
        ):
            upsert_override(raw_row, new_status, new_owner, new_commentary)


def build_updates_dataframe(raw_deck: pd.DataFrame, deck: pd.DataFrame) -> pd.DataFrame:
    overrides = ensure_override_store()
    if not overrides:
        return pd.DataFrame()

    changed_keys = {
        get_override_key(item["sheet"], item["deal_number"])
        for item in overrides.values()
    }

    columns = [
        "deal_key",
        "sheet",
        "deal_name",
        "deal_number",
        "status",
        "owner",
        "commentary",
        "saved_at",
        "status_needs_update_prompt",
    ]
    out = deck[columns].copy()
    out = out[out["deal_key"].isin(changed_keys)].copy()
    if out.empty:
        return pd.DataFrame()

    out = out.rename(
        columns={
            "deal_key": "Deal Key",
            "sheet": "Type",
            "deal_name": "Deal Name",
            "deal_number": "Deal #",
            "status": "Status",
            "owner": "Owner",
            "commentary": "AM Commentary",
            "saved_at": "Last Saved",
            "status_needs_update_prompt": "Update Reminder",
        }
    )
    out["Update Reminder"] = out["Update Reminder"].map(lambda value: "Yes" if bool(value) else "")

    return out[
        [
            "Type",
            "Deal Name",
            "Deal #",
            "Update Reminder",
            "Status",
            "Owner",
            "AM Commentary",
            "Last Saved",
        ]
    ]


def export_overrides_csv(raw_deck: pd.DataFrame, deck: pd.DataFrame) -> bytes:
    export_df = build_updates_dataframe(raw_deck, deck)
    if export_df.empty:
        return b""
    return export_df.to_csv(index=False).encode("utf-8")


def update_workbook_bytes(file_bytes: bytes, overrides: Dict[str, Dict[str, str]]) -> bytes:
    if not overrides:
        return file_bytes

    workbook = load_workbook(io.BytesIO(file_bytes))

    for sheet_name, ws in [(name, workbook[name]) for name in workbook.sheetnames if name in {"Bridge", "Term"}]:
        relevant = [item for item in overrides.values() if item.get("sheet") == sheet_name]
        if not relevant:
            continue

        header_row, last_col = find_header_row(ws)
        headers = [ws.cell(header_row, c).value for c in range(1, last_col + 1)]
        key_col = next(
            c
            for c, header in enumerate(headers, start=1)
            if canon_header(header) == canon_header("Deal Number")
        )
        status_col = next(
            (c for c, header in enumerate(headers, start=1) if canon_header(header) == canon_header("Status")),
            None,
        )
        owner_col = next(
            (
                c
                for c, header in enumerate(headers, start=1)
                if canon_header(header) in {canon_header("Point Person"), canon_header("Asset Manager")}
            ),
            None,
        )
        commentary_col = next(
            (c for c, header in enumerate(headers, start=1) if canon_header(header) == canon_header("AM Commentary")),
            None,
        )

        row_lookup: Dict[str, int] = {}
        for r in range(header_row + 1, ws.max_row + 1):
            deal_number = norm_text(ws.cell(r, key_col).value)
            if deal_number:
                row_lookup[deal_number] = r

        for item in relevant:
            target_row = row_lookup.get(norm_text(item.get("deal_number")))
            if not target_row:
                continue
            if status_col is not None:
                ws.cell(target_row, status_col).value = item.get("status", "")
            if owner_col is not None:
                ws.cell(target_row, owner_col).value = item.get("owner", "")
            if commentary_col is not None:
                ws.cell(target_row, commentary_col).value = item.get("commentary", "")

    out = io.BytesIO()
    workbook.save(out)
    workbook.close()
    out.seek(0)
    return out.getvalue()


def process_status_history_upload() -> None:
    uploaded_history = st.session_state.get("uploaded_status_history")
    if uploaded_history is None:
        return

    signature = f"{uploaded_history.name}:{uploaded_history.size}"
    if signature == st.session_state.get("uploaded_status_history_signature", ""):
        return

    try:
        uploaded_history.seek(0)
        history = read_status_history_csv(uploaded_history)
        st.session_state.status_history = history
        st.session_state.uploaded_status_history_name = uploaded_history.name
        st.session_state.uploaded_status_history_signature = signature
        st.session_state.status_history_load_error = ""
    except Exception as exc:
        st.session_state.status_history_load_error = str(exc)


# -----------------------------------------------------------------------------
# Standard Streamlit controls and overview
# -----------------------------------------------------------------------------
def render_controls_bar(workbook_name: str) -> None:
    with st.sidebar:
        st.title(APP_TITLE)
        st.caption("Meeting navigation, workbook controls, and filters")
        st.caption(f"Workbook: {workbook_name}")

        st.divider()
        st.radio(
            "View",
            options=["Overview", "Presentation", "Review", "Exports"],
            key="view_mode",
            horizontal=False,
        )

        st.divider()
        st.subheader("Workbook")
        st.file_uploader("Upload workbook", type=["xlsx"], key="uploaded_workbook")
        st.file_uploader("Upload status history CSV backup", type=["csv"], key="uploaded_status_history")
        st.date_input("As-of date", key="as_of_date")
        st.toggle("Include hidden rows", key="include_hidden")

        st.divider()
        st.subheader("Filters")
        st.text_input("Search deal # / name / borrower / loan mod", key="search_query")
        st.toggle("Update reminders only", key="stale_only")

        if st.button("Reset filters", use_container_width=True):
            st.session_state.sheet_filter = "All"
            st.session_state.search_query = ""
            st.session_state.stale_only = False
            st.rerun()


def render_status_history_notice(deck: pd.DataFrame) -> None:
    history = normalize_status_history(st.session_state.status_history)
    history_name = display_text(st.session_state.get("uploaded_status_history_name", ""), blank="None loaded")
    warning_count = int(deck["status_needs_update_prompt"].sum()) if "status_needs_update_prompt" in deck.columns else 0

    c1, c2, c3 = st.columns(3)
    c1.metric("History rows", fmt_int(len(history)))
    c2.metric("History source", history_name)
    c3.metric("Update reminders", fmt_int(warning_count))

    if st.session_state.get("status_history_load_error"):
        st.error(f"Could not load status history: {st.session_state.status_history_load_error}")

    if history.empty:
        st.info(
            "No embedded status history is loaded yet. The app will still run; update reminders begin once "
            "you use a workbook generated by the weekly overview builder with the hidden history sheet."
        )


def render_overview_view(deck: pd.DataFrame, as_of_date: dt.date) -> None:
    total_upb = deck["upb"].fillna(0).sum()
    bridge_count = int((deck["sheet"] == "Bridge").sum())
    term_count = int((deck["sheet"] == "Term").sum())
    next_30 = int((deck["days_to_maturity"].fillna(9999) <= 30).sum())
    stale_count = int(deck["status_needs_update_prompt"].sum())

    if "loan_mod_count" in deck.columns:
        loan_mod_count_series = pd.to_numeric(deck["loan_mod_count"], errors="coerce").fillna(0)
    else:
        loan_mod_count_series = pd.Series([0] * len(deck), index=deck.index)
    loan_mod_deal_count = int((loan_mod_count_series.astype(int) > 0).sum())

    m1, m2, m3, m4, m5, m6 = st.columns(6)
    m1.metric("Agenda items", fmt_int(len(deck)))
    m2.metric("Total UPB", fmt_money(total_upb, decimals=0))
    m3.metric("Bridge / Term", f"{fmt_int(bridge_count)} / {fmt_int(term_count)}")
    m4.metric("Maturing in 30d", fmt_int(next_30), as_of_date.strftime("%m/%d/%Y"))
    m5.metric("Update reminders", fmt_int(stale_count))
    m6.metric("Deals w/ loan mods", fmt_int(loan_mod_deal_count))

    render_status_history_notice(deck)

    loan_mod_spotlight = build_loan_mod_spotlight_dataframe(deck)
    if not loan_mod_spotlight.empty:
        st.subheader("Loan modification spotlight")
        st.caption("Prioritizes the highlighted fields from the Loan Modifications sheet. Shows the latest mod per deal; use Presentation to toggle through every modification.")
        st.dataframe(
            loan_mod_spotlight,
            hide_index=True,
            use_container_width=True,
            column_config={
                "Comments": st.column_config.TextColumn("Comments", width="large"),
                "Modification Type": st.column_config.TextColumn("Modification Type", width="medium"),
            },
        )

    if stale_count:
        stale = deck[deck["status_needs_update_prompt"] == True].copy()  # noqa: E712
        st.warning(f"{stale_count} deal(s) have an update reminder before the meeting is closed.")
        st.dataframe(
            stale[[
                "sheet",
                "deal_name",
                "deal_number",
                "status",
                "owner",
            ]].rename(
                columns={
                    "sheet": "Type",
                    "deal_name": "Deal Name",
                    "deal_number": "Deal #",
                    "status": "Status",
                    "owner": "Owner",
                }
            ),
            hide_index=True,
            use_container_width=True,
        )

    st.info(
        "Overview stays in original workbook order. Use Presentation for meeting discussion "
        "and Review for controlled update/commentary edits."
    )

    left, right = st.columns([1.15, 0.85], gap="large")
    with left:
        st.subheader("Agenda queue")
        st.dataframe(build_agenda_dataframe(deck), hide_index=True, use_container_width=True)

    with right:
        st.subheader("Upcoming timing")
        timing = deck[
            [
                "sheet",
                "deal_name",
                "maturity_date",
                "days_to_maturity",
                "next_payment_date",
                "days_to_next_payment",
            ]
        ].copy()
        timing["Maturity"] = timing["maturity_date"].map(fmt_date)
        timing["Maturity Timing"] = timing["days_to_maturity"].map(fmt_day_delta)
        timing["Next Payment"] = timing["next_payment_date"].map(fmt_date)
        timing["Payment Timing"] = timing["days_to_next_payment"].map(fmt_day_delta)
        st.dataframe(
            timing[["sheet", "deal_name", "Maturity", "Maturity Timing", "Next Payment", "Payment Timing"]].rename(
                columns={"sheet": "Type", "deal_name": "Deal Name"}
            ),
            hide_index=True,
            use_container_width=True,
        )


# -----------------------------------------------------------------------------
# Edit dialog
# -----------------------------------------------------------------------------
def maybe_open_edit_dialog(current_deck: pd.DataFrame, raw_deck: pd.DataFrame) -> None:
    target = st.session_state.get("dialog_target")
    if not target:
        return

    current_row = get_row_by_key(current_deck, target)
    raw_row = get_row_by_key(raw_deck, target)
    if current_row is None or raw_row is None:
        st.session_state.dialog_target = None
        return

    @st.dialog("Update current deal", width="large")
    def edit_dialog(current_row_dict: Dict[str, object], raw_row_dict: Dict[str, object]) -> None:
        current_row_local = pd.Series(current_row_dict)
        raw_row_local = pd.Series(raw_row_dict)
        deal_key = str(current_row_local.get("deal_key"))

        st.caption(
            f"{display_text(current_row_local.get('deal_name'))} | "
            f"Deal {display_text(current_row_local.get('deal_number'))}"
        )

        if bool(current_row_local.get("status_needs_update_prompt", False)):
            st.warning(str(current_row_local.get("status_prompt_message") or "Update reminder."))

        with st.form(f"edit-form::{deal_key}"):
            c1, c2 = st.columns(2)
            with c1:
                status = st.text_input("Status", value=str(current_row_local.get("status") or ""))
            with c2:
                owner_label = "Point Person" if current_row_local.get("sheet") == "Bridge" else "Asset Manager"
                owner = st.text_input(owner_label, value=str(current_row_local.get("owner") or ""))

            commentary = st.text_area(
                "AM Commentary",
                value=str(current_row_local.get("commentary") or ""),
                height=240,
            )

            save_clicked = st.form_submit_button("Save changes", type="primary")

            if save_clicked:
                upsert_override(raw_row_local, status, owner, commentary)
                st.session_state.dialog_target = None
                st.rerun()

        if st.button("Cancel", use_container_width=True):
            st.session_state.dialog_target = None
            st.rerun()

    edit_dialog(current_row.to_dict(), raw_row.to_dict())


# -----------------------------------------------------------------------------
# Presentation HTML helpers
# -----------------------------------------------------------------------------
def html_safe(value: object, blank: str = "-") -> str:
    return html.escape(display_text(value, blank=blank), quote=True)


PRESENTATION_CSS = """
<style>
    .rt-hero {
        border: 1px solid rgba(148, 163, 184, 0.28);
        border-radius: 28px;
        padding: 1.28rem 1.38rem;
        background: linear-gradient(135deg, #ffffff 0%, #f7f9fc 48%, #eef4f8 100%);
        box-shadow: 0 18px 48px rgba(30, 41, 59, 0.10);
        margin: 0.70rem 0 0.88rem 0;
    }
    .rt-hero-top { display: flex; align-items: flex-start; justify-content: space-between; gap: 1.1rem; }
    .rt-eyebrow {
        display: inline-flex; border-radius: 999px; background: rgba(71, 85, 105, 0.08);
        color: #475569; border: 1px solid rgba(100, 116, 139, 0.16);
        padding: 0.34rem 0.72rem; font-size: 0.72rem; font-weight: 900;
        letter-spacing: 0.08em; text-transform: uppercase;
    }
    .rt-hero-title {
        max-width: 980px; font-size: clamp(1.45rem, 2.25vw, 2.75rem);
        line-height: 1.05; font-weight: 940; margin: 0.52rem 0 0.34rem 0;
        color: #111827; letter-spacing: -0.04em;
    }
    .rt-hero-subtitle { font-size: 1.02rem; color: #475569; font-weight: 720; line-height: 1.36; margin-bottom: 0.9rem; }
    .rt-chip-row { display: flex; flex-wrap: wrap; gap: 0.45rem; }
    .rt-chip {
        border: 1px solid rgba(100, 116, 139, 0.16); background: rgba(255, 255, 255, 0.72);
        color: #334155; border-radius: 999px; padding: 0.38rem 0.68rem; font-size: 0.82rem; font-weight: 820;
    }
    .rt-status-pill {
        border-radius: 22px; padding: 0.78rem 0.9rem; min-width: 190px; text-align: right;
        border: 1px solid rgba(100, 116, 139, 0.18); background: rgba(255, 255, 255, 0.76);
        box-shadow: 0 12px 24px rgba(30, 41, 59, 0.08);
    }
    .rt-status-pill .rt-label { font-size: 0.68rem; letter-spacing: 0.09em; text-transform: uppercase; color: #64748b; font-weight: 920; }
    .rt-status-pill .rt-value { margin-top: 0.18rem; font-size: 1.11rem; font-weight: 930; color: #111827; overflow-wrap: anywhere; }
    .rt-status-warning { background: rgba(255, 251, 235, 0.88); border-color: rgba(217, 119, 6, 0.28); }
    .rt-status-ok { background: rgba(255, 255, 255, 0.76); border-color: rgba(100, 116, 139, 0.18); }
    .rt-kpi-grid { display: grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap: 0.74rem; margin: 0.5rem 0 0.95rem 0; }
    .rt-kpi {
        border: 1px solid rgba(100, 116, 139, 0.16); border-radius: 22px; background: linear-gradient(180deg, #ffffff, #f8fafc);
        box-shadow: 0 12px 26px rgba(30, 41, 59, 0.07); padding: 1.02rem 1.04rem 0.94rem 1.04rem; min-height: 124px;
    }
    .rt-kpi-label { color: #64748b; font-size: 0.82rem; font-weight: 950; letter-spacing: 0.08em; text-transform: uppercase; margin-bottom: 0.40rem; }
    .rt-kpi-value { color: #111827; font-size: clamp(1.34rem, 1.85vw, 2.12rem); font-weight: 960; line-height: 1.02; overflow-wrap: anywhere; }
    .rt-kpi-helper { margin-top: 0.50rem; color: #64748b; font-size: 0.98rem; font-weight: 900; display: inline-flex; width: fit-content; border-radius: 999px; padding: 0.20rem 0.54rem; background: rgba(100, 116, 139, 0.08); }
    .rt-signal-green { color: #166534; background: #dcfce7; border: 1px solid rgba(22, 101, 52, 0.18); }
    .rt-signal-red { color: #991b1b; background: #fee2e2; border: 1px solid rgba(153, 27, 27, 0.18); }
    .rt-signal-amber { color: #92400e; background: #fef3c7; border: 1px solid rgba(146, 64, 14, 0.18); }
    .rt-signal-muted { color: #64748b; background: rgba(100, 116, 139, 0.08); border: 1px solid rgba(100, 116, 139, 0.12); }
    .rt-panel {
        border: 1px solid rgba(100, 116, 139, 0.16); border-radius: 26px; background: linear-gradient(180deg, #ffffff, #f8fafc);
        box-shadow: 0 14px 30px rgba(30, 41, 59, 0.065); padding: 1.02rem 1.05rem; margin-bottom: 0.85rem;
    }
    .rt-panel-header { display: flex; align-items: center; justify-content: space-between; gap: 0.75rem; margin-bottom: 0.82rem; padding-bottom: 0.68rem; border-bottom: 1px solid rgba(100, 116, 139, 0.14); }
    .rt-panel-title { font-size: 1.14rem; font-weight: 950; color: #111827; letter-spacing: -0.025em; }
    .rt-panel-note { border-radius: 999px; background: rgba(100, 116, 139, 0.09); padding: 0.28rem 0.6rem; color: #475569; font-size: 0.75rem; font-weight: 850; }
    .rt-field-grid { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 0.72rem; }
    .rt-field { border-radius: 18px; background: rgba(241, 245, 249, 0.72); border: 1px solid rgba(100, 116, 139, 0.12); padding: 0.72rem 0.76rem; min-height: 78px; }
    .rt-field-label { font-size: 0.68rem; color: #64748b; font-weight: 920; text-transform: uppercase; letter-spacing: 0.07em; margin-bottom: 0.24rem; }
    .rt-field-value { font-size: 1.04rem; font-weight: 950; color: #111827; line-height: 1.14; overflow-wrap: anywhere; }
    .rt-status-card { border-radius: 24px; padding: 0.98rem 1.02rem; margin-bottom: 0.85rem; border: 1px solid rgba(217, 119, 6, 0.22); background: linear-gradient(135deg, #fffbeb, #ffffff); box-shadow: 0 12px 26px rgba(30, 41, 59, 0.055); }
    .rt-status-card-title { font-size: 1.04rem; font-weight: 950; color: #111827; margin-bottom: 0.30rem; }
    .rt-status-card-body { color: #475569; font-size: 0.93rem; font-weight: 700; line-height: 1.38; }
    .rt-prompt-list { display: flex; flex-direction: column; gap: 0.56rem; }
    .rt-prompt { display: grid; grid-template-columns: 34px minmax(0, 1fr); gap: 0.64rem; border-radius: 18px; border: 1px solid rgba(100, 116, 139, 0.14); background: rgba(248, 250, 252, 0.96); padding: 0.72rem 0.74rem; }
    .rt-prompt-num { width: 32px; height: 32px; display: inline-flex; align-items: center; justify-content: center; border-radius: 12px; background: #dbe6ee; color: #334155; font-weight: 950; font-size: 0.90rem; }
    .rt-prompt-text { color: #334155; font-weight: 760; line-height: 1.30; font-size: 0.95rem; }
    .rt-commentary { border-radius: 20px; border: 1px solid rgba(100, 116, 139, 0.14); background: rgba(255, 255, 255, 0.90); padding: 0.88rem 0.92rem; color: #334155; font-size: 0.96rem; font-weight: 700; line-height: 1.38; white-space: pre-wrap; }
    .rt-mod-notes { display: flex; flex-direction: column; gap: 0.66rem; }
    .rt-mod-note { border-radius: 18px; border: 1px solid rgba(100, 116, 139, 0.14); background: rgba(255, 255, 255, 0.94); padding: 0.78rem 0.84rem; }
    .rt-mod-note-label { color: #475569; font-size: 0.72rem; font-weight: 950; letter-spacing: 0.08em; text-transform: uppercase; margin-bottom: 0.36rem; }
    .rt-mod-note-body { color: #1f2937; font-size: 0.94rem; font-weight: 760; line-height: 1.42; overflow-wrap: anywhere; }
    .rt-mod-note-list { margin: 0.18rem 0 0 1.12rem; padding-left: 0.72rem; color: #1f2937; font-size: 0.94rem; font-weight: 760; line-height: 1.38; }
    .rt-mod-note-list li { margin-bottom: 0.34rem; padding-left: 0.08rem; }
    .rt-mod-banner { border-radius: 22px; border: 1px solid rgba(146, 64, 14, 0.20); background: linear-gradient(135deg, #fffbeb 0%, #ffffff 70%); box-shadow: 0 12px 26px rgba(30, 41, 59, 0.06); padding: 0.72rem 0.92rem; margin: 0.2rem 0 0.6rem 0; display: flex; align-items: center; justify-content: space-between; gap: 0.75rem; }
    .rt-mod-banner-title { font-size: 1.02rem; font-weight: 950; color: #111827; letter-spacing: -0.02em; }
    .rt-mod-banner-note { border-radius: 999px; background: rgba(146, 64, 14, 0.10); color: #92400e; padding: 0.26rem 0.62rem; font-size: 0.78rem; font-weight: 900; }
    @media (max-width: 1100px) { .rt-kpi-grid { grid-template-columns: repeat(2, minmax(0, 1fr)); } .rt-field-grid { grid-template-columns: 1fr; } .rt-hero-top { flex-direction: column; } .rt-status-pill { text-align: left; } }
</style>
"""


def apply_presentation_css() -> None:
    st.markdown(PRESENTATION_CSS, unsafe_allow_html=True)


def timing_signal(days: object) -> Tuple[str, str]:
    text_value = fmt_day_delta(days)
    if text_value == "-":
        return text_value, "rt-signal-muted"

    try:
        days_i = int(round(float(days)))
    except Exception:
        return text_value, "rt-signal-muted"

    if days_i < 0:
        return text_value, "rt-signal-red"
    if days_i == 0:
        return text_value, "rt-signal-amber"
    return text_value, "rt-signal-green"


def render_html_kpi_strip(items: List[Tuple]) -> None:
    allowed_helper_classes = {
        "rt-signal-green",
        "rt-signal-red",
        "rt-signal-amber",
        "rt-signal-muted",
    }
    cards = []

    for item in items:
        if len(item) == 4:
            label, value, helper, helper_class = item
        else:
            label, value, helper = item
            helper_class = "rt-signal-muted"

        helper_text = html_safe(helper, blank="")
        helper_class = str(helper_class or "rt-signal-muted")
        if helper_class not in allowed_helper_classes:
            helper_class = "rt-signal-muted"

        helper_html = (
            f"<div class='rt-kpi-helper {helper_class}'>{helper_text}</div>"
            if helper_text
            else ""
        )
        cards.append(
            "<div class='rt-kpi'>"
            f"<div class='rt-kpi-label'>{html_safe(label)}</div>"
            f"<div class='rt-kpi-value'>{html_safe(value)}</div>"
            f"{helper_html}"
            "</div>"
        )

    st.markdown(f"<div class='rt-kpi-grid'>{''.join(cards)}</div>", unsafe_allow_html=True)


def render_html_field_panel(title: str, items: List[Tuple[str, object]], note: object = "") -> None:
    fields = []
    for label, value in items:
        fields.append(
            "<div class='rt-field'>"
            f"<div class='rt-field-label'>{html_safe(label)}</div>"
            f"<div class='rt-field-value'>{html_safe(value)}</div>"
            "</div>"
        )
    note_text = html_safe(note, blank="")
    note_html = f"<div class='rt-panel-note'>{note_text}</div>" if note_text else ""
    st.markdown(
        "<section class='rt-panel'>"
        "<div class='rt-panel-header'>"
        f"<div class='rt-panel-title'>{html_safe(title)}</div>"
        f"{note_html}"
        "</div>"
        f"<div class='rt-field-grid'>{''.join(fields)}</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def render_html_prompt_panel(prompts: List[str]) -> None:
    blocks = []
    for i, prompt in enumerate(prompts, start=1):
        blocks.append(
            "<div class='rt-prompt'>"
            f"<div class='rt-prompt-num'>{i}</div>"
            f"<div class='rt-prompt-text'>{html_safe(prompt)}</div>"
            "</div>"
        )
    st.markdown(
        "<section class='rt-panel'>"
        "<div class='rt-panel-header'><div class='rt-panel-title'>Presenter prompts</div></div>"
        f"<div class='rt-prompt-list'>{''.join(blocks)}</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def render_html_commentary_panel(commentary: object) -> None:
    st.markdown(
        "<section class='rt-panel'>"
        "<div class='rt-panel-header'><div class='rt-panel-title'>AM commentary</div></div>"
        f"<div class='rt-commentary'>{html_safe(commentary, blank='No commentary entered.')}</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def render_agenda_jump(deck: pd.DataFrame, current_idx: int) -> None:
    option_labels = [
        f"{i + 1}. {row.sheet} | {row.deal_name} | {row.deal_number}"
        for i, row in enumerate(
            deck[["sheet", "deal_name", "deal_number"]].itertuples(index=False, name="DealRow")
        )
    ]
    selected_label = st.selectbox("Jump to deal", options=option_labels, index=current_idx)
    selected_idx = option_labels.index(selected_label)
    if selected_idx != current_idx:
        st.session_state.selected_deal_key = deck.iloc[selected_idx]["deal_key"]
        st.rerun()


def render_html_status_check_panel(current_row: pd.Series) -> None:
    if not bool(current_row.get("status_needs_update_prompt", False)):
        return

    body = str(current_row.get("status_prompt_message") or "Review whether the meeting update should be refreshed.")
    st.markdown(
        "<section class='rt-status-card rt-status-card-warning'>"
        "<div class='rt-status-card-title'>Update reminder</div>"
        f"<div class='rt-status-card-body'>{html_safe(body)}</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def build_capital_items(current_row: pd.Series) -> List[Tuple[str, object]]:
    if current_row.get("sheet") == "Bridge":
        return [
            ("UPB", fmt_money(current_row.get("upb"), decimals=0)),
            ("Funded Amount", fmt_money(current_row.get("funded_amount"), decimals=0)),
            ("Loan Commitment", fmt_money(current_row.get("commitment"), decimals=0)),
            ("Remaining Commitment", fmt_money(current_row.get("remaining_commitment"), decimals=0)),
            ("Maturity Date", fmt_date(current_row.get("maturity_date"))),
            ("Next Payment Date", fmt_date(current_row.get("next_payment_date"))),
            ("Financing", current_row.get("financing")),
            ("Loan Buyer", current_row.get("loan_buyer")),
        ]

    return [
        ("UPB", fmt_money(current_row.get("upb"), decimals=0)),
        ("Funded Amount", fmt_money(first_nonblank_value(current_row.get("funded_amount"), current_row.get("loan_amount")), decimals=0)),
        ("Loan Amount", fmt_money(current_row.get("loan_amount"), decimals=0)),
        ("Maturity Date", fmt_date(current_row.get("maturity_date"))),
        ("Next Payment Date", fmt_date(current_row.get("next_payment_date"))),
        ("Financing", current_row.get("financing")),
        ("Loan Buyer", current_row.get("loan_buyer")),
        ("Segment", current_row.get("segment")),
    ]


def render_html_hero(current_row: pd.Series, current_idx: int, total_count: int) -> None:
    status_class = "rt-status-warning" if bool(current_row.get("status_needs_update_prompt", False)) else "rt-status-ok"
    st.markdown(
        "<section class='rt-hero'>"
        "<div class='rt-hero-top'>"
        "<div>"
        f"<div class='rt-eyebrow'>{html_safe(current_row.get('sheet'))} &bull; Agenda {current_idx + 1} of {total_count}</div>"
        f"<div class='rt-hero-title'>{html_safe(current_row.get('deal_name'))}</div>"
        f"<div class='rt-hero-subtitle'>Deal {html_safe(current_row.get('deal_number'))} &bull; Borrower: {html_safe(current_row.get('borrower'))}</div>"
        "<div class='rt-chip-row'>"
        f"<span class='rt-chip'>Owner: {html_safe(current_row.get('owner'))}</span>"
        f"<span class='rt-chip'>Servicer: {html_safe(current_row.get('servicer'))}</span>"
        f"<span class='rt-chip'>Portfolio: {html_safe(current_row.get('portfolio'))}</span>"
        f"<span class='rt-chip'>Segment: {html_safe(current_row.get('segment'))}</span>"
        "</div>"
        "</div>"
        f"<div class='rt-status-pill {status_class}'>"
        "<div class='rt-label'>Current update</div>"
        f"<div class='rt-value'>{html_safe(current_row.get('status'), blank='No status')}</div>"
        "</div>"
        "</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def split_loan_mod_note_lines(label: str, value: object) -> List[str]:
    text = display_text(value, blank="")
    if not text:
        return []

    # Remove duplicated field labels from Salesforce text, for example "Comments: ...".
    text = re.sub(rf"^{re.escape(label)}\s*:\s*", "", text, flags=re.IGNORECASE).strip()

    # Turn common Salesforce note separators into readable bullets.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\s+[-–—](?=[A-Za-z0-9$])", "\n", text)
    text = re.sub(r"\s*[•]\s+", "\n", text)
    text = re.sub(r"\s+(?=\d+\.\s+[A-Za-z$])", "\n", text)
    text = re.sub(r"\s+\(([ivxlcdm]+)\)\s*", r"\n(\1) ", text, flags=re.IGNORECASE)

    raw_lines = []
    for chunk in text.split("\n"):
        chunk = chunk.strip(" -–—;\t")
        chunk = re.sub(r"^\d+\.\s*", "", chunk)
        chunk = re.sub(r"^\(([ivxlcdm]+)\)\s*", "", chunk, flags=re.IGNORECASE)
        if chunk:
            raw_lines.append(chunk)

    if not raw_lines and text:
        raw_lines = [text]

    return raw_lines


def loan_mod_note_section_html(label: str, value: object) -> str:
    lines = split_loan_mod_note_lines(label, value)
    if not lines:
        return ""

    if len(lines) == 1:
        body_html = f"<div class='rt-mod-note-body'>{html_safe(lines[0])}</div>"
    else:
        items_html = "".join(f"<li>{html_safe(line)}</li>" for line in lines)
        body_html = f"<ul class='rt-mod-note-list'>{items_html}</ul>"

    return (
        "<div class='rt-mod-note'>"
        f"<div class='rt-mod-note-label'>{html_safe(label)}</div>"
        f"{body_html}"
        "</div>"
    )


def render_html_loan_mod_notes_panel(current_row: pd.Series) -> None:
    html_block = loan_mod_note_section_html("Comments", current_row.get("loan_mod_latest_comments"))
    if not html_block:
        return

    st.markdown(
        "<section class='rt-panel'>"
        "<div class='rt-panel-header'><div class='rt-panel-title'>Highlighted loan mod comments</div></div>"
        f"<div class='rt-mod-notes'>{html_block}</div>"
        "</section>",
        unsafe_allow_html=True,
    )


def render_html_loan_modification_panel(current_row: pd.Series) -> None:
    if not has_loan_mods(current_row):
        return

    render_html_field_panel(
        "Loan modification spotlight",
        [
            ("Modification Type", current_row.get("loan_mod_latest_type")),
            ("Mod Status", current_row.get("loan_mod_latest_status")),
            ("Mod Order #", current_row.get("loan_mod_latest_order_number")),
            ("Completed Date", fmt_date(current_row.get("loan_mod_latest_completed_date"))),
            ("Effective Date", fmt_date(current_row.get("loan_mod_latest_effective_date"))),
            ("Previous Maturity", fmt_date(current_row.get("loan_mod_latest_previous_maturity_date"))),
            ("Updated Maturity", fmt_date(current_row.get("loan_mod_latest_updated_maturity_date"))),
            ("Previous Commitment", fmt_money(current_row.get("loan_mod_latest_previous_commitment"), decimals=0)),
            ("Updated Commitment", fmt_money(current_row.get("loan_mod_latest_updated_commitment"), decimals=0)),
        ],
        note=f"{fmt_int(current_row.get('loan_mod_count'))} mod record(s)",
    )

    render_html_loan_mod_notes_panel(current_row)


def render_html_loan_modification_record(
    record: Dict[str, object],
    position: int,
    total: int,
) -> None:
    """Render one specific modification record (used by the presentation toggle)."""
    mod_type = first_nonblank_value(
        record.get("Modification Type"),
        record.get("Loan Mod Type"),
        record.get("Mod Reporting Type"),
    )

    render_html_field_panel(
        "Loan modification detail",
        [
            ("Modification Type", mod_type),
            ("Mod Status", record.get("Status")),
            ("Mod Order #", record.get("Loan Mod Order Number")),
            ("Completed Date", fmt_date(record.get("System Task Completed Date"))),
            ("Effective Date", fmt_date(record.get("Mod Effective Date"))),
            ("Finalized Date", fmt_date(record.get("Modification Finalized Date"))),
            ("Previous Maturity", fmt_date(record.get("Previous Maturity Date"))),
            ("Updated Maturity", fmt_date(record.get("Updated Maturity Date"))),
            ("Previous Commitment", fmt_money(record.get("Previous Loan Commitment"), decimals=0)),
            ("Updated Commitment", fmt_money(record.get("Updated Loan Commitment"), decimals=0)),
        ],
        note=f"Showing modification {position} of {total}",
    )

    comments_html = loan_mod_note_section_html("Comments", record.get("Comments"))
    if comments_html:
        st.markdown(
            "<section class='rt-panel'>"
            "<div class='rt-panel-header'><div class='rt-panel-title'>Highlighted loan mod comments</div></div>"
            f"<div class='rt-mod-notes'>{comments_html}</div>"
            "</section>",
            unsafe_allow_html=True,
        )


def render_presentation_loan_mods_top(
    current_row: pd.Series,
    loan_mods: pd.DataFrame,
    deal_key: str,
) -> None:
    """Loan modifications section at the very top of the presentation.

    When a loan carries more than one modification, a segmented control lets the
    presenter toggle between each modification. The detail panel below reflects
    whichever modification is selected.
    """
    mods = loan_mods_for_deal(loan_mods, current_row.get("deal_number"))
    if mods.empty:
        return

    total = len(mods)
    labels = build_loan_mod_toggle_labels(mods)
    selected_pos = 0

    with st.container(border=True):
        st.markdown(
            "<div class='rt-mod-banner'>"
            "<div class='rt-mod-banner-title'>Loan modifications</div>"
            f"<div class='rt-mod-banner-note'>{fmt_int(total)} modification record(s) on this loan</div>"
            "</div>",
            unsafe_allow_html=True,
        )

        if total > 1:
            widget_key = f"loanmod_toggle::{deal_key}"
            chosen = st.segmented_control(
                "Toggle between modifications",
                options=labels,
                default=labels[0],
                key=widget_key,
            )
            selected_pos = labels.index(chosen) if chosen in labels else 0
        else:
            st.caption(f"Single modification on file: {labels[0]}")

        record = mods.iloc[selected_pos].to_dict()
        render_html_loan_modification_record(record, selected_pos + 1, total)


def render_presentation_view(deck: pd.DataFrame, raw_deck: pd.DataFrame, loan_mods: pd.DataFrame) -> None:
    apply_presentation_css()
    sync_selected_deal(deck)
    current_row = get_selected_row(deck)
    current_idx = int(deck.index[deck["deal_key"] == current_row["deal_key"]][0])
    total_count = len(deck)
    progress_value = 0.0 if total_count <= 1 else (current_idx + 1) / total_count
    deal_key = str(current_row.get("deal_key"))

    with st.container(border=True):
        top_left, top_mid, top_filter, top_jump, top_prompt, top_edit = st.columns(
            [0.72, 0.72, 1.12, 2.55, 1.05, 0.98],
            vertical_alignment="bottom",
        )
        with top_left:
            st.button(
                "Previous",
                use_container_width=True,
                disabled=current_idx == 0,
                on_click=move_selection,
                args=(deck, -1),
            )
        with top_mid:
            st.button(
                "Next",
                use_container_width=True,
                disabled=current_idx >= total_count - 1,
                on_click=move_selection,
                args=(deck, 1),
            )
        with top_filter:
            st.segmented_control(
                "Portfolio type",
                options=["All", "Bridge", "Term"],
                key="sheet_filter",
                width="stretch",
            )
        with top_jump:
            render_agenda_jump(deck, current_idx)
        with top_prompt:
            with st.popover("Prompts", icon=":material/lightbulb:", width="stretch"):
                render_html_prompt_panel(build_presenter_prompts(current_row))
        with top_edit:
            if st.button("Edit status", type="primary", use_container_width=True):
                st.session_state.dialog_target = deal_key
                st.rerun()

        st.progress(progress_value, text=f"Agenda item {current_idx + 1} of {total_count}")

    render_html_hero(current_row, current_idx, total_count)

    # Loan modifications moved to the top, with a toggle to switch between each
    # modification on the current loan.
    render_presentation_loan_mods_top(current_row, loan_mods, deal_key)

    funded_kpi_value = first_nonblank_value(current_row.get("funded_amount"), current_row.get("loan_amount"))
    maturity_timing, maturity_signal_class = timing_signal(current_row.get("days_to_maturity"))
    payment_timing, payment_signal_class = timing_signal(current_row.get("days_to_next_payment"))
    loan_mod_count = int(current_row.get("loan_mod_count") or 0)
    loan_mod_helper = display_text(current_row.get("loan_mod_latest_type"), blank="No mods")
    loan_mod_signal = "rt-signal-amber" if loan_mod_count else "rt-signal-muted"

    render_html_kpi_strip(
        [
            ("UPB", fmt_money(current_row.get("upb"), decimals=0), "Current balance", "rt-signal-muted"),
            ("Funded Amount", fmt_money(funded_kpi_value, decimals=0), "Funded / loan amount", "rt-signal-muted"),
            ("Maturity", fmt_date(current_row.get("maturity_date")), maturity_timing, maturity_signal_class),
            ("Next Payment", fmt_date(current_row.get("next_payment_date")), payment_timing, payment_signal_class),
            ("Loan Mods", fmt_int(loan_mod_count), loan_mod_helper, loan_mod_signal),
        ]
    )

    left, right = st.columns([1.34, 0.86], gap="large", vertical_alignment="top")

    with left:
        render_html_field_panel(
            "Deal snapshot",
            [
                ("Borrower", current_row.get("borrower")),
                ("Account", current_row.get("account_display")),
                ("Deal #", current_row.get("deal_number")),
                ("Servicer", current_row.get("servicer")),
                ("Portfolio", current_row.get("portfolio")),
                ("Segment", current_row.get("segment")),
                ("Owner", current_row.get("owner")),
                ("NPL / Delinquency", current_row.get("npl_raw")),
            ],
            note=f"Days Past Due: {fmt_int(current_row.get('days_past_due'))}",
        )
        render_html_field_panel("Capital profile", build_capital_items(current_row))

    with right:
        render_html_status_check_panel(current_row)
        render_html_commentary_panel(current_row.get("commentary"))

        with st.expander("Nearby agenda", expanded=False):
            start = max(current_idx - 3, 0)
            end = min(current_idx + 4, total_count)
            agenda_slice = deck.iloc[start:end].copy()
            agenda_slice.insert(0, "Agenda #", range(start + 1, end + 1))
            agenda_slice["Current"] = agenda_slice["deal_key"].map(lambda value: "Current" if value == deal_key else "")
            agenda_slice["Update Reminder"] = agenda_slice["status_needs_update_prompt"].map(lambda value: "Yes" if bool(value) else "")
            agenda_slice["Loan Mods"] = agenda_slice["loan_mod_count"].map(lambda value: fmt_int(value) if int(value or 0) else "")
            agenda_slice["UPB"] = agenda_slice["upb"].map(lambda value: fmt_money(value, decimals=0))
            st.dataframe(
                agenda_slice[
                    [
                        "Agenda #",
                        "Current",
                        "sheet",
                        "deal_name",
                        "deal_number",
                        "Update Reminder",
                        "Loan Mods",
                        "status",
                        "owner",
                        "UPB",
                    ]
                ].rename(
                    columns={
                        "sheet": "Type",
                        "deal_name": "Deal Name",
                        "deal_number": "Deal #",
                        "status": "Status",
                        "owner": "Owner",
                    }
                ),
                hide_index=True,
                use_container_width=True,
            )


# -----------------------------------------------------------------------------
# Review and exports
# -----------------------------------------------------------------------------
def render_review_view(deck: pd.DataFrame, raw_deck: pd.DataFrame) -> None:
    st.subheader("Review and update")
    st.caption("Edit Status, Owner, and AM Commentary. Disabled columns are for reference.")

    review_df = build_review_dataframe(deck)
    edited_df = st.data_editor(
        review_df,
        key="review_editor",
        use_container_width=True,
        hide_index=False,
        disabled=[
            "Agenda #",
            "Type",
            "Deal Name",
            "Deal #",
            "Maturity",
            "Next Payment",
            "UPB",
            "Update Reminder",
        ],
        column_config={
            "AM Commentary": st.column_config.TextColumn("AM Commentary", width="large"),
            "Status": st.column_config.TextColumn("Status", width="medium"),
            "Owner": st.column_config.TextColumn("Owner", width="medium"),
        },
    )

    c1, c2 = st.columns([0.4, 0.6])
    with c1:
        if st.button("Apply review edits", type="primary", use_container_width=True):
            apply_review_edits(edited_df, deck, raw_deck)
            st.success("Review edits applied.")
            st.rerun()
    with c2:
        st.caption("Download the updated workbook from Exports after applying edits.")


def render_exports_view(
    file_bytes: bytes,
    workbook_name: str,
    raw_deck: pd.DataFrame,
    deck: pd.DataFrame,
    as_of_date: dt.date,
) -> None:
    st.subheader("Exports")

    updated_history = merge_status_history(
        existing_history=st.session_state.status_history,
        deck=deck,
        meeting_date=as_of_date,
        workbook_name=workbook_name,
    )
    update_bytes = export_overrides_csv(raw_deck, deck)
    workbook_bytes = update_workbook_bytes(file_bytes, ensure_override_store())
    workbook_bytes = write_status_history_sheet_to_workbook_bytes(workbook_bytes, updated_history)
    agenda_bytes = build_agenda_dataframe(deck).to_csv(index=False).encode("utf-8")

    c1, c2 = st.columns(2)
    with c1:
        st.download_button(
            "Download changed statuses/commentary CSV",
            data=update_bytes,
            file_name="weekly_portfolio_meeting_updates.csv",
            mime="text/csv",
            use_container_width=True,
            disabled=(len(update_bytes) == 0),
        )
        st.download_button(
            "Download agenda CSV",
            data=agenda_bytes,
            file_name="weekly_portfolio_agenda.csv",
            mime="text/csv",
            use_container_width=True,
        )

    with c2:
        st.download_button(
            "Download updated workbook with embedded history",
            data=workbook_bytes,
            file_name=f"updated_{workbook_name}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        st.download_button(
            "Download status history CSV backup",
            data=status_history_to_csv_bytes(updated_history),
            file_name="status_history_backup.csv",
            mime="text/csv",
            use_container_width=True,
        )

    st.info(
        "The updated workbook includes the hidden _Status History sheet. Use that workbook or the "
        "Jupyter weekly overview generator as the normal path; the CSV is only a backup/export."
    )

    st.subheader("Recommended .gitignore entries")
    st.code(
        "status_history.csv\n"
        "*_status_history.csv\n"
        "weekly_portfolio_status_history*.csv\n"
        "*.xlsx\n",
        language="gitignore",
    )


# -----------------------------------------------------------------------------
# Workbook source and main app
# -----------------------------------------------------------------------------
def get_workbook_source() -> Tuple[Optional[bytes], str]:
    uploaded = st.session_state.get("uploaded_workbook")
    if uploaded is not None:
        return uploaded.getvalue(), uploaded.name

    default_path = first_existing_file(DEFAULT_SAMPLE_FILES)
    if default_path is None:
        return None, "No workbook loaded"

    return default_path.read_bytes(), default_path.name


def render_no_workbook_loaded() -> None:
    st.info("Upload a Portfolio Overview workbook from the sidebar to start the meeting deck.")
    st.write("The app expects workbook sheets named Bridge and/or Term with a Deal Number header row.")
    st.subheader("Status history")
    st.write(
        "Status history is automatic when the workbook contains the hidden _Status History sheet "
        "created by the weekly overview generator. A CSV upload remains available in the sidebar only as a backup."
    )
    st.subheader("Loan modifications")
    st.write(
        "Loan modification details load automatically when the uploaded workbook includes a visible "
        f"{LOAN_MOD_SHEET_NAME!r} sheet with a Deal Number column."
    )


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, layout="wide")
    initialize_state()
    load_default_status_history_once()

    file_bytes, workbook_name = get_workbook_source()
    render_controls_bar(workbook_name)
    process_status_history_upload()
    if file_bytes is not None:
        load_embedded_status_history_if_available(file_bytes, workbook_name)

    if file_bytes is None:
        render_no_workbook_loaded()
        return

    try:
        raw_deck, _metadata = load_portfolio_workbook(
            file_bytes=file_bytes,
            as_of_date_iso=st.session_state.as_of_date.isoformat(),
            include_hidden=bool(st.session_state.include_hidden),
        )
    except Exception as exc:
        st.error(f"Could not load workbook: {exc}")
        return

    if raw_deck.empty:
        st.warning("No Bridge or Term rows were found in the workbook.")
        return

    loan_mods = load_loan_modifications_from_workbook(file_bytes)
    raw_deck = add_loan_modification_columns(raw_deck, loan_mods)

    deck = apply_overrides(raw_deck, ensure_override_store())
    deck = add_status_staleness_columns(
        deck=deck,
        status_history=st.session_state.status_history,
        meeting_date=st.session_state.as_of_date,
        workbook_name=workbook_name,
    )
    filtered_deck = apply_filters(deck)

    if filtered_deck.empty:
        st.warning("No deals match the current filters.")
        return

    maybe_open_edit_dialog(deck, raw_deck)

    selected_view = st.session_state.view_mode
    if selected_view == "Overview":
        render_overview_view(filtered_deck, st.session_state.as_of_date)
    elif selected_view == "Presentation":
        render_presentation_view(filtered_deck, raw_deck, loan_mods)
    elif selected_view == "Review":
        render_review_view(filtered_deck, raw_deck)
    elif selected_view == "Exports":
        render_exports_view(
            file_bytes=file_bytes,
            workbook_name=workbook_name,
            raw_deck=raw_deck,
            deck=deck,
            as_of_date=st.session_state.as_of_date,
        )


if __name__ == "__main__":
    main()
