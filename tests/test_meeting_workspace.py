"""Synthetic meeting model, recovery, export and UI regressions."""
import copy
import datetime as dt
from email import message_from_bytes
import hashlib
import io
import json
import threading
from types import SimpleNamespace
import zipfile
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
import meeting_workspace as mw

SHA = "a" * 64
DATE = dt.date(2026, 10, 20)
ROW = {"deal_key": "Bridge::10001", "deal_name": "Synthetic Harbor", "deal_number": "10001", "sheet": "Bridge",
       "status": "Current", "owner": "Alice", "commentary": "Review draw", "upb": 1000000,
       "commitment": 2000000, "maturity_date": "2027-01-31", "next_payment_date": "2026-11-01"}

def model():
    return mw.new_workspace(SHA, "synthetic.xlsx", DATE)

def ctx_for(workspace, raw=None, rows=None, parsed=None, wb=None):
    rows = rows or [ROW]
    return SimpleNamespace(wb_sha=workspace["workbook_sha"], workbook_name="synthetic.xlsx", as_of=DATE,
        deck=pd.DataFrame(rows), raw_deck=pd.DataFrame(raw or rows), parsed=parsed or {}, wb=wb or {})

def test_checkpoint_undo_redo_and_clear_are_lossless():
    ws = mw.checkpoint(model(), {})
    first = {ROW["deal_key"]: {"status": "Surveillance", "owner": "Alice", "commentary": "A"}}
    ws = mw.checkpoint(ws, first, base_rows={ROW["deal_key"]: ROW}, now=10)
    ws = mw.checkpoint(ws, {ROW["deal_key"]: {**first[ROW["deal_key"]], "status": "Workout"}}, now=20)
    assert len(ws["undo"]) == 2
    assert ws["changes"][0]["before"] == "Current"
    same = mw.checkpoint(ws, ws["overrides"], now=30)
    assert same == ws
    undone = mw.undo_edit(ws)
    assert undone["overrides"] == first
    assert mw.undo_edit(undone, redo=True)["overrides"] == ws["overrides"]
    assert mw.undo_edit(mw.undo_edit(ws))["overrides"] == {}
    assert ws["overrides"][ROW["deal_key"]]["status"] == "Workout"

def test_new_edit_after_undo_invalidates_redo():
    ws = mw.checkpoint(model(), {})
    ws = mw.checkpoint(ws, {"a": {"status": "One"}})
    ws = mw.undo_edit(ws)
    assert ws["redo"]
    assert mw.checkpoint(ws, {"a": {"status": "Two"}})["redo"] == []

def test_undo_history_is_bounded():
    ws = mw.checkpoint(model(), {})
    for n in range(80):
        ws = mw.checkpoint(ws, {"a": {"status": str(n)}})
    assert len(ws["undo"]) == mw.MAX_HISTORY

def test_action_default_schema_close_reopen_and_overdue():
    ws = mw.add_item(model(), item_type="Action", text="Obtain inspection", owner="Alice", due_date="2026-10-19", row=ROW, now=10)
    item = ws["items"][0]
    assert list(item) == mw.ACTION_ITEM_COLUMNS
    assert mw.followups(ws, mode="Overdue")[0]["item_id"] == item["item_id"]
    closed = mw.update_item(ws, item["item_id"], status="Done", now=20)
    assert closed["items"][0]["closed_meeting_date"] == DATE.isoformat()
    assert not mw.followups(closed)
    reopened = mw.update_item(closed, item["item_id"], status="Open", now=30)
    assert reopened["items"][0]["closed_at"] == ""
    assert item["status"] == "Open"

@pytest.mark.parametrize("kind", ["Decision", "Note", "Parking lot"])
def test_capture_types(kind):
    ws = mw.add_item(model(), item_type=kind, text="Synthetic note")
    assert ws["items"][0]["status"] == ("Open" if kind == "Parking lot" else "Logged")
    assert ws["items"][0]["due_date"] == ""

@pytest.mark.parametrize("bad", ["", " ", "x" * 1001])
def test_blank_or_oversized_capture_is_rejected(bad):
    with pytest.raises(ValueError):
        mw.add_item(model(), item_type="Action", text=bad)

def test_item_merge_newest_wins_and_closure_survives():
    ws = mw.add_item(model(), item_type="Action", text="Inspect", now=10)
    closed = mw.update_item(ws, ws["items"][0]["item_id"], status="Done", now=20)
    assert mw.merge_items(closed["items"], ws["items"]) == closed["items"]
    assert mw.merge_items(ws["items"], closed["items"]) == closed["items"]

def test_clock_excludes_pauses_and_stays_stopped():
    ws = mw.control_clock(model(), "start", now=100)
    ws = mw.control_clock(ws, "pause", now=160)
    assert mw.clock_seconds(ws["clock"], now=1000) == 60
    ws = mw.control_clock(ws, "resume", now=200)
    assert mw.clock_seconds(ws["clock"], now=240) == 100
    ws = mw.control_clock(ws, "end", now=250)
    assert mw.clock_seconds(ws["clock"], now=9999) == 110

def test_ending_while_paused_freezes_elapsed():
    ws = mw.control_clock(model(), "start", now=100)
    ws = mw.control_clock(ws, "pause", now=160)
    ws = mw.control_clock(ws, "end", now=200)
    assert mw.clock_seconds(ws["clock"], now=9999) == 60

def test_recovery_roundtrip_is_scoped_and_has_no_workbook_bytes():
    ws = mw.add_item(model(), item_type="Decision", text="Approve extension", row=ROW)
    ws = mw.set_outcome(ws, ROW["deal_key"], "Discussed")
    payload = mw.recovery_bytes(ws)
    restored = mw.parse_recovery(payload, SHA)
    assert restored["items"] == ws["items"]
    assert restored["progress"] == ws["progress"]
    assert b"workbook_bytes" not in payload
    with pytest.raises(ValueError, match="different workbook"):
        mw.parse_recovery(payload, "b" * 64)
    assert len(payload) < 200_000

@pytest.mark.parametrize("payload", [b"broken", b"[]", b'{"schema":"unknown"}', b'{"schema":"wpo-workspace/1","workbook_sha":NaN}'])
def test_malformed_recovery_rejected_without_mutation(payload):
    ws = model()
    before = copy.deepcopy(ws)
    with pytest.raises(ValueError):
        mw.parse_recovery(payload, SHA)
    assert ws == before

def test_tampered_item_clock_and_foreign_override_rejected():
    for key, value in [("items", [{"item_id": "1", "type": "Other"}]), ("clock", {"started_at": "bad"}), ("overrides", {"x": {"workbook_sha": "b" * 64}})]:
        ws = model()
        ws[key] = value
        with pytest.raises(ValueError):
            mw.parse_recovery(json.dumps(ws).encode(), SHA)

def test_atomic_journal_and_failure_preserves_previous(tmp_path, monkeypatch):
    ws = model()
    target = mw.write_journal(tmp_path, ws)
    original = target.read_bytes()
    assert mw.read_journal(tmp_path, SHA)["workbook_sha"] == SHA
    def fail(*args):
        raise OSError("simulated")
    monkeypatch.setattr(mw.os, "replace", fail)
    with pytest.raises(OSError):
        mw.write_journal(tmp_path, mw.add_item(ws, item_type="Note", text="new"))
    assert target.read_bytes() == original
    assert not list(tmp_path.glob(".wpo-*.tmp"))

def test_memory_journal_copies_prunes_and_caps():
    store = {"lock": threading.RLock(), "entries": {}}
    payload = {"session": {"overrides": {"x": "one"}}}
    mw.memory_save("old", payload, store=store, now=0)
    payload["session"]["overrides"]["x"] = "changed"
    assert mw.memory_load("old", store=store, now=1)["session"]["overrides"]["x"] == "one"
    for n in range(12):
        mw.memory_save(str(n), {"session": {}}, store=store, now=100000 + n)
    assert len(store["entries"]) == 8
    assert mw.memory_load("old", store=store, now=100100) is None
    assert mw.memory_load("11", store=store, now=200000) is None

def test_snapshot_comparison_detects_fields_new_and_missing_deals():
    previous = mw.build_snapshot([ROW], model(), now=10)
    current = mw.build_snapshot([{**ROW, "status": "Workout", "upb": 1200000},
                                 {**ROW, "deal_key": "Term::2", "deal_name": "Synthetic Term"}], model(), now=20)
    changes = mw.compare_snapshots(previous, current)
    assert {r["Field"] for r in changes} == {"status", "upb", "Deal"}
    assert len(mw.compare_snapshots(current, previous)) == 3

def test_minutes_email_pack_and_csv_are_safe_and_complete():
    ws = mw.add_item(model(), item_type="Decision", text="<script>alert(1)</script>", row=ROW)
    ws = mw.add_item(ws, item_type="Action", text="=Inspect", owner="Alice", due_date="2026-10-27", row=ROW)
    ws = mw.checkpoint(ws, {})
    ws = mw.checkpoint(ws, {ROW["deal_key"]: {"status": "Workout", "owner": "Alice", "commentary": ""}}, base_rows={ROW["deal_key"]: ROW})
    markdown, document, plain = mw.build_minutes(ws, [ROW], data_as_of=DATE)
    assert "Alice" in markdown and "2026-10-27" in markdown and "Current → Workout" in markdown
    assert "<script>" not in document and "Confidential" in document
    message = message_from_bytes(mw.minutes_email(ws, [ROW]))
    assert message["X-Unsent"] == "1"
    assert {p.get_content_type() for p in message.walk()} >= {"text/plain", "text/html"}
    with zipfile.ZipFile(io.BytesIO(mw.meeting_pack(ws, [ROW]))) as zf:
        assert set(zf.namelist()) == {"meeting_recovery.json", "meeting_snapshot.json", "minutes.md", "minutes.html", "outlook_draft.eml", "actions.csv", "changes.csv"}
        assert "'=Inspect" in zf.read("actions.csv").decode("utf-8-sig")
        assert mw.parse_recovery(zf.read("meeting_recovery.json"), SHA)["items"] == ws["items"]

def test_workbook_register_roundtrip_same_date_replaces_and_keeps_formula_cache(sample_bytes, app_module):
    ws = mw.add_item(model(), item_type="Action", text="=Synthetic inspection\nSecond line", owner="Alice", due_date=DATE)
    ws = mw.set_outcome(ws, ROW["deal_key"], "Discussed")
    ctx = ctx_for(ws)
    tables = mw.workbook_tables(ctx, ws)
    from workbook_io import replace_hidden_sheets
    output = replace_hidden_sheets(sample_bytes, tables)
    parsed = app_module.parse_workbook_bytes(output)
    ctx.parsed = parsed
    assert len(mw._embedded_rows(ctx, "_Action Items")) == 1
    reloaded = mw.workbook_tables(ctx, ws)
    assert len(reloaded["_Meeting Log"]) == 1 and len(reloaded["_Meeting Summary"]) == 1
    assert reloaded["_Action Items"].iloc[0]["text"] == ws["items"][0]["text"]
    assert reloaded["_Meeting Log"].iloc[0]["outcome"] == "Discussed"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(output), data_only=True)
    for name, frame in tables.items():
        assert wb[name].sheet_state == "hidden"
        assert [c.value for c in wb[name][1]] == list(frame.columns)

def workspace_harness():
    from types import SimpleNamespace
    import datetime as dt
    import pandas as pd
    import streamlit as st
    import meeting_workspace as mw
    mw.restore_session()
    row = {"deal_key": "Bridge::10001", "deal_name": "Synthetic Harbor", "deal_number": "10001",
           "sheet": "Bridge", "owner": "Alice", "status": "Current", "commentary": "", "upb": 1000}
    ctx = SimpleNamespace(wb_sha="a"*64, workbook_name="synthetic.xlsx", as_of=dt.date(2026,10,20),
        deck=pd.DataFrame([row]), raw_deck=pd.DataFrame([row]), parsed={}, wb={})
    st.session_state.setdefault("meeting_overrides", {})
    mw.checkpoint_overrides(ctx)
    view = st.radio("View", ["Deal", "Follow-ups", "Wrap-up"], key="mw_test_view")
    with st.sidebar:
        mw.render_settings(ctx)
    if view == "Deal":
        mw.render_deal(ctx, row)
    elif view == "Follow-ups":
        mw.render_followups(ctx)
    else:
        mw.render_wrapup(ctx)

def test_ui_capture_complete_and_wrapup():
    at = AppTest.from_function(workspace_harness).run()
    assert not at.exception
    assert "m" in at.query_params
    at.text_input[0].set_value("Request inspection")
    at.button(key=next(b.key for b in at.button if b.label == "Add item")).click().run()
    assert not at.exception
    assert len(at.session_state["_mw_model"]["items"]) == 1
    at.checkbox(key=next(c.key for c in at.checkbox if "Request inspection" in c.label)).check().run()
    assert not at.exception
    assert at.session_state["_mw_model"]["items"][0]["status"] == "Done"
    at.radio(key="mw_test_view").set_value("Wrap-up").run()
    assert not at.exception
    at.button(key="mw_remaining_not_reached").click().run()
    assert not at.exception
    assert at.session_state["_mw_model"]["progress"]["Bridge::10001"]["outcome"] == "Not reached"


def refresh_harness():
    import datetime as dt
    import hashlib
    from types import SimpleNamespace
    import pandas as pd
    import streamlit as st
    import meeting_workspace as mw
    mw.restore_session()
    raw = b"synthetic workbook memory only"
    sha = hashlib.sha256(raw).hexdigest()
    row = {"deal_key": "Bridge::10001", "deal_name": "Synthetic Harbor", "deal_number": "10001",
           "sheet": "Bridge", "owner": "Alice", "status": "Current", "commentary": "", "upb": 1000}
    st.session_state.setdefault("wb", {"bytes": raw, "sha256": sha, "name": "synthetic.xlsx", "source": "upload"})
    st.session_state.setdefault("meeting_overrides", {})
    st.session_state.setdefault("meeting_date", dt.date(2026, 10, 20))
    ctx = SimpleNamespace(wb_sha=sha, workbook_name="synthetic.xlsx", as_of=dt.date(2026,10,15),
        deck=pd.DataFrame([row]), raw_deck=pd.DataFrame([row]), parsed={}, wb=st.session_state["wb"])
    def save():
        st.session_state["meeting_overrides"] = {row["deal_key"]: {**row, "status": "Workout", "workbook_sha": sha}}
        st.session_state["cursor"] = row["deal_key"]
        ws = mw.add_item(mw.ensure_workspace(ctx), item_type="Action", text="Recover this action", owner="Alice", due_date="2026-10-27", row=row)
        ws = mw.set_outcome(ws, row["deal_key"], "Discussed")
        mw._put(ctx, ws)
    st.button("Make changes", on_click=save)
    mw.checkpoint_overrides(ctx)
    st.write(str(len(mw.ensure_workspace(ctx)["items"])))

def test_refresh_recovers_uploaded_workbook_edits_items_and_progress():
    first = AppTest.from_function(refresh_harness).run()
    assert not first.exception
    first.button[0].click().run()
    assert not first.exception
    token = first.query_params["m"]
    second = AppTest.from_function(refresh_harness)
    second.query_params["m"] = token
    second.run()
    assert not second.exception
    assert second.session_state["wb"]["bytes"] == b"synthetic workbook memory only"
    assert second.session_state["wb"]["source"] == "journal"
    assert second.session_state["meeting_overrides"]["Bridge::10001"]["status"] == "Workout"
    assert second.session_state["_mw_model"]["items"][0]["text"] == "Recover this action"
    assert second.session_state["_mw_model"]["progress"]["Bridge::10001"]["outcome"] == "Discussed"
    assert second.session_state["cursor"] == "Bridge::10001"
    assert second.query_params["m"] == token
    assert "Restored meeting" in second.session_state["_mw_restore_notice"]

def test_fresh_workbook_imports_saved_action_register(sample_bytes, app_module, monkeypatch):
    ws = mw.add_item(model(), item_type="Action", text="Carry me forward", owner="Alice", due_date="2026-10-27", row=ROW)
    from workbook_io import replace_hidden_sheets
    output = replace_hidden_sheets(sample_bytes, mw.workbook_tables(ctx_for(ws), ws))
    parsed = app_module.parse_workbook_bytes(output)
    ctx = ctx_for(ws, parsed=parsed)
    monkeypatch.setattr(mw.st, "session_state", {})
    fresh = mw.ensure_workspace(ctx)
    assert fresh["items"][0]["text"] == "Carry me forward"
    assert fresh["items"][0]["owner"] == "Alice"
    assert not fresh["overrides"]

def test_settings_default_does_not_write_local_disk(monkeypatch):
    calls = []
    monkeypatch.setattr(mw, "write_journal", lambda *a, **k: calls.append((a, k)))
    at = AppTest.from_function(workspace_harness).run()
    assert not at.exception
    at.button(key=next(b.key for b in at.button if b.label == "Discussed")).click().run()
    assert not at.exception
    assert calls == []


def test_summary_uses_canonical_delinquency_reminder_fields():
    ws = model()
    rows = [{**ROW, "days_past_due": 31, "upb": 1234567, "status_needs_update_prompt": True},
            {**ROW, "deal_key": "Term::2", "days_past_due": 0, "upb": 500000,
             "status_needs_update_prompt": True, "reminder_addressed": True}]
    summary = mw.workbook_tables(ctx_for(ws, rows=rows), ws)["_Meeting Summary"].iloc[0]
    assert summary["dpd30_upb"] == 1234567
    assert summary["reminders"] == 1
    assert summary["total_upb"] == 1734567
    assert summary["maturing_30d_count"] == 0

def test_existing_local_journal_not_overwritten_before_explicit_start(tmp_path, monkeypatch):
    ws = mw.add_item(model(), item_type="Note", text="Do not overwrite")
    target = mw.write_journal(tmp_path, ws)
    original = target.read_bytes()
    state = {"_mw_local_enabled": True, "_mw_local_folder": str(tmp_path)}
    monkeypatch.setattr(mw.st, "session_state", state)
    monkeypatch.setattr(mw, "_memory_checkpoint", lambda ctx: None)
    ctx = ctx_for(ws)
    mw._put(ctx, model())
    assert target.read_bytes() == original
    mw._enable_local_cb(ctx, str(tmp_path))
    assert mw.read_journal(tmp_path, SHA)["items"] == []

def test_memory_payload_tampered_workbook_rejects_before_restoring_state(monkeypatch):
    state = {}
    monkeypatch.setattr(mw.st, "session_state", state)
    monkeypatch.setattr(mw.st, "query_params", {"m": "a" * 32})
    monkeypatch.setattr(mw, "memory_load", lambda token: {"session": {"meeting_overrides": {"bad": "value"}},
        "workbook": {"bytes": b"wrong", "sha256": SHA}})
    assert mw.restore_session() is False
    assert "meeting_overrides" not in state


def test_clear_deletes_token_and_does_not_recreate_journal_on_next_checkpoint(monkeypatch):
    token = "z" * 32
    store = {"lock": threading.RLock(), "entries": {token: {"payload": {}, "touched": 1}}}
    state = {"_mw_token": token, "meeting_overrides": {"x": {}}, "wb": {"bytes": b"synthetic"},
             "_app_context": "synthetic source", "_data_context": "synthetic source",
             "status_history": "synthetic source", "prepared_export": b"synthetic export"}
    query = {"m": token}
    monkeypatch.setattr(mw.st, "session_state", state)
    monkeypatch.setattr(mw.st, "query_params", query)
    monkeypatch.setattr(mw, "journal_store", lambda: store)
    monkeypatch.setattr(mw, "_notify", lambda *args: None)
    ctx = ctx_for(model())
    mw.clear_meeting(ctx)
    assert state["meeting_overrides"] == {} and state["wb"] is None
    assert not {"_app_context", "_data_context", "status_history", "prepared_export"}.intersection(state)
    assert "m" not in query and token not in store["entries"]
    mw.checkpoint_overrides(ctx)
    assert not store["entries"] and "m" not in query
    mw._start_new_cb(ctx)
    assert query["m"] != token
    assert query["m"] in store["entries"]
