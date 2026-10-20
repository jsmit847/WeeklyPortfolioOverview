"""Synthetic Portfolio Overview workbook for tests and local demos.

The layout mirrors what the weekly overview generator produces: title rows above
the header, dated "M/D UPB" columns, a Loan Modifications sheet with Salesforce
fields, and the hidden _Status History sheet. It deliberately includes messy
real-world cases (float deal numbers, Excel serial dates, hidden rows, blank
subtotal rows, multi-line commentary, multi-mod loans).

The default workbook is the contract reference for acceptance numbers (10 visible
deals, total UPB 66,769,850 at as-of 2026-10-15) and must not change. Messier cases
are opt-in through keyword-only variant flags on build_sample_workbook(); see its
docstring (cell_overrides sets any cell by sheet, deal number and header).
build_scaled_workbook_bytes() builds large books for performance tests, and
with_styled_empty_rows() / with_cached_formula_values() patch the sheet XML of a built
workbook (formatted empty rows; saved results for formula cells, as Excel would write).

Usage:
    python tests/sample_workbook.py [output_path] [--as-of YYYY-MM-DD]
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import random
import re
import zipfile
from pathlib import Path
from typing import Dict, List, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Border, Side

DEFAULT_OUTPUT = Path("portfolio_overview.xlsx")


def _excel_serial(day: dt.date) -> int:
    return (day - dt.date(1899, 12, 30)).days


def _hdr(day: dt.date, suffix: str, with_year: bool = False) -> str:
    if with_year:
        return f"{day.month}/{day.day}/{day.strftime('%y')} {suffix}"
    return f"{day.month}/{day.day} {suffix}"


BRIDGE_DEALS = [
    # deal_number, name, borrower, point_person, status, commentary, commitment, funded, maturity_offset, payment_offset, dpd, npl, hidden
    (10001.0, "Harbor View Lofts", "Harbor View LLC", "Alice Chen", "Surveillance",
     "Construction 82% complete.\nSponsor injected $250k of additional equity.", 12_500_000, 10_250_000, 45, 5, 0, "", False),
    ("10002", "Maple Street Rehab", "Maple St Partners", "Bob Diaz", "Upcoming Maturity",
     "Extension request received; reviewing updated budget.", 3_400_000, 3_400_000, 12, -3, 3, "", False),
    (10003, "Sunset Ridge Townhomes", "Sunset Ridge Dev Co", "Alice Chen", "Modification",
     "Mod approved 3/15. Rate step-up to 9.0", 8_750_000, 7_900_000, -10, -15, 15, "30+", False),
    (10004.0, "Oak Creek SFR Portfolio", "Oak Creek Holdings", "Carla Gomez", "Foreclosure",
     "NOD filed. Trustee sale scheduled; receiver in place.", 5_200_000, 5_200_000, -120, -95, 95, "90+", False),
    ("10005", "Pine Hill Duplexes", "Pine Hill LLC", "", "",
     "", 1_150_000, 900_000, 200, 20, 0, "", False),
    ("10006", "Hidden Row Deal", "Ghost Borrower LLC", "Bob Diaz", "Surveillance",
     "This row is hidden in Excel.", 2_000_000, 1_500_000, 90, 10, 0, "", True),
    ("10007", "Cedar Point Condos", "Cedar Point Ventures (LLC)", "Dan Evans", "Litigation",
     "Borrower filed suit 2/1. Counsel engaged; hearing set for 4/20.", 15_000_000, 14_250_000, 30, 0, 0, "", False),
    ("10008", "Willow Bend Infill", "Willow Bend Builders", "Alice Chen", "Surveillance",
     "On track.", 950_000, 600_000, 365, 12, 0, "", False),
]

TERM_DEALS = [
    # deal_number, loan_name, borrower_entity, asset_manager, status, commentary, loan_amount, maturity_offset, payment_offset, dpd, npl
    ("20001", "Bayside Rental Portfolio", "Bayside SFR I LLC", "Erin Fox", "Surveillance",
     "Performing. DSCR 1.35x.", 22_000_000, 1500, 15, 0, ""),
    ("20002", "Lakeshore DSCR Loan", "Lakeshore Rentals", "Erin Fox", "Probate Court",
     "Guarantor deceased; estate in probate. Payments current via estate.", 640_000, 900, -2, 2, ""),
    ("20003", "Midtown 5-Unit", "Midtown Five LLC", "Frank Green", "Receivership",
     "Receiver appointed 3/1; collecting rents.", 1_875_000, 700, -60, 60, "60+"),
]

LOAN_MODS = [
    # deal_number, order, mod_type, status, completed_offset, effective_offset, prev_mat_offset, new_mat_offset, prev_commit, new_commit, comments
    ("10002", "2", "Maturity Extension", "Completed", -40, -38, -30, 12,
     3_400_000, 3_400_000, "Comments: Extended maturity 6 months to 2026. Borrower paid 0.5% extension fee - Updated budget approved"),
    ("10002", "1", "Maturity Extension", "Completed", -220, -215, -200, -30,
     3_000_000, 3_400_000, "Initial extension (i) paid down $100k (ii) added interest reserve"),
    ("10003", "1", "Rate Modification", "In Progress", None, -16, -10, 170,
     8_750_000, 8_750_000, "1. Rate step-up approved 2. Sponsor guaranty reaffirmed 3. Cash management triggered"),
    ("10007", "3", "Forbearance", "Completed", -5, -5, 30, 30,
     15_000_000, 15_000_000, "Forbearance through 4/30; Cedar Point Ventures (LLC) to cure by then."),
    ("10007", "2", "Commitment Increase", "Completed", -150, -148, 30, 30,
     13_500_000, 15_000_000, "Upsized for cost overruns."),
    ("10007", "1", "Maturity Extension", "Completed", -400, -398, -200, 30,
     13_500_000, 13_500_000, ""),
    ("20003", "1", "Forbearance", "Cancelled", None, None, 700, 700,
     1_875_000, 1_875_000, "Cancelled after receiver appointment."),
    ("99999", "1", "Maturity Extension", "Completed", -10, -10, 0, 180,
     500_000, 500_000, "Not on the overview; should be ignored."),
]

# Variant: the newest mod on 10008 was cancelled; the older one is in force (U42).
CANCELLED_LATEST_MODS = [
    ("10008", "2", "Maturity Extension", "Cancelled", None, None, 365, 900,
     950_000, 950_000, "Extension request withdrawn by sponsor."),
    ("10008", "1", "Rate Modification", "Completed", -60, -58, 365, 365,
     900_000, 950_000, "Rate cut -0.50% and reserve top-up."),
]
# Variant: a new 10007 request with no order number yet (U48).
UNNUMBERED_NEW_MOD = ("10007", "", "Maturity Extension", "In Progress", None, None, 30, 210,
                      15_000_000, 15_000_000, "New extension request under review.")

MATCH_ENCODINGS = {
    # variant name -> (value for rows on the overview, value for rows that are not)
    "yes_no": ("Yes", "No"),
    "bool": (True, False),
    "yn": ("Y", "N"),
    "10": (1, 0),
    "x": ("X", ""),
    "blank": ("", ""),
}


def _history_rows(as_of: dt.date) -> List[list]:
    """Three prior weekly meetings so update reminders trigger for some deals."""
    rows = []
    meetings = [as_of - dt.timedelta(days=7 * k) for k in (3, 2, 1)]
    statuses = {
        "Bridge::10001": ["Surveillance", "Surveillance", "Surveillance"],
        "Bridge::10002": ["Surveillance", "Upcoming Maturity", "Upcoming Maturity"],
        "Bridge::10003": ["Surveillance", "Surveillance", "Modification"],
        "Bridge::10004": ["Foreclosure", "Foreclosure", "Foreclosure"],
        "Bridge::10007": ["Litigation", "Litigation", "Litigation"],
        "Term::20001": ["Surveillance", "Surveillance", "Surveillance"],
        "Term::20003": ["Surveillance", "Receivership", "Receivership"],
    }
    for idx, meeting in enumerate(meetings):
        for deal_key, seq in statuses.items():
            sheet, deal_number = deal_key.split("::")
            rows.append([
                meeting.isoformat(), sheet, deal_number, deal_key, seq[idx],
                "Alice Chen" if sheet == "Bridge" else "Erin Fox",
                f"Commentary from {meeting:%m/%d}", "Portfolio Overview.xlsx",
                f"{meeting.isoformat()} 10:00:00",
            ])
    return rows


def _rename_sheet(wb: Workbook, old: str, new: str) -> None:
    """Rename via a temporary title: openpyxl would turn 'Term' -> 'TERM' into 'TERM1'."""
    ws = wb[old]
    ws.title = "__tmp_rename__"
    ws.title = new


def _bridge_row(as_of: dt.date, i: int, deal: tuple) -> list:
    num, name, borrower, pp, status, comm, commit, funded, mat_off, pay_off, dpd, npl, _hidden = deal
    maturity = as_of + dt.timedelta(days=mat_off)
    payment = as_of + dt.timedelta(days=pay_off)
    # Mix real dates and raw Excel serials, as seen in exported workbooks.
    maturity_cell = _excel_serial(maturity) if i % 3 == 1 else maturity
    return [
        num, name, borrower, borrower, "Servicer A" if i % 2 else "Servicer B", "Bridge",
        "Fix & Flip" if i % 2 else "Construction", "Warehouse" if i % 2 else "Balance Sheet",
        "Buyer X" if i % 3 == 0 else "", pp, status, comm, commit, funded, commit - funded,
        funded * 0.98, funded, "", npl, dpd, maturity_cell, payment,
    ]


def build_sample_workbook(
    as_of: Optional[dt.date] = None,
    include_history: bool = True,
    *,
    duplicate_deal: Optional[str] = None,
    junk_rows: bool = False,
    duplicate_headers: bool = False,
    term_owner_header: Optional[str] = None,
    sheet_titles: Optional[Mapping[str, str]] = None,
    term_sheet: str = "normal",
    bridge_header_row: Optional[int] = None,
    bad_dates: bool = False,
    mods_key: str = "both",
    match_encoding: str = "yes_no",
    sf_utc_stamps: bool = False,
    cancelled_latest_mod: bool = False,
    unnumbered_new_mod: bool = False,
    formula_deal_numbers: Sequence[str] = (),
    styled_empty_rows: int = 0,
    term_commitment: Optional[Mapping[str, Optional[float]]] = None,
    term_loan_size: str = "normal",
    cell_overrides: Optional[Mapping[tuple, object]] = None,
) -> Workbook:
    """Build the sample workbook; every variant flag defaults to the plain sample.

    duplicate_deal: 'visible' | 'hidden' | 'adjacent' -- another Bridge row for 10002
        ('Maple Street Rehab - Phase 2', deal number written as 10002.0). 'visible' and
        'hidden' append it after the Bridge total row (Excel row 14); 'adjacent' inserts
        it right below the original (row 7).
    junk_rows: the Bridge total row gets 'Total' in Deal Number and a 44,000,000 UPB,
        and a '#N/A' error row is added; Term gets a 'Grand Total' row.
    duplicate_headers: Bridge gets a second 'Status' column (values 'Funded') and a
        second, blank 'Days Past Due' column at the far right.
    term_owner_header: rename the Term 'Asset Manager' header (e.g. 'Active RM').
    sheet_titles: rename tabs, e.g. {'Term': 'TERM ', 'Loan Modifications': 'Loan modifications'}.
    term_sheet: 'normal' | 'blank' (empty tab) | 'note' (a note, no header) | 'missing'.
    bridge_header_row: push the Bridge header down to this row with note rows.
    bad_dates: a Salesforce UTC string in a Term Next Payment Date, a 3026 Term maturity,
        a '12/31/9999' loan-mod Updated Maturity Date and a UTC System Task Completed Date.
    mods_key: 'both' | 'opportunity_only' (no 'Deal Number' column) | 'zero_padded'
        (Deal Number text like '00010002').
    match_encoding: key of MATCH_ENCODINGS for 'Portfolio Overview Match'.
    sf_utc_stamps: Created/Last Modified Date as 'YYYY-MM-DD HH:MM:SS UTC' strings.
    cancelled_latest_mod / unnumbered_new_mod: extra loan mods (see constants above).
    formula_deal_numbers: Bridge deal numbers written as formulas without cached values.
    styled_empty_rows: border-styled empty rows below the Bridge data (slow-path tests;
        for tens of thousands of rows use with_styled_empty_rows(), which is much faster).
    term_commitment: add a Term 'Loan Commitment' column, {deal_number: value or None}.
    term_loan_size: how the Term sheet carries loan size: 'normal' ('Loan Amount' numbers) |
        'renamed' (header 'Orig. Loan Amt ($)') | 'text' (values as '$22,000,000.00' text) |
        'missing' (no loan size column at all) | 'blank_first' (a blank 'Loan Amount' column
        plus an 'Original Loan Amount' column at the far right) | 'formula' (loan amounts as
        formulas with no saved value) | 'zero_commitment' (an extra 'Total Commitment'
        column of zeros).
    cell_overrides: {(sheet, deal_number, header): value} set after the book is built, e.g.
        {('Bridge', '10001', '10/15 NPL'): datetime(2026, 10, 1)}; '=...' text is a formula.
    """
    as_of = as_of or dt.date.today()
    prev = as_of - dt.timedelta(days=7)
    wb = Workbook()

    # ---------------- Bridge ----------------
    ws = wb.active
    ws.title = "Bridge"
    ws.append([f"Portfolio Overview - Bridge as of {as_of:%m/%d/%Y}"])
    ws.append(["Confidential - internal use only"])
    ws.append([])
    rows_above_header = 3
    if bridge_header_row:
        while rows_above_header < bridge_header_row - 1:
            rows_above_header += 1
            ws.append([f"Note row {rows_above_header}: see the deal memos for context."])
    bridge_headers = [
        "Deal Number", "Deal Name", "Borrower Name", "Account Name", "Servicer", "Portfolio",
        "Segment", "Financing", "Loan Buyer", "Point Person", "Status", "AM Commentary",
        "Loan Commitment", "Active Funded Amount", "Remaining Commitment",
        _hdr(prev, "UPB"), _hdr(as_of, "UPB"), _hdr(prev, "NPL"), _hdr(as_of, "NPL"),
        "Days Past Due", "Current Maturity Date", "Next Payment Date",
    ]
    if duplicate_headers:
        bridge_headers += ["Status", "Days Past Due"]
    ws.append(bridge_headers)
    for i, deal in enumerate(BRIDGE_DEALS):
        row = _bridge_row(as_of, i, deal)
        if duplicate_headers:
            row += ["Funded", None]
        if str(deal[0]).replace(".0", "") in set(formula_deal_numbers):
            row[0] = f"={str(deal[0]).replace('.0', '')}"
        ws.append(row)
        if deal[12]:
            ws.row_dimensions[ws.max_row].hidden = True
        if duplicate_deal == "adjacent" and str(deal[0]) == "10002":
            dup = list(row)
            dup[0], dup[1] = 10002.0, "Maple Street Rehab - Phase 2"
            ws.append(dup)
    # A subtotal row with no deal number must be skipped.
    total_row = ["", "Bridge Total", "", "", "", "", "", "", "", "", "", "",
                 sum(d[6] for d in BRIDGE_DEALS), sum(d[7] for d in BRIDGE_DEALS)]
    if junk_rows:
        total_row[0] = "Total"
        total_row += [None, None, 44_000_000]
    ws.append(total_row)
    if junk_rows:
        error_row = _bridge_row(as_of, 9, ("#N/A", "Broken Lookup Deal", "Unknown Borrower", "", "", "",
                                           1_000_000, 1_000_000, 50, 10, 0, "", False))
        ws.append(error_row)
    if duplicate_deal in {"visible", "hidden"}:
        dup = _bridge_row(as_of, 1, BRIDGE_DEALS[1])
        if duplicate_headers:
            dup += ["Funded", None]
        dup[0], dup[1] = 10002.0, "Maple Street Rehab - Phase 2"
        ws.append(dup)
        if duplicate_deal == "hidden":
            ws.row_dimensions[ws.max_row].hidden = True
    if styled_empty_rows:
        border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
        start = ws.max_row + 1
        for r in range(start, start + styled_empty_rows):
            for c in range(1, len(bridge_headers) + 1):
                ws.cell(r, c).border = border

    # ---------------- Term ----------------
    ws = wb.create_sheet("Term")
    if term_sheet == "normal":
        ws.append([f"Portfolio Overview - Term as of {as_of:%m/%d/%Y}"])
        ws.append([])
        loan_amount_header = "Orig. Loan Amt ($)" if term_loan_size == "renamed" else "Loan Amount"
        term_headers = [
            "Deal Number", "Loan Name", "Borrower Entity", term_owner_header or "Asset Manager", "Servicer",
            "Segment", "Status", "AM Commentary", loan_amount_header, _hdr(prev, "UPB", True),
            _hdr(as_of, "UPB", True), _hdr(as_of, "NPL", True), "Days Past Due", "Maturity Date", "Next Payment Date",
        ]
        if term_commitment is not None:
            term_headers.append("Loan Commitment")
        if term_loan_size == "blank_first":
            term_headers.append("Original Loan Amount")
        if term_loan_size == "zero_commitment":
            term_headers.append("Total Commitment")
        if term_loan_size == "missing":
            del term_headers[8]
        ws.append(term_headers)
        for idx, (num, name, borrower, am, status, comm, amount, mat_off, pay_off, dpd, npl) in enumerate(TERM_DEALS):
            maturity: object = as_of + dt.timedelta(days=mat_off)
            payment: object = as_of + dt.timedelta(days=pay_off)
            if bad_dates and idx == 0:
                payment = "2026-10-30T00:00:00.000+0000"
            if bad_dates and idx == 1:
                maturity = dt.datetime(3026, 6, 1)
            amount_cell: object = amount
            if term_loan_size == "text":
                amount_cell = f"${amount:,.2f}"
            elif term_loan_size == "formula":
                amount_cell = f"={amount}"
            elif term_loan_size == "blank_first":
                amount_cell = None
            row = [
                num, name, borrower, am, "Servicer C", "DSCR", status, comm, amount_cell,
                amount * 0.995, amount * 0.99, npl, dpd, maturity, payment,
            ]
            if term_commitment is not None:
                row.append(term_commitment.get(num))
            if term_loan_size == "blank_first":
                row.append(amount)
            if term_loan_size == "zero_commitment":
                row.append(0)
            if term_loan_size == "missing":
                del row[8]
            ws.append(row)
        if junk_rows:
            ws.append(["Grand Total", "", "", "", "", "", "", "", sum(d[6] for d in TERM_DEALS),
                       None, sum(d[6] for d in TERM_DEALS) * 0.99])
    elif term_sheet == "note":
        ws.append(["No Term loans on the watch list this week."])
    elif term_sheet == "missing":
        wb.remove(ws)

    # ---------------- Loan Modifications ----------------
    ws = wb.create_sheet("Loan Modifications")
    ws.append(["Loan modifications pulled from Salesforce"])
    mod_headers = [
        "Portfolio Overview Match", "Visible Overview Sheet(s)", "Deal Number", "Deal Name",
        "Opportunity Deal Loan Number", "Opportunity Deal Name", "Loan Mod Order Number",
        "Modification Type", "System Task Completed Date", "Mod Effective Date",
        "Previous Maturity Date", "Updated Maturity Date", "Previous Loan Commitment",
        "Updated Loan Commitment", "Comments", "Loan Modification Name", "Status", "Loan Mod Type",
        "Mod Reporting Type", "Modification Finalized Date", "Deal Type", "Last Modified Date",
        "Created Date", "Irrelevant Salesforce Field",
    ]
    drop_deal_number = mods_key == "opportunity_only"
    if drop_deal_number:
        mod_headers.remove("Deal Number")
    ws.append(mod_headers)
    names = {str(d[0]).replace(".0", ""): d[1] for d in BRIDGE_DEALS}
    names.update({d[0]: d[1] for d in TERM_DEALS})
    match_yes, match_no = MATCH_ENCODINGS[match_encoding]

    def _d(offset):
        return None if offset is None else as_of + dt.timedelta(days=offset)

    def _stamp(offset):
        day = _d(offset)
        if sf_utc_stamps:
            return f"{day:%Y-%m-%d} 14:22:10 UTC"
        return day

    mods = list(LOAN_MODS)
    if cancelled_latest_mod:
        mods = mods[:-1] + CANCELLED_LATEST_MODS + mods[-1:]
    if unnumbered_new_mod:
        mods = [UNNUMBERED_NEW_MOD] + mods
    for num, order, mtype, status, comp, eff, pmat, nmat, pcom, ncom, comments in mods:
        on_overview = num in names
        deal_cell: object = num
        if mods_key == "zero_padded":
            deal_cell = num.zfill(8)
        completed: object = _d(comp)
        updated_maturity: object = _d(nmat)
        if bad_dates and num == "10002" and order == "2":
            updated_maturity = "12/31/9999"
        if bad_dates and num == "10007" and order == "3":
            completed = "2026-09-05T17:22:10.000+0000"
        created_offset = 0 if order == "" else -300
        row = [
            match_yes if on_overview else match_no, "Bridge" if num.startswith("1") else "Term", deal_cell,
            names.get(num, "Unknown Deal"), num, names.get(num, "Unknown Deal"), order, mtype,
            completed, _d(eff), _d(pmat), updated_maturity, pcom, ncom, comments,
            f"LM-{num}-{order or 'new'}", status, mtype, "Material", _d(comp), "Bridge", _stamp(-1),
            _stamp(created_offset), "x",
        ]
        if drop_deal_number:
            del row[2]
        ws.append(row)

    # ---------------- Hidden status history ----------------
    if include_history:
        ws = wb.create_sheet("_Status History")
        ws.sheet_state = "hidden"
        ws.append(["meeting_date", "sheet", "deal_number", "deal_key", "status", "owner",
                   "commentary", "workbook_name", "saved_at"])
        for row in _history_rows(as_of):
            ws.append(row)

    if term_sheet == "blank":
        pass  # the Term tab exists with no cells at all
    if cell_overrides:
        _apply_cell_overrides(wb, cell_overrides)
    for old, new in (sheet_titles or {}).items():
        _rename_sheet(wb, old, new)
    return wb


def _apply_cell_overrides(wb: Workbook, overrides: Mapping[tuple, object]) -> None:
    """Set cells by (sheet, deal number, header text) on the sheets built above."""
    for (sheet, deal_number, header), value in overrides.items():
        ws = wb[sheet]
        header_row = next(
            r for r in range(1, 61) if any(ws.cell(r, c).value == "Deal Number" for c in range(1, ws.max_column + 1))
        )
        columns = {ws.cell(header_row, c).value: c for c in range(ws.max_column, 0, -1)}  # left-most wins
        column = columns[header]
        target = next(
            r for r in range(header_row + 1, ws.max_row + 1)
            if str(ws.cell(r, columns["Deal Number"]).value).replace(".0", "") == str(deal_number)
        )
        ws.cell(target, column).value = value


def workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_sample_workbook_bytes(
    as_of: Optional[dt.date] = None,
    include_history: bool = True,
    **variants,
) -> bytes:
    return workbook_to_bytes(build_sample_workbook(as_of=as_of, include_history=include_history, **variants))


def with_cached_formula_values(xlsx_bytes: bytes, sheet_title: str, values: Mapping[str, object]) -> bytes:
    """Give formula cells a cached value, as Excel would after a save.

    openpyxl writes formulas without cached values (<f>..</f><v/>); this fills the <v>
    of the given cells (e.g. {'A11': 10007}) in the named sheet's XML part.
    """
    source = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    part = _sheet_part(source, sheet_title)
    xml = source.read(part).decode("utf-8")
    for ref, value in values.items():
        pattern = re.compile(r'(<c r="' + re.escape(ref) + r'"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/>(</c>)')
        xml, count = pattern.subn(lambda m: f"{m.group(1)}<v>{value}</v>{m.group(2)}", xml)
        if count != 1:
            raise ValueError(f"no formula cell {ref} in {sheet_title}")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source.infolist():
            data = xml.encode("utf-8") if item.filename == part else source.read(item.filename)
            target_zip.writestr(item, data)
    return out.getvalue()


def _sheet_part(source: zipfile.ZipFile, sheet_title: str) -> str:
    workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
    rels_xml = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_id = re.search(r'<sheet[^>]*name="' + re.escape(sheet_title) + r'"[^>]*r:id="([^"]+)"', workbook_xml).group(1)
    target = re.search(r'<Relationship[^>]*Id="' + re.escape(rel_id) + r'"[^>]*Target="([^"]+)"', rels_xml)
    if target is None:
        target = re.search(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="' + re.escape(rel_id) + r'"', rels_xml)
    part = target.group(1).lstrip("/")
    return part if part.startswith("xl/") else f"xl/{part}"


def with_styled_empty_rows(xlsx_bytes: bytes, sheet_title: str, n_rows: int, n_cols: int = 22) -> bytes:
    """Append n_rows formatted-but-empty rows (cells with a style and no value) to a sheet.

    Writes the sheet XML directly, so 60,000 rows take about a second instead of the
    minutes openpyxl needs; the sheet dimension grows to cover them, as it would after
    someone formats whole columns in Excel.
    """
    source = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    part = _sheet_part(source, sheet_title)
    xml = source.read(part).decode("utf-8")
    rows = re.findall(r'<row[^>]*\br="(\d+)"', xml)
    first = (max(int(r) for r in rows) if rows else 0) + 1
    letters = []
    for index in range(1, n_cols + 1):
        name, number = "", index
        while number:
            number, remainder = divmod(number - 1, 26)
            name = chr(65 + remainder) + name
        letters.append(name)
    chunks = []
    for r in range(first, first + n_rows):
        cells = "".join(f'<c r="{letter}{r}" s="0"/>' for letter in letters)
        chunks.append(f'<row r="{r}">{cells}</row>')
    if "</sheetData>" in xml:
        xml = xml.replace("</sheetData>", "".join(chunks) + "</sheetData>", 1)
    else:
        xml = xml.replace("<sheetData/>", "<sheetData>" + "".join(chunks) + "</sheetData>", 1)
    last = first + n_rows - 1
    xml = re.sub(r'<dimension ref="([A-Z]+\d+)(?::[A-Z]+\d+)?"', lambda m: f'<dimension ref="{m.group(1)}:{letters[-1]}{last}"', xml, count=1)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source.infolist():
            data = xml.encode("utf-8") if item.filename == part else source.read(item.filename)
            target_zip.writestr(item, data)
    return out.getvalue()


STATUS_CYCLE = ["Surveillance", "Upcoming Maturity", "Modification", "Foreclosure", "Litigation",
                "Receivership", "Needs Follow-Up"]
OWNERS = ["Alice Chen", "Bob Diaz", "Carla Gomez", "Dan Evans", "Erin Fox", "Frank Green"]


def build_scaled_workbook(
    n_bridge: int = 120,
    n_term: int = 30,
    n_weeks: int = 52,
    as_of: Optional[dt.date] = None,
    seed: int = 7,
) -> Workbook:
    """A large book in the sample's layout: n_bridge + n_term deals, n_weeks of weekly
    _Status History (one row per deal per meeting) and a few loan mods per deal."""
    as_of = as_of or dt.date.today()
    prev = as_of - dt.timedelta(days=7)
    rng = random.Random(seed)
    wb = Workbook()

    ws = wb.active
    ws.title = "Bridge"
    ws.append([f"Portfolio Overview - Bridge as of {as_of:%m/%d/%Y}"])
    ws.append([])
    ws.append([
        "Deal Number", "Deal Name", "Borrower Name", "Account Name", "Servicer", "Portfolio",
        "Segment", "Financing", "Loan Buyer", "Point Person", "Status", "AM Commentary",
        "Loan Commitment", "Active Funded Amount", "Remaining Commitment",
        _hdr(prev, "UPB"), _hdr(as_of, "UPB"), _hdr(as_of, "NPL"),
        "Days Past Due", "Current Maturity Date", "Next Payment Date",
    ])
    statuses: Dict[str, str] = {}
    for i in range(n_bridge):
        number = str(30001 + i)
        commit = rng.randrange(500_000, 25_000_000, 50_000)
        funded = int(commit * rng.uniform(0.4, 1.0))
        status = rng.choice(STATUS_CYCLE)
        statuses[f"Bridge::{number}"] = status
        dpd = rng.choice([0, 0, 0, 0, 5, 35, 95])
        ws.append([
            number, f"Bridge Deal {number}", f"Borrower {number} LLC", f"Borrower {number} LLC",
            rng.choice(["Servicer A", "Servicer B"]), "Bridge", rng.choice(["Fix & Flip", "Construction"]),
            "Warehouse", "", rng.choice(OWNERS), status, f"Line one for {number}.\nLine two.",
            commit, funded, commit - funded, funded * 0.99, funded, "30+" if dpd >= 30 else "", dpd,
            as_of + dt.timedelta(days=rng.randrange(-200, 900)), as_of + dt.timedelta(days=rng.randrange(-30, 30)),
        ])

    ws = wb.create_sheet("Term")
    ws.append([f"Portfolio Overview - Term as of {as_of:%m/%d/%Y}"])
    ws.append([])
    ws.append([
        "Deal Number", "Loan Name", "Borrower Entity", "Asset Manager", "Servicer", "Segment", "Status",
        "AM Commentary", "Loan Amount", _hdr(prev, "UPB", True), _hdr(as_of, "UPB", True),
        _hdr(as_of, "NPL", True), "Days Past Due", "Maturity Date", "Next Payment Date",
    ])
    for i in range(n_term):
        number = str(40001 + i)
        amount = rng.randrange(300_000, 30_000_000, 25_000)
        status = rng.choice(STATUS_CYCLE)
        statuses[f"Term::{number}"] = status
        ws.append([
            number, f"Term Loan {number}", f"Entity {number} LLC", rng.choice(OWNERS), "Servicer C", "DSCR",
            status, f"Performing loan {number}.", amount, amount * 0.996, amount * 0.995, "", 0,
            as_of + dt.timedelta(days=rng.randrange(300, 3000)), as_of + dt.timedelta(days=rng.randrange(-5, 30)),
        ])

    ws = wb.create_sheet("Loan Modifications")
    ws.append(["Loan modifications pulled from Salesforce"])
    ws.append([
        "Portfolio Overview Match", "Visible Overview Sheet(s)", "Deal Number", "Deal Name",
        "Opportunity Deal Loan Number", "Opportunity Deal Name", "Loan Mod Order Number",
        "Modification Type", "System Task Completed Date", "Mod Effective Date",
        "Previous Maturity Date", "Updated Maturity Date", "Previous Loan Commitment",
        "Updated Loan Commitment", "Comments", "Loan Modification Name", "Status", "Loan Mod Type",
        "Mod Reporting Type", "Modification Finalized Date", "Deal Type", "Last Modified Date", "Created Date",
    ])
    for key in statuses:
        if rng.random() > 0.3:
            continue
        sheet, number = key.split("::")
        for order in range(rng.randrange(1, 4), 0, -1):
            done = as_of - dt.timedelta(days=60 * order)
            ws.append([
                "Yes", sheet, number, f"Deal {number}", number, f"Deal {number}", str(order),
                rng.choice(["Maturity Extension", "Rate Modification", "Forbearance"]), done, done,
                done, done + dt.timedelta(days=180), 1_000_000, 1_100_000, "1. Extended 2. Fee paid",
                f"LM-{number}-{order}", rng.choice(["Completed", "Completed", "Cancelled"]), "Extension",
                "Material", done, sheet, f"{done:%Y-%m-%d} 10:00:00 UTC", f"{done:%Y-%m-%d} 09:00:00 UTC",
            ])

    ws = wb.create_sheet("_Status History")
    ws.sheet_state = "hidden"
    ws.append(["meeting_date", "sheet", "deal_number", "deal_key", "status", "owner",
               "commentary", "workbook_name", "saved_at"])
    current = dict(statuses)
    weekly: List[List[object]] = []
    for week in range(1, n_weeks + 1):
        meeting = as_of - dt.timedelta(days=7 * week)
        for key in statuses:
            if rng.random() < 0.08:
                current[key] = rng.choice(STATUS_CYCLE)
            sheet, number = key.split("::")
            weekly.append([meeting.isoformat(), sheet, number, key, current[key], rng.choice(OWNERS),
                           f"Note {week} for {number}", "Portfolio Overview.xlsx", f"{meeting.isoformat()} 10:00:00"])
    for row in reversed(weekly):
        ws.append(row)
    return wb


def build_scaled_workbook_bytes(
    n_bridge: int = 120,
    n_term: int = 30,
    n_weeks: int = 52,
    as_of: Optional[dt.date] = None,
    seed: int = 7,
) -> bytes:
    return workbook_to_bytes(build_scaled_workbook(n_bridge, n_term, n_weeks, as_of=as_of, seed=seed))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", nargs="?", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--as-of", default=None, help="As-of date (YYYY-MM-DD); defaults to today")
    parser.add_argument("--no-history", action="store_true")
    args = parser.parse_args()
    as_of = dt.date.fromisoformat(args.as_of) if args.as_of else None
    Path(args.output).write_bytes(build_sample_workbook_bytes(as_of, include_history=not args.no_history))
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
