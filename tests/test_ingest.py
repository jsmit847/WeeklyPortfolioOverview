"""Workbook ingest: robust dates, the shared sheet layout, bounded reading, tolerant sheet
names, junk rows, duplicate deal keys, dated columns and the as-of date, NaN DPD, loan
mods, formula cells, loan size, and the parse/derive split (all on synthetic data)."""
from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import sample_workbook as sw
from app_helpers import call_counts, upload_workbook

AS_OF = dt.date(2026, 10, 15)
SAMPLE_TOTAL_UPB = 66_769_850


def _parse(app_module, **variants):
    return app_module.parse_workbook_bytes(sw.build_sample_workbook_bytes(as_of=AS_OF, **variants))


def _deck(app_module, parsed, as_of=AS_OF, include_hidden=False):
    return app_module.build_deck(parsed, as_of, include_hidden)


def _row(deck, deal_number):
    rows = deck[deck["deal_number"] == deal_number]
    assert len(rows) == 1, (deal_number, len(rows))
    return rows.iloc[0]


def _notes(app_module, parsed, level=None):
    return [message for lvl, message in app_module.workbook_notes(parsed) if level is None or lvl == level]


# -----------------------------------------------------------------------------
# Default sample: the contract numbers
# -----------------------------------------------------------------------------
def test_default_sample_contract_numbers(app_module, sample_bytes):
    parsed = app_module.parse_workbook_bytes(sample_bytes)
    assert parsed["error"] is None
    assert parsed["reference_date"] == AS_OF.isoformat() and parsed["reference_source"] == "title"
    deck, info = _deck(app_module, parsed)
    assert len(deck) == 10
    assert deck["upb"].sum() == SAMPLE_TOTAL_UPB
    assert info["sheets"]["Bridge"]["upb_header"] == "10/15 UPB"
    assert info["sheets"]["Term"]["upb_header"] == "10/15/26 UPB"
    assert deck["deal_key"].is_unique
    assert not deck["is_duplicate"].any()
    assert app_module.workbook_notes(parsed) == []
    due, past = app_module.maturity_buckets(deck["days_to_maturity"])
    assert (due, past) == (2, 2)
    with_hidden, _ = _deck(app_module, parsed, include_hidden=True)
    assert len(with_hidden) == 11 and "Bridge::10006" in set(with_hidden["deal_key"])


# -----------------------------------------------------------------------------
# 1. Robust dates (U06)
# -----------------------------------------------------------------------------
def test_bad_dates_load_and_are_counted(app_module):
    parsed = _parse(app_module, bad_dates=True)
    assert parsed["error"] is None
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 10
    assert _row(deck, "20001")["next_payment_date"] == pd.Timestamp("2026-10-30")  # Salesforce UTC string
    assert pd.isna(_row(deck, "20002")["maturity_date"])  # 3026-06-01 typo
    assert parsed["rejected_dates"] >= 2
    assert parsed["sheets"]["Term"]["rejected_dates"][0]["header"] == "Maturity Date"
    mods = parsed["loan_mods"]
    sentinel = mods[(mods["Deal Number"] == "10002") & (mods["Loan Mod Order Number"] == "2")].iloc[0]
    assert pd.isna(sentinel["Updated Maturity Date"])  # 12/31/9999
    utc = mods[(mods["Deal Number"] == "10007") & (mods["Loan Mod Order Number"] == "3")].iloc[0]
    assert utc["System Task Completed Date"] == pd.Timestamp("2026-09-05 17:22:10")
    for column in ("maturity_date", "next_payment_date", "loan_mod_latest_updated_maturity_date"):
        assert str(deck[column].dtype) == "datetime64[ns]", column
    assert any("unreadable or out-of-range date" in note for note in _notes(app_module, parsed))


def test_salesforce_utc_stamps_and_sentinels_in_loan_mods(app_module):
    parsed = _parse(app_module, sf_utc_stamps=True)
    mods = parsed["loan_mods"]
    assert str(mods["Created Date"].dtype) == "datetime64[ns]"
    assert mods["Created Date"].dt.tz is None
    assert mods["Last Modified Date"].notna().all()
    deck, _ = _deck(app_module, parsed)
    assert _row(deck, "10007")["loan_mod_count"] == 3


# -----------------------------------------------------------------------------
# 2. One layout for loader and writer (U18 / U23)
# -----------------------------------------------------------------------------
def test_active_rm_owner_edit_round_trips(app_module):
    data = sw.build_sample_workbook_bytes(as_of=AS_OF, term_owner_header="Active RM")
    deck, meta = app_module.load_portfolio_workbook(data, AS_OF.isoformat(), False)
    assert meta["Term"]["owner_header"] == "Active RM"
    row = _row(deck, "20001")
    assert row["owner"] == "Erin Fox"
    override = {
        "sheet": "Term", "deal_number": "20001", "status": row["status"], "owner": "Gina Hart",
        "commentary": row["commentary"],
    }
    exported = app_module.update_workbook_bytes(data, {"Term::20001": override})
    reloaded, _ = app_module.load_portfolio_workbook(exported, AS_OF.isoformat(), False)
    assert _row(reloaded, "20001")["owner"] == "Gina Hart"
    assert _row(reloaded, "20002")["owner"] == "Erin Fox"


def test_duplicate_headers_read_and_write_the_first_column(app_module):
    data = sw.build_sample_workbook_bytes(as_of=AS_OF, duplicate_headers=True)
    parsed = app_module.parse_workbook_bytes(data)
    layout = parsed["sheets"]["Bridge"]["layout"]
    assert layout["roles"]["status"] == "Status" and "Status (2)" in layout["labels"]
    assert len(layout["dpd_cols"]) == 2
    deck, _ = _deck(app_module, parsed)
    bridge = deck[deck["sheet"] == "Bridge"]
    expected = {"10001": "Surveillance", "10002": "Upcoming Maturity", "10003": "Modification",
                "10004": "Foreclosure", "10005": "", "10007": "Litigation", "10008": "Surveillance"}
    assert dict(zip(bridge["deal_number"], bridge["status"])) == expected
    assert _row(deck, "10004")["days_past_due"] == 95
    assert any("'Status' appears in columns K, W" in note for note in _notes(app_module, parsed, "warning"))

    row = _row(deck, "10001")
    override = {"sheet": "Bridge", "deal_number": "10001", "status": "Watchlist", "owner": row["owner"],
                "commentary": row["commentary"]}
    exported = app_module.update_workbook_bytes(data, {"Bridge::10001": override})
    ws = load_workbook(io.BytesIO(exported))["Bridge"]
    header_row = layout["header_row"]
    status_cols = [c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value == "Status"]
    assert status_cols == [11, 23]
    assert ws.cell(row["_excel_row"], 11).value == "Watchlist"
    assert ws.cell(row["_excel_row"], 23).value == "Funded"


def test_writer_and_loader_share_resolve_sheet_layout(app_module, sample_bytes):
    read_only = load_workbook(io.BytesIO(sample_bytes), read_only=True, data_only=True)
    normal = load_workbook(io.BytesIO(sample_bytes))
    for sheet in ("Bridge", "Term"):
        a = app_module.resolve_sheet_layout(read_only[sheet])
        b = app_module.resolve_sheet_layout(normal[sheet])
        assert a["header_row"] == b["header_row"] and a["role_cols"] == b["role_cols"], sheet
    assert app_module.resolve_sheet_layout(normal["Bridge"])["header_row"] == 4
    assert app_module.resolve_sheet_layout(normal["Loan Modifications"]) is not None  # has 'Deal Number'
    blank = Workbook().active
    assert app_module.resolve_sheet_layout(blank) is None


# -----------------------------------------------------------------------------
# 3. Bounded reading
# -----------------------------------------------------------------------------
def test_reading_stops_after_blank_rows(app_module):
    data = sw.with_styled_empty_rows(sw.build_sample_workbook_bytes(as_of=AS_OF), "Bridge", 800)
    parsed = app_module.parse_workbook_bytes(data)
    report = parsed["sheets"]["Bridge"]
    assert report["blank_stop_row"] is not None and report["values_below_stop"] is False
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 10 and deck["upb"].sum() == SAMPLE_TOTAL_UPB
    assert app_module.workbook_notes(parsed) == []  # formatted-but-empty rows are not news


def test_values_below_a_long_gap_are_reported(app_module):
    wb = sw.build_sample_workbook(as_of=AS_OF)
    ws = wb["Bridge"]
    ws.cell(ws.max_row + app_module.BLANK_ROW_STOP + 50, 1).value = "10099"
    parsed = app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb))
    assert parsed["sheets"]["Bridge"]["values_below_stop"] is True
    deck, _ = _deck(app_module, parsed)
    assert "10099" not in set(deck["deal_number"])
    assert any("were not read" in note for note in _notes(app_module, parsed, "warning"))


def test_layout_is_bounded_by_the_last_header_column(app_module):
    wb = sw.build_sample_workbook(as_of=AS_OF)
    ws = wb["Bridge"]
    ws.cell(5, 40).value = "stray note far to the right"
    parsed = app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb))
    layout = parsed["sheets"]["Bridge"]["layout"]
    assert layout["max_col"] == 22
    assert list(parsed["sheets"]["Bridge"]["raw"].columns[-2:]) == ["_excel_row", "_hidden"]


# -----------------------------------------------------------------------------
# 4. Tolerant sheet names; sheets without a header (U16 / U17)
# -----------------------------------------------------------------------------
def test_renamed_tabs_load(app_module):
    parsed = _parse(app_module, sheet_titles={"Term": "TERM ", "Loan Modifications": "Loan modifications"})
    deck, _ = _deck(app_module, parsed)
    assert int((deck["sheet"] == "Term").sum()) == 3
    assert set(deck.loc[deck["sheet"] == "Term", "deal_key"]) == {"Term::20001", "Term::20002", "Term::20003"}
    assert len(parsed["loan_mods"]) == 7
    assert parsed["sheets"]["Term"]["tab"] == "TERM " and parsed["loan_mod_info"]["tab"] == "Loan modifications"
    captions = _notes(app_module, parsed, "caption")
    assert "Loaded tab 'TERM ' as Term." in captions
    history = _parse(app_module, sheet_titles={"_Status History": "_status history"})
    assert len(app_module.embedded_status_history(history)) == 21


@pytest.mark.parametrize("term_sheet, phrase", [("blank", "tab is empty"), ("note", "no 'Deal Number' header")])
def test_term_sheet_without_header_still_loads_bridge(app_module, term_sheet, phrase):
    parsed = _parse(app_module, term_sheet=term_sheet)
    assert parsed["error"] is None
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 7 and set(deck["sheet"]) == {"Bridge"}
    assert any(phrase in note for note in _notes(app_module, parsed, "warning"))


def test_missing_term_tab_is_reported(app_module):
    parsed = _parse(app_module, term_sheet="missing")
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 7
    warning = next(n for n in _notes(app_module, parsed, "warning") if "No 'Term' tab" in n)
    assert "'Bridge'" in warning and "_Status History" not in warning


def test_header_on_row_32_is_found(app_module):
    parsed = _parse(app_module, bridge_header_row=32)
    assert parsed["sheets"]["Bridge"]["layout"]["header_row"] == 32
    deck, _ = _deck(app_module, parsed)
    assert int((deck["sheet"] == "Bridge").sum()) == 7


def test_ambiguous_tab_names_load_nothing_and_warn(app_module):
    wb = sw.build_sample_workbook(as_of=AS_OF)
    wb.copy_worksheet(wb["Term"]).title = "term "
    sw._rename_sheet(wb, "Term", "TERM")
    parsed = app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb))
    assert parsed["sheets"]["Term"]["status"] == "ambiguous"
    assert any("Several tabs could be Term" in note for note in _notes(app_module, parsed, "warning"))


# -----------------------------------------------------------------------------
# 5. Junk rows (U19)
# -----------------------------------------------------------------------------
def test_total_and_error_rows_are_skipped(app_module):
    parsed = _parse(app_module, junk_rows=True)
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 10
    assert deck["upb"].sum() == SAMPLE_TOTAL_UPB
    assert not deck["deal_number"].isin(["Total", "#N/A", "Grand Total"]).any()
    assert parsed["sheets"]["Bridge"]["skipped_count"] == 2
    assert "3 non-deal rows skipped (Total, #N/A, Grand Total)" in _notes(app_module, parsed, "caption")


@pytest.mark.parametrize(
    "value, kind",
    [("Total", "total"), ("Grand Total", "total"), ("Sub-Total", "total"), ("Subtotal", "total"),
     ("Bridge Total", "total"), ("TOTALS", "total"), ("#N/A", "error"), ("#REF!", "error"), ("#DIV/0!", "error"),
     ("#VALUE!", "error"), ("#NAME?", "error"), ("#NULL!", "error"), ("#NUM!", "error"), ("#SPILL!", "error"),
     (10001.0, "deal"), ("BR-10009", "deal"), ("", "blank"), (None, "blank")],
)
def test_deal_cell_classifier(app_module, value, kind):
    assert app_module._classify_deal_cell(value)[0] == kind


# -----------------------------------------------------------------------------
# 6. Duplicate deal numbers get their own key (U12)
# -----------------------------------------------------------------------------
@pytest.mark.parametrize("mode, expected_rows", [("visible", 11), ("adjacent", 11), ("hidden", 10)])
def test_duplicate_deal_numbers_get_row_keys(app_module, mode, expected_rows):
    parsed = _parse(app_module, duplicate_deal=mode)
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == expected_rows
    assert deck["deal_key"].is_unique
    full, _ = _deck(app_module, parsed, include_hidden=True)
    copies = full[full["deal_number"] == "10002"].sort_values("_excel_row")
    first, second = copies.iloc[0], copies.iloc[1]
    assert first["deal_key"] == "Bridge::10002" and not first["is_duplicate"]
    assert second["deal_key"] == f"Bridge::10002#r{second['_excel_row']}" and second["is_duplicate"]
    # keys do not depend on the hidden-rows toggle
    assert set(deck["deal_key"]) <= set(full["deal_key"])
    warning = next(n for n in _notes(app_module, parsed, "warning") if n.startswith("Duplicate deal numbers"))
    assert f"rows {first['_excel_row']}, {second['_excel_row']}" in warning


def test_duplicate_rows_stay_out_of_history_and_streaks(app_module):
    parsed = _parse(app_module, duplicate_deal="visible")
    deck, _ = _deck(app_module, parsed)
    snapshot = app_module.current_status_history_snapshot(deck, AS_OF, "x.xlsx")
    assert "Bridge::10002#r14" not in set(snapshot["deal_key"])
    assert (snapshot["deal_key"] == "Bridge::10002").sum() == 1
    enriched = app_module.add_status_staleness_columns(deck, app_module.embedded_status_history(parsed), AS_OF)
    dup = enriched[enriched["deal_key"] == "Bridge::10002#r14"].iloc[0]
    assert dup["status_same_meeting_count"] == 0 and not dup["status_needs_update_prompt"]
    assert enriched.loc[enriched["deal_key"] == "Bridge::10002", "status_same_meeting_count"].item() == 3


# -----------------------------------------------------------------------------
# 7. Dated columns, reference date and the as-of date (U20)
# -----------------------------------------------------------------------------
def test_each_sheet_uses_its_newest_column_on_or_before_as_of(app_module, sample_bytes):
    parsed = app_module.parse_workbook_bytes(sample_bytes)
    deck14, info14 = _deck(app_module, parsed, as_of=dt.date(2026, 10, 14))
    assert info14["sheets"]["Bridge"]["upb_header"] == "10/8 UPB"
    assert info14["sheets"]["Term"]["upb_header"] == "10/8/26 UPB"
    assert info14["sheets"]["Bridge"]["uses_older_column"] and info14["sheets"]["Term"]["uses_older_column"]
    deck15, info15 = _deck(app_module, parsed)
    assert (info15["sheets"]["Bridge"]["upb_header"], info15["sheets"]["Term"]["upb_header"]) == ("10/15 UPB", "10/15/26 UPB")
    assert deck15["upb"].sum() == SAMPLE_TOTAL_UPB
    assert not info15["sheets"]["Bridge"]["uses_older_column"]
    assert info15["sheets"]["Bridge"]["dated_columns"]["UPB"] == [("2026-10-08", "10/8 UPB"), ("2026-10-15", "10/15 UPB")]
    assert info15["sheets"]["Term"]["dated_columns"]["UPB"] == [("2026-10-08", "10/8/26 UPB"), ("2026-10-15", "10/15/26 UPB")]
    # derive only: the parse is independent of the as-of date
    assert deck14["deal_key"].tolist() == deck15["deal_key"].tolist()


def test_year_boundary_headers(app_module):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bridge"
    ws.append(["Deal Number", "Deal Name", "12/31 UPB", "1/7/27 UPB"])
    ws.append(["1", "Year End Deal", 100, 200])
    parsed = app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb))
    assert parsed["reference_date"] == "2027-01-07" and parsed["reference_source"] == "header"
    deck, info = _deck(app_module, parsed, as_of=dt.date(2027, 1, 7))
    assert info["sheets"]["Bridge"]["dated_columns"]["UPB"] == [("2026-12-31", "12/31 UPB"), ("2027-01-07", "1/7/27 UPB")]
    assert deck["upb"].tolist() == [200]
    older, _ = _deck(app_module, parsed, as_of=dt.date(2027, 1, 6))
    assert older["upb"].tolist() == [100]


@pytest.mark.parametrize(
    "title, expected",
    [("Portfolio Overview - Bridge as of 10/15/2026", dt.date(2026, 10, 15)),
     ("Portfolio Overview as of 10/15/26", dt.date(2026, 10, 15)),
     ("Report as of 2026-10-15", dt.date(2026, 10, 15)),
     ("As of October 15, 2026", dt.date(2026, 10, 15)),
     ("As of: Oct. 15th, 2026", dt.date(2026, 10, 15)),
     ("As of next week", None),
     ("Portfolio Overview", None)],
)
def test_title_row_as_of(app_module, title, expected):
    assert app_module._title_as_of([(title,), ()]) == expected


def test_title_as_of_in_a_neighbouring_cell(app_module):
    assert app_module._title_as_of([("As of", dt.datetime(2026, 10, 8))]) == dt.date(2026, 10, 8)


def test_reference_date_without_title_uses_newest_dated_header_with_a_year(app_module):
    wb = sw.build_sample_workbook(as_of=AS_OF)
    for name in ("Bridge", "Term"):
        wb[name]["A1"] = "Portfolio Overview"
    parsed = app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb))
    assert parsed["reference_date"] == "2026-10-15" and parsed["reference_source"] == "header"


def test_first_run_takes_the_as_of_date_from_the_workbook(make_at):
    at = make_at(as_of=dt.date(2026, 10, 8))
    at.run()
    assert not at.exception
    assert at.session_state["as_of_date"] == dt.date(2026, 10, 8)
    assert any("From workbook title row: 10/08/2026" in c.value for c in at.sidebar.caption)


def test_as_of_differs_warning_offers_the_workbook_date(make_at):
    at = make_at()
    at.run()
    assert at.session_state["as_of_date"] == AS_OF
    assert not [w for w in at.sidebar.warning if w.proto.title == "As-of differs from the workbook"]
    at.sidebar.date_input(key="as_of_date").set_value(dt.date(2026, 10, 14)).run()
    assert at.session_state["as_of_date"] == dt.date(2026, 10, 14)
    warning = next(w for w in at.sidebar.warning if w.proto.title == "As-of differs from the workbook")
    assert "'10/8 UPB'" in warning.value and "'10/8/26 UPB'" in warning.value
    at.run()  # a rerun keeps the user's date for this workbook
    assert at.session_state["as_of_date"] == dt.date(2026, 10, 14)
    next(b for b in at.sidebar.button if b.label == "Use 10/15/2026").click().run()
    assert at.session_state["as_of_date"] == AS_OF
    assert not [w for w in at.sidebar.warning if w.proto.title == "As-of differs from the workbook"]


def test_a_new_workbook_brings_its_own_as_of_date(make_at):
    at = make_at()
    at.run()
    at.sidebar.date_input(key="as_of_date").set_value(dt.date(2026, 10, 14)).run()
    upload_workbook(at, sw.build_sample_workbook_bytes(as_of=dt.date(2026, 10, 22)), "next_week.xlsx")
    assert not at.exception
    assert at.session_state["as_of_date"] == dt.date(2026, 10, 22)


# -----------------------------------------------------------------------------
# 8. Days past due: NaN, not 0, when nothing is numeric
# -----------------------------------------------------------------------------
def test_dpd_is_nan_when_no_dpd_cell_is_numeric(app_module):
    parsed = _parse(
        app_module,
        cell_overrides={("Term", "20001", "Days Past Due"): None, ("Term", "20002", "Days Past Due"): "N/A",
                        ("Term", "20003", "Days Past Due"): "60"},
    )
    deck, _ = _deck(app_module, parsed)
    assert np.isnan(_row(deck, "20001")["days_past_due"])
    assert np.isnan(_row(deck, "20002")["days_past_due"])
    assert _row(deck, "20003")["days_past_due"] == 60
    assert _row(deck, "10004")["days_past_due"] == 95
    assert app_module.fmt_int(_row(deck, "20001")["days_past_due"]) == "-"


def test_dpd_takes_the_largest_value_across_dpd_columns(app_module):
    parsed = _parse(app_module, duplicate_headers=True,
                    cell_overrides={("Bridge", "10001", "Days Past Due"): None})
    deck, _ = _deck(app_module, parsed)
    assert np.isnan(_row(deck, "10001")["days_past_due"])  # both DPD cells blank
    wb = sw.build_sample_workbook(as_of=AS_OF, duplicate_headers=True)
    ws = wb["Bridge"]
    ws.cell(5, ws.max_column).value = 40  # second DPD column for 10001
    deck2, _ = _deck(app_module, app_module.parse_workbook_bytes(sw.workbook_to_bytes(wb)))
    assert _row(deck2, "10001")["days_past_due"] == 40


# -----------------------------------------------------------------------------
# 9. Loan mods: key headers, match flag, canonical key, one merge (U21/U46/U47/U42/U48/U60)
# -----------------------------------------------------------------------------
def test_opportunity_number_only_extract(app_module):
    parsed = _parse(app_module, mods_key="opportunity_only")
    assert parsed["loan_mod_info"]["key_header"] == "Opportunity Deal Loan Number"
    assert len(parsed["loan_mods"]) == 7
    deck, _ = _deck(app_module, parsed)
    assert dict(zip(deck["deal_number"], deck["loan_mod_count"]))["10007"] == 3


def test_zero_padded_mod_numbers_attach(app_module):
    parsed = _parse(app_module, mods_key="zero_padded")
    deck, _ = _deck(app_module, parsed)
    assert _row(deck, "10002")["loan_mod_count"] == 2
    assert parsed["loan_mods"]["Deal Number"].iloc[0] == "00010002"  # stored value unchanged
    assert len(app_module.loan_mods_for_deal(parsed["loan_mods"], "10002")) == 2


@pytest.mark.parametrize("encoding", ["yes_no", "bool", "yn", "10", "x"])
def test_match_flag_encodings_filter_rows(app_module, encoding):
    parsed = _parse(app_module, match_encoding=encoding)
    mods = parsed["loan_mods"]
    assert len(mods) == 7 and "99999" not in set(mods["Deal Number"])
    assert parsed["loan_mod_info"]["match_filter_applied"] is True


def test_match_flag_fails_open_when_nothing_is_recognised(app_module):
    parsed = _parse(app_module, match_encoding="blank")
    assert len(parsed["loan_mods"]) == 8 and parsed["loan_mod_info"]["match_filter_applied"] is False
    deck, _ = _deck(app_module, parsed)
    assert int(deck["loan_mod_count"].sum()) == 7  # 99999 is not on the deck


def test_loan_mod_columns_come_from_one_join(app_module, sample_bytes):
    parsed = app_module.parse_workbook_bytes(sample_bytes)
    deck, _ = _deck(app_module, parsed)
    counts = dict(zip(deck["deal_number"], deck["loan_mod_count"]))
    assert counts == {"10001": 0, "10002": 2, "10003": 1, "10004": 0, "10005": 0, "10007": 3, "10008": 0,
                      "20001": 0, "20002": 0, "20003": 1}
    assert str(deck["loan_mod_count"].dtype) == "int64"
    assert str(deck["loan_mod_latest_completed_date"].dtype) == "datetime64[ns]"
    assert deck["loan_mod_latest_updated_commitment"].dtype == float
    row = _row(deck, "10002")
    assert row["loan_mod_latest_order_number"] == "2" and row["loan_mod_latest_type"] == "Maturity Extension"
    assert "New maturity" in row["loan_mod_latest_summary"]
    blank = _row(deck, "10001")
    assert blank["loan_mod_latest_summary"] == "" and pd.isna(blank["loan_mod_latest_updated_maturity_date"])
    # idempotent: applying the join again gives the same columns
    again = app_module.add_loan_modification_columns(deck, parsed["loan_mods"])
    pd.testing.assert_frame_equal(again, deck)
    # no mods at all
    none = app_module.add_loan_modification_columns(deck.drop(columns=[c for c in deck if c.startswith("loan_mod_")]),
                                                    app_module.empty_loan_modifications())
    assert (none["loan_mod_count"] == 0).all() and (none["loan_mod_latest_summary"] == "").all()


def test_cancelled_latest_mod_is_not_the_headline(app_module):
    parsed = _parse(app_module, cancelled_latest_mod=True)
    deck, _ = _deck(app_module, parsed)
    row = _row(deck, "10008")
    assert (row["loan_mod_count"], row["loan_mod_active_count"], row["loan_mod_inactive_count"]) == (2, 1, 1)
    assert row["loan_mod_latest_order_number"] == "1" and row["loan_mod_latest_status"] == "Completed"
    assert row["loan_mod_latest_request_note"] == "Mod #2 Maturity Extension: Cancelled"
    mods = app_module.loan_mods_for_deal(parsed["loan_mods"], "10008")
    assert app_module.build_loan_mod_toggle_labels(mods) == ["Mod #2 · Cancelled", "Mod #1"]
    assert app_module.headline_mod_position(mods) == 1
    only_cancelled = _row(deck, "20003")
    assert only_cancelled["loan_mod_active_count"] == 0
    assert "maturity" not in only_cancelled["loan_mod_latest_summary"].lower()
    assert app_module.loan_mod_kpi_helper(only_cancelled) == ("Cancelled request", "rt-signal-muted")


def test_new_unnumbered_mod_ranks_first(app_module):
    parsed = _parse(app_module, unnumbered_new_mod=True)
    mods = app_module.loan_mods_for_deal(parsed["loan_mods"], "10007")
    assert mods["Loan Mod Order Number"].tolist() == ["", "3", "2", "1"]
    deck, _ = _deck(app_module, parsed)
    row = _row(deck, "10007")
    assert row["loan_mod_latest_status"] == "In Progress"
    assert "Requested maturity" in row["loan_mod_latest_summary"]


def test_one_mod_type_everywhere(app_module):
    parsed = _parse(app_module, cell_overrides={("Loan Modifications", "10002", "Loan Mod Type"): "Extension"})
    deck, _ = _deck(app_module, parsed)
    row = _row(deck, "10002")
    record = app_module.loan_mods_for_deal(parsed["loan_mods"], "10002").iloc[0].to_dict()
    assert record["Loan Mod Type"] == "Extension"
    assert row["loan_mod_latest_type"] == app_module.loan_mod_type(record) == "Maturity Extension"
    assert row["loan_mod_latest_summary"].startswith("Maturity Extension")
    no_type = row.copy()
    no_type["loan_mod_latest_type"] = ""
    assert app_module.loan_mod_kpi_helper(no_type) == ("Type not set", "rt-signal-amber")


def test_loan_mod_comments_keep_line_breaks(app_module):
    # Excel stores a CRLF inside a cell as '_x000D_' plus a line break.
    parsed = _parse(app_module, cell_overrides={("Loan Modifications", "10002", "Comments"): "Line one_x000D_\nLine two"})
    record = app_module.loan_mods_for_deal(parsed["loan_mods"], "10002").iloc[0]
    assert record["Comments"] == "Line one\nLine two"
    assert app_module.split_loan_mod_note_lines("Comments", record["Comments"]) == ["Line one", "Line two"]
    deck, _ = _deck(app_module, parsed)
    assert "Line one Line two" in _row(deck, "10002")["loan_mod_search_text"]


# -----------------------------------------------------------------------------
# 10. Formula cells without saved values (Deal Number and every other column)
# -----------------------------------------------------------------------------
def test_formula_deal_numbers_without_saved_values_are_counted(app_module):
    parsed = _parse(app_module, formula_deal_numbers=("10007", "10008"))
    report = parsed["sheets"]["Bridge"]
    assert report["formula_deal_uncached"] == [11, 12] and report["formula_deal_cached"] == 0
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 8
    assert any("2 Deal Number cells are formulas" in note for note in _notes(app_module, parsed, "warning"))


def test_formula_deal_numbers_with_saved_values_load(app_module):
    data = sw.build_sample_workbook_bytes(as_of=AS_OF, formula_deal_numbers=("10007",))
    data = sw.with_cached_formula_values(data, "Bridge", {"A11": 10007})
    parsed = app_module.parse_workbook_bytes(data)
    assert parsed["sheets"]["Bridge"]["formula_deal_cached"] == 1
    assert parsed["sheets"]["Bridge"]["formula_deal_uncached"] == []
    deck, _ = _deck(app_module, parsed)
    assert len(deck) == 10 and "10007" in set(deck["deal_number"])


def test_other_formula_columns_without_saved_values_are_reported(app_module):
    parsed = _parse(app_module, term_loan_size="formula",
                    cell_overrides={("Term", "20001", "Days Past Due"): "=0"})
    columns = {item["header"]: item for item in parsed["sheets"]["Term"]["formula_uncached_columns"]}
    assert columns["Loan Amount"]["count"] == 3 and columns["Loan Amount"]["column"] == "I"
    assert columns["Days Past Due"]["rows"] == [4]
    warning = next(n for n in _notes(app_module, parsed, "warning") if n.startswith("Formulas with no saved results"))
    assert "Term 'Loan Amount', 'Days Past Due' (4 cells)" in warning
    assert "Open the workbook in Excel and save it" in warning
    deck, _ = _deck(app_module, parsed)
    assert np.isnan(_row(deck, "20001")["days_past_due"])  # blank, not 0


def test_formula_scan_reads_saved_results(app_module):
    xml = (
        b'<worksheet><sheetData>'
        b'<row r="1"><c r="A1" t="inlineStr"><is><t>Deal Number</t></is></c><c r="B1"><f>TODAY()</f><v/></c></row>'
        b'<row r="2" hidden="1"><c r="A2"><v>1</v></c><c r="B2"><f>A2*2</f><v>2</v></c>'
        b'<c r="C2" t="str"><f>IF(A2&gt;5,"x","")</f><v></v></c></row>'
        b'<row r="3"><c r="A3"><v>2</v></c><c r="B3"><f>A3*2</f><v/></c><x:c r="D3"><x:f>1</x:f></x:c></row>'
        b'<row r="4"><c r="A4"/></row>'
        b'<row r="9"><c r="A9"><v>7</v></c></row>'
        b'</sheetData></worksheet>'
    )
    scan = app_module._scan_sheet_xml(xml, header_row=1)
    assert scan["hidden_rows"] == {2}
    assert scan["formulas"]["B"] == {"cached": 1, "uncached": [3]}  # B1 is above the header
    assert scan["formulas"]["C"] == {"cached": 1, "uncached": []}  # Excel's saved empty string
    assert scan["formulas"]["D"] == {"cached": 0, "uncached": [3]}  # namespaced cell
    assert scan["values_below"] is False
    stopped = app_module._scan_sheet_xml(xml, header_row=1, last_row=4)
    assert stopped["values_below"] is True and 9 not in stopped["hidden_rows"]


# -----------------------------------------------------------------------------
# Loan size: the 'Loan Commitment' tile (A3; the user's Term report)
# -----------------------------------------------------------------------------
TERM_AMOUNTS = [22_000_000.0, 640_000.0, 1_875_000.0]


@pytest.mark.parametrize(
    "mode, header",
    [("normal", "Loan Amount"), ("renamed", "Orig. Loan Amt ($)"), ("text", "Loan Amount"),
     ("blank_first", "Original Loan Amount"), ("zero_commitment", "Loan Amount")],
)
def test_term_loan_size_is_found(app_module, mode, header):
    parsed = _parse(app_module, term_loan_size=mode)
    deck, _ = _deck(app_module, parsed)
    term = deck[deck["sheet"] == "Term"]
    assert term["commitment"].tolist() == TERM_AMOUNTS
    assert term["loan_amount"].tolist() == TERM_AMOUNTS
    assert set(term["commitment_source"]) == {"Original loan amount"}
    assert set(term["commitment_header"]) == {header}
    tile = app_module.loan_size_kpi(term.iloc[0])
    assert tile[:3] == ("Loan Commitment", "$22.0MM", "Original loan amount")
    assert header in tile[4]
    assert not [n for n in _notes(app_module, parsed) if "loan size" in n]


def test_explicit_term_commitment_wins_per_row(app_module):
    parsed = _parse(app_module, term_commitment={"20001": 25_000_000, "20002": None, "20003": 2_000_000})
    deck, _ = _deck(app_module, parsed)
    term = deck[deck["sheet"] == "Term"]
    assert term["commitment"].tolist() == [25_000_000.0, 640_000.0, 2_000_000.0]
    assert term["commitment_source"].tolist() == ["Total commitment", "Original loan amount", "Total commitment"]
    assert term["loan_amount"].tolist() == TERM_AMOUNTS


def test_missing_term_loan_size_is_explained(app_module):
    parsed = _parse(app_module, term_loan_size="missing")
    deck, _ = _deck(app_module, parsed)
    row = _row(deck, "20001")
    assert np.isnan(row["commitment"]) and row["commitment_source"] == ""
    assert app_module.loan_size_kpi(row)[1:3] == ("-", "Not in workbook")
    warning = next(n for n in _notes(app_module, parsed, "warning") if "no loan size column" in n)
    assert warning.startswith("Term:") and "'Loan Amount'" in warning


def test_partly_blank_loan_size_is_listed(app_module):
    parsed = _parse(app_module, cell_overrides={("Term", "20002", "Loan Amount"): None})
    caption = next(n for n in _notes(app_module, parsed, "caption") if "no loan size" in n)
    assert caption == "Term: no loan size for 1 of 3 deals (row 5); 'Loan Amount' is blank there."


def test_bridge_commitment_keeps_its_source(app_module, sample_bytes):
    deck, _ = _deck(app_module, app_module.parse_workbook_bytes(sample_bytes))
    bridge = deck[deck["sheet"] == "Bridge"]
    assert set(bridge["commitment_source"]) == {"Total commitment"}
    assert set(bridge["commitment_header"]) == {"Loan Commitment"}
    assert bridge["loan_amount"].isna().all()
    assert app_module.loan_size_kpi(_row(deck, "10001"))[:3] == ("Loan Commitment", "$12.5MM", "Total commitment")


# -----------------------------------------------------------------------------
# Text fidelity in the loader and history (U15)
# -----------------------------------------------------------------------------
def test_loader_keeps_free_text_as_written(app_module, sample_bytes):
    deck, _ = _deck(app_module, app_module.parse_workbook_bytes(sample_bytes))
    assert "\n" in _row(deck, "10001")["commentary"]
    assert _row(deck, "10003")["commentary"].endswith("9.0")
    assert _row(deck, "10001")["deal_number"] == "10001"  # 10001.0 float cell
    assert _row(deck, "10007")["borrower"] == "Cedar Point Ventures (LLC)"


def test_history_normalization_keeps_line_breaks_and_ids(app_module):
    raw = pd.DataFrame(
        {
            "meeting_date": ["2026-10-08", 46296, dt.datetime(2026, 9, 24), "09/17/2026"],
            "sheet": ["Bridge"] * 4,
            "deal_number": [10001.0, "10001", "10001.0", "00123"],
            "deal_key": ["Bridge::10001.0", "Bridge::10001", "Bridge::10001", "Bridge::00123"],
            "status": ["Rate step-up to 9.0", "Surveillance", "Surveillance", "NA"],
            "owner": [" Alice  Chen ", "Alice Chen", "Alice Chen", "None"],
            "commentary": ["Line one\r\nLine two", "a\x0bb", "", "x"],
            "workbook_name": ["w.xlsx"] * 4,
            "saved_at": ["2026-10-08 10:00:00"] * 4,
        }
    )
    history = app_module.normalize_status_history(raw)
    assert history["meeting_date"].tolist() == ["2026-09-17", "2026-09-24", "2026-10-01", "2026-10-08"]
    by_date = history.set_index("meeting_date")
    assert by_date.loc["2026-10-08", "status"] == "Rate step-up to 9.0"
    assert by_date.loc["2026-10-08", "commentary"] == "Line one\nLine two"
    assert by_date.loc["2026-10-01", "commentary"] == "a\nb"
    assert by_date.loc["2026-10-08", "owner"] == "Alice Chen"
    assert by_date.loc["2026-10-08", "deal_key"] == "Bridge::10001"
    assert by_date.loc["2026-09-17", "deal_number"] == "00123"
    assert list(history.columns) == app_module.STATUS_HISTORY_COLUMNS


def test_npl_dates_and_blanks_display_cleanly(app_module):
    parsed = _parse(app_module, cell_overrides={("Bridge", "10001", "10/15 NPL"): dt.datetime(2026, 10, 1)})
    deck, _ = _deck(app_module, parsed)
    assert _row(deck, "10001")["npl_raw"] == "10/01/2026"
    assert _row(deck, "10002")["npl_raw"] == ""
    assert _row(deck, "10004")["npl_raw"] == "90+"
    assert not deck["npl_raw"].isin(["nan", "NaT", "None"]).any()


# -----------------------------------------------------------------------------
# Parse/derive split: toggles and dates never re-read the file (U09)
# -----------------------------------------------------------------------------
def test_toggles_and_dates_do_not_reparse_but_a_new_file_does(patched_at):
    at = patched_at(count=("parse_workbook_bytes", "build_deck"))
    at.run()
    assert not at.exception
    assert call_counts()["parse_workbook_bytes"] == 1
    builds = call_counts()["build_deck"]
    at.sidebar.toggle(key="include_hidden").set_value(True).run()
    at.sidebar.date_input(key="as_of_date").set_value(dt.date(2026, 10, 14)).run()
    at.sidebar.toggle(key="include_hidden").set_value(False).run()
    assert call_counts()["parse_workbook_bytes"] == 1
    # (hidden, 10/15), (hidden, 10/14) and (visible, 10/14) are new derive keys
    assert call_counts()["build_deck"] == builds + 3
    for view in ("Overview", "Review", "Exports", "Presentation"):
        at.sidebar.radio(key="view_mode").set_value(view).run()
    assert call_counts()["parse_workbook_bytes"] == 1
    assert call_counts()["build_deck"] == builds + 3
    upload_workbook(at, sw.build_sample_workbook_bytes(as_of=dt.date(2026, 10, 22)))
    assert not at.exception
    assert call_counts()["parse_workbook_bytes"] == 2


def test_parse_never_raises_on_garbage(app_module):
    parsed = app_module.parse_workbook_bytes(b"not a workbook")
    assert parsed["error"] and parsed["sheets"] == {}
    assert app_module.workbook_notes(parsed) == []
    with pytest.raises(ValueError):
        app_module.load_portfolio_workbook(b"not a workbook", AS_OF.isoformat(), False)


@pytest.mark.parametrize("deal_number", ["10001", "10003", "10005", "20002"])
def test_saving_the_edit_dialog_unchanged_creates_no_override(make_at, deal_number):
    """Loaded text round-trips through the edit form unchanged (line breaks, '9.0', blanks).
    The edit path itself is rebuilt in P1b; this pins the text-helper half of the rule."""
    from app_helpers import button_by_label, jump_to

    at = make_at()
    at.run()
    jump_to(at, deal_number)
    button_by_label(at, "Edit status").click().run()
    assert [t.label for t in at.text_area] == ["AM Commentary"]
    button_by_label(at, "Save changes").click().run()
    assert not at.exception
    assert at.session_state["meeting_overrides"] == {}


def test_term_loan_commitment_tile_end_to_end(make_at):
    from app_helpers import jump_to, kpi_tiles

    at = make_at()
    at.run()
    jump_to(at, "20001")
    tile = kpi_tiles(at)["Loan Commitment"]
    assert (tile.value, tile.helper, tile.tooltip) == ("$22.0MM", "Original loan amount", "From the Term column 'Loan Amount'")
    jump_to(at, "10001")
    tile = kpi_tiles(at)["Loan Commitment"]
    assert (tile.value, tile.helper) == ("$12.5MM", "Total commitment")


def test_missing_term_loan_size_is_explained_on_screen(make_at):
    from app_helpers import jump_to, kpi_tiles

    at = make_at(term_loan_size="missing")
    at.run()
    assert any("Term: no loan size column found" in w.value for w in at.sidebar.warning)
    jump_to(at, "20001")
    tile = kpi_tiles(at)["Loan Commitment"]
    assert (tile.value, tile.helper) == ("-", "Not in workbook")
    assert "No loan commitment or loan amount" in tile.tooltip


def test_formulas_in_totals_rows_are_not_reported(app_module):
    parsed = _parse(app_module, junk_rows=True, cell_overrides={("Bridge", "Total", "Loan Commitment"): "=SUM(M5:M12)"})
    assert parsed["sheets"]["Bridge"]["formula_uncached_columns"] == []
    assert not [n for n in _notes(app_module, parsed, "warning") if n.startswith("Formulas with no saved results")]


def test_formula_warnings_from_both_sheets_are_one_note(app_module):
    parsed = _parse(app_module, cell_overrides={("Bridge", "10001", "Days Past Due"): "=0",
                                                ("Term", "20001", "Days Past Due"): "=0",
                                                ("Term", "20002", "Days Past Due"): "=2"})
    warnings = [n for n in _notes(app_module, parsed, "warning") if n.startswith("Formulas with no saved results")]
    assert warnings == [
        "Formulas with no saved results: Bridge 'Days Past Due' (1 cell); Term 'Days Past Due' (2 cells). "
        "Those fields show as blank ('-') in the deck. Open the workbook in Excel and save it, then upload it again."
    ]
