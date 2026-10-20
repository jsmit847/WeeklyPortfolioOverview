"""One named regression test per defect closed by the data-layer package (P1a).

Ids refer to the verified bug list (discover/bugs_confirmed.md). The tests go through
entry points that also exist in the pre-fix app (load_portfolio_workbook,
load_loan_modifications_from_workbook, add_loan_modification_columns, ...) or through
AppTest, so `WPO_APP_PATH=<old app> pytest tests/test_bug_regressions.py` shows each one
failing on commit 10afc49.
"""
from __future__ import annotations

import datetime as dt
import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import sample_workbook as sw
from app_helpers import (
    button_by_label,
    call_counts,
    goto_view,
    jump_to,
    kpi_tiles,
    metric_by_label,
    progress_text,
)

AS_OF = dt.date(2026, 10, 15)
SAMPLE_TOTAL_UPB = 66_769_850


def _load(app_module, data: bytes, as_of: dt.date = AS_OF, include_hidden: bool = False):
    deck, meta = app_module.load_portfolio_workbook(data, as_of.isoformat(), include_hidden)
    mods = app_module.load_loan_modifications_from_workbook(data)
    deck = app_module.add_loan_modification_columns(deck, mods)
    return deck, meta, mods


def _sample(**variants) -> bytes:
    return sw.build_sample_workbook_bytes(as_of=AS_OF, **variants)


def _row(deck: pd.DataFrame, deal_number: str) -> pd.Series:
    rows = deck[deck["deal_number"].astype(str) == deal_number]
    assert len(rows) == 1, (deal_number, len(rows))
    return rows.iloc[0]


def _streak_deck(rows):
    return pd.DataFrame(
        [{"sheet": "Bridge", "deal_number": key.split("::")[1], "deal_key": key, "status": status, "owner": "",
          "commentary": ""} for key, status in rows]
    )


def _history(rows):
    return pd.DataFrame(
        [{"meeting_date": day, "sheet": "Bridge", "deal_number": key.split("::")[1], "deal_key": key,
          "status": status, "owner": "", "commentary": "", "workbook_name": "w.xlsx", "saved_at": ""}
         for day, key, status in rows]
    )


# -----------------------------------------------------------------------------
def test_U03_note_splitter_keeps_numbers_signs_and_parentheticals(app_module):
    split = app_module.split_loan_mod_note_lines
    for text in ("Rate cut -0.50% and reserve top-up.", "Floor -0.25% per IC", "-$12,345.67 credit",
                 "Unit (V) kitchen remodel", "1.5% exit fee waived", "Extended maturity to 2026. Borrower paid fee",
                 "Cedar Point Ventures (LLC) signed"):
        assert split("Comments", text) == [text]
    assert split("Comments", "1. Rate step-up approved 2. Sponsor guaranty reaffirmed 3. Cash management triggered") == [
        "Rate step-up approved", "Sponsor guaranty reaffirmed", "Cash management triggered"]
    assert split("Comments", "Initial extension (i) paid down $100k (ii) added interest reserve") == [
        "Initial extension", "paid down $100k", "added interest reserve"]
    assert split("Comments", "Line one\nLine two") == ["Line one", "Line two"]


def test_U05_history_meeting_dates_in_any_format_load(app_module):
    history = _history([
        (46289, "Bridge::1", "Surveillance"),  # Excel serial for 09/24/2026
        ("10/01/2026", "Bridge::1", "Surveillance"),
        (dt.datetime(2026, 10, 8), "Bridge::1", "Surveillance"),
        ("2026-10-08 00:00:00", "Bridge::2", "Surveillance"),
    ])
    normalized = app_module.normalize_status_history(history)
    assert len(normalized) == 4
    assert sorted(set(normalized["meeting_date"])) == ["2026-09-24", "2026-10-01", "2026-10-08"]


def test_U06_timezone_far_future_and_sentinel_dates_do_not_crash(app_module):
    deck, _, mods = _load(app_module, _sample(bad_dates=True, sf_utc_stamps=True))
    assert len(deck) == 10
    assert pd.Timestamp(_row(deck, "20001")["next_payment_date"]) == pd.Timestamp("2026-10-30")
    assert pd.isna(_row(deck, "20002")["maturity_date"])
    sentinel = mods[(mods["Deal Number"] == "10002") & (mods["Loan Mod Order Number"] == "2")].iloc[0]
    assert pd.isna(sentinel["Updated Maturity Date"])
    completed = mods[(mods["Deal Number"] == "10007") & (mods["Loan Mod Order Number"] == "3")].iloc[0]
    assert pd.Timestamp(completed["System Task Completed Date"]) == pd.Timestamp("2026-09-05 17:22:10")
    assert app_module.fmt_date(5) == "-" and app_module.fmt_date(45000) == "03/15/2023"
    assert app_module.fmt_date(20260331) == "03/31/2026"


def test_U08_status_streaks_are_cached_across_reruns(patched_at):
    at = patched_at(count=("status_streak_table",))
    at.run()
    assert not at.exception
    first = call_counts()["status_streak_table"]
    for _ in range(3):
        button_by_label(at, "Next").click().run()
    goto_view(at, "Overview")
    assert not at.exception
    assert call_counts()["status_streak_table"] == first


def test_U09_hidden_rows_toggle_and_as_of_change_do_not_reparse(patched_at):
    at = patched_at(count=("parse_workbook_bytes",))
    at.run()
    assert not at.exception
    at.sidebar.toggle(key="include_hidden").set_value(True).run()
    at.sidebar.date_input(key="as_of_date").set_value(dt.date(2026, 10, 14)).run()
    assert call_counts()["parse_workbook_bytes"] == 1


def test_U12_duplicate_deal_numbers_get_unique_keys(app_module):
    deck, _, _ = _load(app_module, _sample(duplicate_deal="visible"))
    assert len(deck) == 11
    assert deck["deal_key"].is_unique
    assert {"Bridge::10002", "Bridge::10002#r14"} <= set(deck["deal_key"])


def test_U12_next_reaches_every_agenda_item_with_a_duplicate(make_at):
    at = make_at(duplicate_deal="visible")
    at.run()
    assert not at.exception
    seen = [progress_text(at)]
    for _ in range(10):
        button_by_label(at, "Next").click().run()
        seen.append(progress_text(at))
    assert not at.exception
    assert seen == [f"Agenda item {i} of 11" for i in range(1, 12)]
    assert button_by_label(at, "Next").disabled
    box = next(sb for sb in at.selectbox if sb.label == "Jump to deal")
    box.set_value(box.options[7]).run()  # the second 10002 row
    assert not at.exception
    assert progress_text(at) == "Agenda item 8 of 11"


def test_U15_free_text_keeps_line_breaks_and_trailing_zero(app_module):
    data = _sample(cell_overrides={("Loan Modifications", "10002", "Comments"): "Extended 6 months.\nFee waived"})
    deck, _, mods = _load(app_module, data)
    assert "\n" in _row(deck, "10001")["commentary"]
    assert _row(deck, "10003")["commentary"] == "Mod approved 3/15. Rate step-up to 9.0"
    comments = mods.loc[mods["Deal Number"] == "10002", "Comments"].tolist()
    assert "Extended 6 months.\nFee waived" in comments


def test_U16_sheet_without_header_does_not_block_the_workbook(app_module):
    for variants in ({"term_sheet": "blank"}, {"term_sheet": "note"}):
        deck, _, _ = _load(app_module, _sample(**variants))
        assert len(deck) == 7 and set(deck["sheet"]) == {"Bridge"}, variants
    deck, _, _ = _load(app_module, _sample(bridge_header_row=32))
    assert len(deck) == 10


def test_U17_renamed_tabs_still_load(app_module):
    deck, _, mods = _load(app_module, _sample(sheet_titles={"Term": "TERM ", "Loan Modifications": "Loan modifications"}))
    assert int((deck["sheet"] == "Term").sum()) == 3
    assert len(mods) == 7
    assert _row(deck, "10007")["loan_mod_count"] == 3


def test_U18_duplicate_headers_read_and_write_the_same_first_column(app_module):
    data = _sample(duplicate_headers=True)
    deck, _, _ = _load(app_module, data)
    assert _row(deck, "10001")["status"] == "Surveillance"
    assert _row(deck, "10004")["status"] == "Foreclosure"
    assert _row(deck, "10004")["days_past_due"] == 95
    row = _row(deck, "10001")
    override = {"sheet": "Bridge", "deal_number": "10001", "status": "Watchlist", "owner": row["owner"],
                "commentary": row["commentary"]}
    exported = app_module.update_workbook_bytes(data, {"Bridge::10001": override})
    reloaded, _, _ = _load(app_module, exported)
    assert _row(reloaded, "10001")["status"] == "Watchlist"
    ws = load_workbook(io.BytesIO(exported))["Bridge"]
    assert (ws.cell(5, 11).value, ws.cell(5, 23).value) == ("Watchlist", "Funded")


def test_U19_total_and_error_rows_are_not_deals(app_module):
    deck, _, _ = _load(app_module, _sample(junk_rows=True))
    assert len(deck) == 10
    assert deck["upb"].fillna(0).sum() == SAMPLE_TOTAL_UPB
    assert not deck["deal_number"].astype(str).isin(["Total", "#N/A", "Grand Total"]).any()


def test_U20_both_sheets_use_the_same_week_when_as_of_is_a_day_early(app_module):
    _, meta, _ = _load(app_module, _sample(), as_of=dt.date(2026, 10, 14))
    assert meta["Bridge"]["upb_header"] == "10/8 UPB"
    assert meta["Term"]["upb_header"] == "10/8/26 UPB"
    _, meta15, _ = _load(app_module, _sample())
    assert (meta15["Bridge"]["upb_header"], meta15["Term"]["upb_header"]) == ("10/15 UPB", "10/15/26 UPB")


def test_U21_extract_keyed_only_by_opportunity_number_loads(app_module):
    deck, _, mods = _load(app_module, _sample(mods_key="opportunity_only"))
    assert len(mods) == 7
    assert _row(deck, "10002")["loan_mod_count"] == 2


def test_U27_money_signs_rounding_and_exact_capital_profile(app_module):
    fmt = app_module.fmt_money
    assert fmt(-1_500_000) == "-$1.5MM"
    assert fmt(-250_000) == "-$250,000"
    assert fmt(999_999.6) == "$1.0MM"
    assert fmt(-0.0) == "$0"
    assert fmt(float("inf")) == "-"
    deck, _, _ = _load(app_module, _sample())
    items = dict(app_module.build_capital_items(_row(deck, "10001")))
    assert (items["Loan Commitment"], items["Funded Amount"], items["Remaining Commitment"]) == (
        "$12,500,000", "$10,250,000", "$2,250,000")


def test_U31_a_missed_meeting_breaks_the_streak(app_module):
    history = _history([
        ("2026-09-24", "Bridge::1", "Surveillance"), ("2026-10-08", "Bridge::1", "Surveillance"),
        ("2026-09-24", "Bridge::2", "Surveillance"), ("2026-10-01", "Bridge::2", "Surveillance"),
        ("2026-10-08", "Bridge::2", "Surveillance"),
    ])
    deck = _streak_deck([("Bridge::1", "Surveillance"), ("Bridge::2", "Surveillance")])
    out = app_module.add_status_staleness_columns(deck, history, AS_OF, "w.xlsx")
    counts = dict(zip(out["deal_key"], out["status_same_meeting_count"]))
    assert counts == {"Bridge::1": 2, "Bridge::2": 4}


def test_U32_history_after_the_meeting_date_is_ignored(app_module):
    history = _history([
        ("2026-09-24", "Bridge::10004", "Foreclosure"), ("2026-10-01", "Bridge::10004", "Foreclosure"),
        ("2026-10-08", "Bridge::10004", "Foreclosure"), ("2026-10-22", "Bridge::10004", "Litigation"),
    ])
    deck = _streak_deck([("Bridge::10004", "Foreclosure")])
    row = app_module.add_status_staleness_columns(deck, history, AS_OF, "w.xlsx").iloc[0]
    assert int(row["status_same_meeting_count"]) == 4
    assert row["status_same_since"] == "09/24/2026"
    assert bool(row["status_needs_update_prompt"])


def test_U42_cancelled_mod_is_not_the_headline(app_module):
    deck, _, _ = _load(app_module, _sample(cancelled_latest_mod=True))
    row = _row(deck, "10008")
    assert row["loan_mod_latest_status"] == "Completed"
    assert row["loan_mod_latest_order_number"] == "1"
    prompts = " ".join(app_module.build_presenter_prompts(row))
    assert "04/02/2029" not in prompts and "Cancelled" not in row["loan_mod_latest_summary"]
    cancelled_only = _row(deck, "20003")
    assert "New maturity" not in cancelled_only["loan_mod_latest_summary"]


def test_U43_maturity_inside_30_days_is_amber(make_at):
    at = make_at()
    at.run()
    jump_to(at, "10002")
    tiles = kpi_tiles(at)
    assert tiles["Maturity"].helper == "In 12d" and tiles["Maturity"].signal == "rt-signal-amber"
    jump_to(at, "10008")
    assert kpi_tiles(at)["Maturity"].signal == "rt-signal-green"


def test_U44_overview_counts_past_and_upcoming_maturities_separately(make_at):
    at = make_at()
    at.run()
    goto_view(at, "Overview")
    assert not at.exception
    assert metric_by_label(at, "Past maturity").value == "2"
    assert metric_by_label(at, "Maturing ≤30d").value == "2"
    for metric in at.metric:
        assert "/" not in str(metric.delta or ""), (metric.label, metric.delta)
        assert metric.proto.show_border, metric.label


def test_U45_term_kpi_shows_the_loan_size(make_at):
    at = make_at()
    at.run()
    jump_to(at, "20001")
    tile = kpi_tiles(at)["Loan Commitment"]
    assert (tile.value, tile.helper) == ("$22.0MM", "Original loan amount")


@pytest.mark.parametrize("mode", ["renamed", "text", "blank_first", "zero_commitment"])
def test_U45_term_loan_size_survives_header_and_value_variants(app_module, mode):
    """The user's 'still seeing it on Term deals' report: loan size under a variant header,
    stored as text, in a second column, or behind a zero commitment."""
    deck, _, _ = _load(app_module, _sample(term_loan_size=mode))
    term = deck[deck["sheet"] == "Term"]
    assert term["commitment"].tolist() == [22_000_000.0, 640_000.0, 1_875_000.0]
    assert app_module.fmt_money(_row(deck, "20001")["commitment"]) == "$22.0MM"


def test_U46_match_flag_accepts_booleans_and_letters(app_module):
    for encoding in ("bool", "yn", "10", "x"):
        _, _, mods = _load(app_module, _sample(match_encoding=encoding))
        assert len(mods) == 7 and "99999" not in set(mods["Deal Number"].astype(str)), encoding


def test_U47_zero_padded_mod_deal_numbers_attach(app_module):
    deck, _, _ = _load(app_module, _sample(mods_key="zero_padded"))
    assert _row(deck, "10002")["loan_mod_count"] == 2
    assert _row(deck, "10007")["loan_mod_count"] == 3


def test_U48_new_unnumbered_mod_is_the_latest(app_module):
    deck, _, mods = _load(app_module, _sample(unnumbered_new_mod=True))
    assert _row(deck, "10007")["loan_mod_latest_status"] == "In Progress"
    first = app_module.loan_mods_for_deal(mods, "10007").iloc[0]
    assert str(first["Status"]) == "In Progress"


def test_U60_one_modification_type_everywhere(app_module):
    deck, _, mods = _load(app_module, _sample(cell_overrides={("Loan Modifications", "10002", "Loan Mod Type"): "Extension"}))
    row = _row(deck, "10002")
    assert row["loan_mod_latest_type"] == "Maturity Extension"
    assert row["loan_mod_latest_summary"].startswith("Maturity Extension")


def test_U61_term_capital_profile_has_no_funded_amount(app_module):
    deck, _, _ = _load(app_module, _sample())
    labels = [label for label, _value in app_module.build_capital_items(_row(deck, "20001"))]
    assert "Funded Amount" not in labels
    assert "Loan Commitment" in labels and "Loan Amount" in labels
    values = dict(app_module.build_capital_items(_row(deck, "20001")))
    assert values["Loan Commitment"] == values["Loan Amount"] == "$22,000,000"


def _edit_deal(at, deal_number: str, status=None):
    """Open the edit dialog on a deal and save it (optionally with a new status)."""
    jump_to(at, deal_number)
    button_by_label(at, "Edit status").click().run()
    if status is not None:
        next(t for t in at.text_input if t.label == "Status").set_value(status)
    button_by_label(at, "Save changes").click().run()
    assert not at.exception
    return at.session_state["meeting_overrides"]


def test_U15_status_edit_keeps_commentary_verbatim(make_at):
    at = make_at()
    at.run()
    overrides = _edit_deal(at, "10003", status="Modification - signed")
    saved = next(item for item in overrides.values() if str(item["deal_number"]) == "10003")
    assert saved["status"] == "Modification - signed"
    assert saved["commentary"] == "Mod approved 3/15. Rate step-up to 9.0"
    overrides = _edit_deal(at, "10001", status="Watchlist")
    saved = next(item for item in overrides.values() if str(item["deal_number"]) == "10001")
    assert saved["commentary"] == "Construction 82% complete.\nSponsor injected $250k of additional equity."
