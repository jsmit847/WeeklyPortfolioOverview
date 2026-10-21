import datetime as dt
from types import SimpleNamespace

import pandas as pd

import portfolio_exports as exports


def context(app, data):
    parsed = app.parse_workbook_bytes(data)
    deck, info = app.build_deck(parsed, dt.date(2026, 10, 15), False)
    return SimpleNamespace(wb_sha=app.bytes_sha256(data), wb={"bytes": data}, workbook_name="synthetic.xlsx",
                           as_of=dt.date(2026, 10, 15), raw_deck=deck, deck=deck, parsed=parsed)


def test_prepare_export_round_trip(app_module, sample_bytes):
    ctx = context(app_module, sample_bytes)
    state = {"meeting_overrides": {}, "status_history": app_module.embedded_status_history(ctx.parsed)}
    bundle = exports.prepare_export(ctx, app_module, state)
    assert bundle["verification"]["rows"] >= 10
    result = app_module.parse_workbook_bytes(bundle["workbook"])
    history = app_module.embedded_status_history(result)
    assert set(history.meeting_date) == set(state["status_history"].meeting_date)
    assert len(history) == len(state["status_history"])
    assert bundle["fingerprint"] == exports.fingerprint(ctx, state)


def test_fingerprint_changes_when_saved_edits_change(app_module, sample_bytes):
    ctx = context(app_module, sample_bytes)
    state = {"meeting_overrides": {}}
    before = exports.fingerprint(ctx, state)
    state["meeting_overrides"] = {"Bridge::10001": {"status": "Synthetic edit"}}
    assert before != exports.fingerprint(ctx, state)


def test_repeat_export_does_not_duplicate_history(app_module, sample_bytes):
    ctx = context(app_module, sample_bytes)
    state = {"meeting_overrides": {}, "status_history": app_module.embedded_status_history(ctx.parsed)}
    first = exports.prepare_export(ctx, app_module, state)
    state["status_history"] = app_module.read_status_history_csv(__import__('io').BytesIO(first["history"]))
    second = exports.prepare_export(ctx, app_module, state)
    first_frame = app_module.read_status_history_csv(__import__('io').BytesIO(first["history"]))
    second_frame = app_module.read_status_history_csv(__import__('io').BytesIO(second["history"]))
    assert len(first_frame) == len(second_frame)
    assert not second_frame.duplicated(["meeting_date", "deal_key"]).any()


def test_export_without_source_history_creates_no_tracking_sheets(app_module):
    import io
    from openpyxl import load_workbook
    from sample_workbook import build_sample_workbook_bytes
    data = build_sample_workbook_bytes(include_history=False)
    ctx = context(app_module, data)
    bundle = exports.prepare_export(ctx, app_module, {"meeting_overrides": {}})
    before = load_workbook(io.BytesIO(data), read_only=True)
    after = load_workbook(io.BytesIO(bundle["workbook"]), read_only=True)
    try:
        assert after.sheetnames == before.sheetnames
    finally:
        before.close()
        after.close()
