"""State, navigation and export regressions for the completed meeting runtime."""
from __future__ import annotations
import datetime as dt
import io
import pandas as pd
from openpyxl import load_workbook
import sample_workbook as sw

AS_OF = dt.date(2026, 10, 15)

class State(dict):
    def __getattr__(self, name):
        try:
            return self[name]
        except KeyError as exc:
            raise AttributeError(name) from exc
    def __setattr__(self, name, value):
        self[name] = value


def test_single_edit_path_preserves_text_clears_and_noops(app_module, sample_bytes, monkeypatch):
    app = app_module
    parsed = app.parse_workbook_bytes(sample_bytes)
    base_rows = app.build_base_rows(parsed)
    state = State()
    monkeypatch.setattr(app.st, "session_state", state)
    app.init_state()
    sha = app.bytes_sha256(sample_bytes)
    state.wb = {"sha256": sha}
    state["_base_rows"] = base_rows
    key = "Bridge::10001"
    first = app.apply_edit(key, commentary="a\x0bb\x07c", owner="")
    assert first.changed and state.state_version == 1
    assert state.meeting_overrides[key]["commentary"] == "a\nb c"
    assert state.meeting_overrides[key]["owner"] == ""
    again = app.apply_edit(key, commentary="a\nb c", owner="")
    assert not again.changed and state.state_version == 1
    app.apply_edit(key, commentary="x" * 5000)
    assert len(state.meeting_overrides[key]["commentary"]) == 4000
    assert any("Trimmed to" in item["body"] for item in state["_toasts"])
    base = base_rows[key]
    app.apply_edit(key, status=base["status"], owner=base["owner"], commentary=base["commentary"])
    assert key not in state.meeting_overrides and state.state_version == 3


def test_review_delta_uses_rendered_row_snapshot_and_supports_clear(app_module):
    keys = ["Bridge::1", "Bridge::2"]
    current = {key: {"status": "Old", "owner": "Alice", "commentary": "Text"} for key in keys}
    changes = app_module.review_delta_to_changes({1: {"Status": "New"}, 0: {"AM Commentary": None}, 99: {"Status": "Wrong"}}, keys, current)
    assert changes == [(keys[0], {"commentary": ""}), (keys[1], {"status": "New"})]


def test_csv_history_roundtrip_preserves_na_labels_and_leading_zeros(app_module):
    frame = pd.DataFrame([{"meeting_date": "2026-10-15", "sheet": "Bridge", "deal_number": "00123", "deal_key": "Bridge::00123", "status": "NA", "owner": "None", "commentary": "Line one\nLine two"}])
    loaded = app_module.read_status_history_csv(io.BytesIO(app_module.status_history_to_csv_bytes(frame)))
    assert len(loaded) == 1
    assert loaded.iloc[0][["deal_number", "status", "owner", "commentary"]].tolist() == ["00123", "NA", "None", "Line one\nLine two"]


def test_duplicate_row_export_and_formula_text_are_exact(app_module):
    app = app_module
    data = sw.build_sample_workbook_bytes(as_of=AS_OF, duplicate_deal="visible")
    parsed = app.parse_workbook_bytes(data)
    base = app.build_base_rows(parsed)
    key = next(key for key in base if key.startswith("Bridge::10002#r"))
    values = app.base_values(base[key])
    values.update(status="=1+1", owner="", commentary="Line one\nLine two")
    payload = app.override_payload(base[key], values, app.bytes_sha256(data), "now")
    exported = app.update_workbook_bytes(data, {key: payload})
    exported = app.write_status_history_sheet_to_workbook_bytes(exported, app.empty_status_history())
    rows = app.build_base_rows(app.parse_workbook_bytes(exported))
    assert rows[key]["status"] == "=1+1"
    assert rows[key]["owner"] == "" and rows[key]["commentary"] == "Line one\nLine two"
    assert rows["Bridge::10002"]["status"] == base["Bridge::10002"]["status"]
    workbook = load_workbook(io.BytesIO(exported), data_only=False)
    layout = parsed["sheets"]["Bridge"]["layout"]
    assert workbook["Bridge"].cell(base[key]["excel_row"], layout["role_cols"]["status"]).data_type == "s"
    workbook.close()


def test_cursor_remains_on_deal_outside_filter_and_dialog_priority(app_module):
    app = app_module
    assert app.resolve_cursor("b", ["a", "b", "c"], ["a", "c"]) == "b"
    assert app.displayed_agenda(["a", "c"], "b", {"a": (0, 1), "b": (0, 2), "c": (0, 3)}) == (["a", "b", "c"], True)
    assert app.dialog_to_render({"kind": "edit", "deal_key": "b"}, view="Presentation", deck_keys=["b"], has_foreign=True)[0] == app.DIALOG_WORKBOOK_SWITCH
    assert app.dialog_to_render({"kind": "edit", "deal_key": "missing"}, view="Presentation", deck_keys=["b"], has_foreign=False) == (None, True)


def test_sold_loan_stale_payment_is_a_data_check(app_module):
    row = pd.Series({"sheet": "Term", "portfolio": "Sold Term", "status": "Sold", "days_past_due": 0, "days_to_next_payment": -522, "next_payment_date": pd.Timestamp("2025-05-01")})
    payment_tile = next(tile for tile in app_module.build_presentation_kpis(row) if tile[0] == "Next Payment")
    assert payment_tile[2:4] == ("Date not current?", "rt-signal-amber")

def test_switching_to_workbook_without_embedded_history_clears_prior_source(app_module, monkeypatch):
    app = app_module
    state = State()
    monkeypatch.setattr(app.st, "session_state", state)
    app.init_state()
    old = pd.DataFrame([{"meeting_date": "2026-10-08", "sheet": "Bridge", "deal_number": "1", "deal_key": "Bridge::1", "status": "Watchlist"}])
    app.store_status_history(app.normalize_status_history(old), "Embedded in first.xlsx")
    state.embedded_status_history_signature = "embedded:first"
    app.load_embedded_status_history_if_available({"hidden_sheets": {}}, "second", "second.xlsx")
    assert state.status_history.empty
    assert state.uploaded_status_history_name == ""
    assert state.embedded_status_history_signature == "embedded:second"
    app.store_status_history(app.normalize_status_history(old), "manual.csv")
    state.uploaded_status_history = object()
    app.load_embedded_status_history_if_available({"hidden_sheets": {}}, "third", "third.xlsx")
    assert len(state.status_history) == 1 and state.uploaded_status_history_name == "manual.csv"

def test_recognized_all_false_mod_match_flags_exclude_rows(app_module):
    raw = pd.DataFrame({"Deal Number": ["10001", "10002"], "Portfolio Overview Match": ["No", False], "Status": ["Completed", "Completed"]})
    mods, info = app_module.normalize_loan_modifications_with_info(raw)
    assert mods.empty
    assert info["match_filter_applied"] is True and info["excluded_by_match"] == 2


def test_confirmation_and_undo_close_and_reopen_current_reminder(make_at):
    from app_helpers import button_by_label
    at = make_at()
    at.session_state["show_update_reminders"] = True
    at.run()
    assert not at.exception
    key = at.session_state["cursor"]
    button_by_label(at, "Confirm still accurate").click().run()
    assert not at.exception
    ctx = at.session_state["_app_context"]
    row = ctx.deck.loc[ctx.deck["deal_key"] == key].iloc[0]
    assert row["reminder_addressed"] == "Confirmed"
    assert not bool(row["status_needs_update_prompt"])
    assert at.session_state["cursor"] == key
    button_by_label(at, "Undo confirmation").click().run()
    assert not at.exception
    ctx = at.session_state["_app_context"]
    row = ctx.deck.loc[ctx.deck["deal_key"] == key].iloc[0]
    assert bool(row["status_needs_update_prompt"])


def test_top_filters_select_portfolio_search_and_reset_in_sync_with_sidebar(make_at):
    at = make_at().run()
    assert not at.exception
    at.button_group(key="top_sheet_filter").set_value("Term").run()
    assert not at.exception
    assert at.session_state["sheet_filter"] == "Term"
    assert at.sidebar.button_group(key="sheet_filter").value == "Term"
    ctx = at.session_state["_app_context"]
    assert set(ctx.filtered.sheet) == {"Term"} and ctx.cursor.startswith("Term::")
    at.text_input(key="top_search_query").set_value("Bayside").run()
    assert not at.exception
    assert at.sidebar.text_input(key="search_query").value == "Bayside"
    assert set(at.session_state["_app_context"].filtered.deal_number) == {"20001"}
    at.sidebar.text_input(key="search_query").set_value("no matching loan").run()
    assert not at.exception
    assert at.text_input(key="top_search_query").value == "no matching loan"
    assert at.session_state["_app_context"].filtered.empty
    at.button(key="top_reset_filters").click().run()
    assert not at.exception
    assert at.button_group(key="top_sheet_filter").value == "All"
    assert at.sidebar.button_group(key="sheet_filter").value == "All"
    assert at.text_input(key="top_search_query").value == ""
    assert len(at.session_state["_app_context"].filtered) == 10
