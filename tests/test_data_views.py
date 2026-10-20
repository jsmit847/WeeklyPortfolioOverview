"""Synthetic-data contracts for Loan Today and data source discovery."""
import datetime as dt
import io

import pandas as pd
import pytest
from openpyxl import Workbook

import portfolio_data_views as views

AS_OF = dt.date(2026, 10, 15)


def test_loan_today_sold_payment_is_stale():
    row = {"sheet": "Term", "portfolio": "Sold Term", "upb": 0, "days_past_due": 0,
           "next_payment_date": dt.date(2025, 5, 9), "commitment": 12_000_000}
    today = views.build_loan_today(row, as_of=AS_OF)
    message = " ".join(today.sentences)
    assert "Sold loan" in message and "looks stale" in message
    assert "overdue" not in message
    assert not any(x in message for x in ("nan", "NaT", "None"))


def test_loan_today_term_never_uses_funded_percent():
    row = {"sheet": "Term", "upb": 10_000_000, "commitment": 20_000_000,
           "funded_amount": 10_000_000, "days_past_due": 0}
    assert "funded" not in " ".join(views.build_loan_today(row, as_of=AS_OF).sentences)


def test_supplemental_canonical_match_and_hidden_exclusion():
    wb = Workbook()
    ws = wb.active
    ws.title = "SF Deal Snapshot"
    ws.append(["Extract"])
    ws.append(["Deal Number", "Stage", "Last Activity Date", "Salesforce Link"])
    ws.append(["00010001", "Extension", dt.date(2026, 10, 14), "https://example.force.com/deal"])
    hidden = wb.create_sheet("Secret")
    hidden.sheet_state = "hidden"
    hidden.append(["Deal Number"])
    hidden.append([10001])
    unrelated = wb.create_sheet("Other")
    unrelated.append(["Summary"])
    frames = views.parse_supplemental(wb)
    assert set(frames) == {"SF Deal Snapshot"}
    assert frames["SF Deal Snapshot"]["_deal_match"].tolist() == ["10001"]
    assert pd.api.types.is_datetime64_any_dtype(frames["SF Deal Snapshot"]["Last Activity Date"])


def test_record_is_exact_excel_row_and_contains_all_headers(app_module, sample_bytes):
    parsed = app_module.parse_workbook_bytes(sample_bytes)
    deck, _ = app_module.build_deck(parsed, AS_OF, False)
    row = deck[deck["deal_number"] == "10001"].iloc[0]
    record = views.raw_record(row, parsed)
    assert set(record) == {c for c in parsed["sheets"]["Bridge"]["raw"] if not c.startswith("_")}
    assert record["Deal Name"] == "Harbor View Lofts"


def test_meeting_comparison_contract(app_module, sample_bytes):
    parsed = app_module.parse_workbook_bytes(sample_bytes)
    deck, _ = app_module.build_deck(parsed, AS_OF, False)
    history = app_module.embedded_status_history(parsed)
    columns = {sheet: list(report["raw"].columns) for sheet, report in parsed["sheets"].items()}
    changes = views.compute_changes(views.comparison_deck(deck, parsed), history, AS_OF, columns).set_index("deal_key")
    assert set(changes.index[changes["owner_changed"]]) == {
        "Bridge::10002", "Bridge::10004", "Bridge::10007", "Term::20003"}
    assert set(changes.index[changes["no_status_last_meeting"]]) == {
        "Bridge::10005", "Bridge::10008", "Term::20002"}
    assert changes.loc["Bridge::10001", "upb_delta"] == 205_000
    assert changes.loc["Term::20001", "upb_delta"] == -110_000
    assert changes.loc["Bridge::10003", "npl_changed"]
    assert changes.loc["Bridge::10004", "npl_changed"]
    assert pd.isna(changes.loc["Term::20001", "npl_prev"])


def test_quick_links_validate_templates_and_schemes():
    row = {"deal_number": "10001", "servicer": "Example"}
    assert views.quick_links({"Salesforce Link": "https://example.force.com/deal"}, row, {}) == [
        ("Open in Salesforce", "https://example.force.com/deal")]
    assert views.quick_links({"URL": "javascript:alert(1)"}, row, {}) == []
    assert views.quick_links({}, row, {"salesforce_deal_url": "https://example.com/{bad}"}) == []
    assert views.quick_links({}, row, {"salesforce_deal_url": "https://example.com/{deal_number}"}) == [
        ("Open in Salesforce", "https://example.com/10001")]
    assert not views.valid_url("https://user:password@example.com")


def test_diff_escapes_currency_and_markup():
    rendered = views.commentary_diff("Request $250k", "Approve $250k")
    assert r"\$250k" in rendered and ":green[Approve]" in rendered


def test_percent_format_is_safe():
    assert views.format_field("LTV", .8) == "0.8%"
    assert views.format_field("Rate", .09) == "9%"
    assert views.format_field("Rate", .09, pd.Series([.09, 9.0])) == "0.09%"


def test_box_scan_metadata_only_and_whole_number_matching(tmp_path, monkeypatch):
    folder = tmp_path / "Asset Management" / "10001 - Harbor View Lofts"
    folder.mkdir(parents=True)
    document = folder / "Inspection.pdf"
    document.write_bytes(b"synthetic")
    (folder / ".hidden.pdf").write_bytes(b"hidden")
    (folder / "~$draft.docx").write_bytes(b"hidden")
    (tmp_path / "110001-other.pdf").write_bytes(b"other")
    def no_open(*args, **kwargs):
        raise AssertionError("file content must not be opened")
    monkeypatch.setattr("builtins.open", no_open)
    index = views.scan_box_index.__wrapped__(str(tmp_path), max_depth=4)
    assert not any(index["name"].str.startswith((".", "~$")))
    matches = views.match_box_documents(index, deal_number="10001")
    assert matches["name"].tolist() == ["Inspection.pdf"]
    assert matches.iloc[0]["Matched by"] == "Deal number"
    limited = views.scan_box_index.__wrapped__(str(tmp_path), max_entries=1)
    assert len(limited) == 1 and limited.attrs["truncated"]


def test_box_name_match_requires_six_characters(tmp_path):
    index = pd.DataFrame([{"path": str(tmp_path / "Bayside Rental" / "Note.pdf"),
                           "rel_path": str(__import__("pathlib").Path("Bayside Rental") / "Note.pdf"),
                           "name": "Note.pdf", "is_dir": False, "modified": pd.Timestamp("2026-10-12"),
                           "size": 10, "ext": ".pdf"}])
    assert len(views.match_box_documents(index, deal_number="20001", borrower="Bayside Rental")) == 1
    assert views.match_box_documents(index, deal_number="20001", borrower="Bay").empty



def test_repeated_links_across_source_rows_render_once(make_at, sample_bytes):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(sample_bytes))
    bridge = workbook["Bridge"]
    header = next(row for row in bridge.iter_rows(max_row=60)
                  if any(cell.value == "Deal Number" for cell in row))
    link_column = bridge.max_column + 1
    bridge.cell(header[0].row, link_column, "Salesforce Link")
    number_column = next(cell.column for cell in header if cell.value == "Deal Number")
    deal_row = next(row for row in bridge.iter_rows(min_row=header[0].row + 1)
                    if views.canonical(row[number_column - 1].value) == "10001")
    bridge.cell(deal_row[0].row, link_column, "https://example.force.com/deal/10001")
    supplemental = workbook.create_sheet("SF Link Extract")
    supplemental.append(["Deal Number", "Salesforce Link"])
    supplemental.append([10001, "https://example.force.com/deal/10001"])
    supplemental.append([10001, "https://example.force.com/deal/10001"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    at = make_at(buffer.getvalue()).run()
    assert not at.exception
    assert any("All workbook fields" in element.label for element in at.expander)
    assert len(at.get("link_button")) == 1
