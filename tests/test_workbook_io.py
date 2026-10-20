import io
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook

import workbook_io as wio


def formula_workbook():
    wb = Workbook()
    wb.active.title = "Bridge"
    wb.active.append(["Deal Number", "Status", "UPB"])
    wb.active.append(["=10000+1", "Current", "=5+2"])
    stream = io.BytesIO()
    wb.save(stream)
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as src, zipfile.ZipFile(result, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(wio.tag("c")):
                    if cell.get("r") in ("A2", "C2"):
                        cell.find(wio.tag("v")).text = "10001" if cell.get("r") == "A2" else "7"
                data = ET.tostring(root)
            dst.writestr(info, data)
    return result.getvalue()


def test_edits_keep_cached_formula_values_and_force_text():
    original = formula_workbook()
    patched = wio.patch_text_cells(original, {"Bridge": {"B2": "=literal\nsecond line"}})
    wb = load_workbook(io.BytesIO(patched), data_only=True)
    assert wb["Bridge"]["A2"].value == 10001
    assert wb["Bridge"]["C2"].value == 7
    assert wb["Bridge"]["B2"].value == "=literal\nsecond line"
    wb.close()
    formula = load_workbook(io.BytesIO(patched), data_only=False)
    assert formula["Bridge"]["B2"].data_type == "s"
    assert formula["Bridge"]["A2"].data_type == "f"


def test_hidden_sheets_keep_formulas_and_can_be_replaced():
    original = formula_workbook()
    frame = pd.DataFrame({"meeting_date": ["2026-10-15"], "commentary": ["=literal"], "number": [0]})
    first = wio.replace_hidden_sheets(original, {"_Status History": frame, "_Action Items": frame})
    second = wio.replace_hidden_sheets(first, {"_Status History": frame.iloc[:0]})
    wb = load_workbook(io.BytesIO(second), data_only=True)
    assert wb["Bridge"]["A2"].value == 10001
    assert wb["Bridge"]["C2"].value == 7
    assert wb["_Status History"].max_row == 1
    assert wb["_Action Items"]["B2"].value == "=literal"
    assert wb["_Action Items"]["C2"].value == 0
    assert wb["_Status History"].sheet_state == "hidden"
    assert len(wb.sheetnames) == 3


def test_untouched_zip_members_identical():
    original = formula_workbook()
    result = wio.patch_text_cells(original, {"Bridge": {"B2": "New"}})
    with zipfile.ZipFile(io.BytesIO(original)) as src, zipfile.ZipFile(io.BytesIO(result)) as dst:
        for name in src.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert src.read(name) == dst.read(name)
