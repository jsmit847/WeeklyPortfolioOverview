"""Prepare, reopen and verify meeting exports before offering a download."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

import portfolio_cockpit as cockpit


def fingerprint(ctx, state) -> str:
    payload = {"workbook": ctx.wb_sha, "as_of": str(ctx.as_of),
               "meeting_date": str(state.get("meeting_date", ctx.as_of)),
               "history": hashlib.sha256(pd.DataFrame(state.get("status_history")).to_csv(index=False).encode()).hexdigest()}
    for key in ("meeting_overrides", "meeting_items", "meeting_progress", "reminder_confirmed", "attendees", "next_meeting", "agenda_order", "_mw_model"):
        payload[key] = state.get(key)
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()


def prepare_export(ctx, app, state, workspace=None) -> dict:
    meeting_date = state.get("meeting_date", ctx.as_of)
    if not isinstance(meeting_date, dt.date):
        meeting_date = dt.date.fromisoformat(str(meeting_date))
    edits = {key: item for key, item in (state.get("meeting_overrides") or {}).items() if item.get("workbook_sha") == ctx.wb_sha}
    history = app.merge_status_history(state.get("status_history"), ctx.deck, meeting_date, ctx.workbook_name)
    result = app.update_workbook_bytes(ctx.wb["bytes"], edits)
    result = app.write_status_history_sheet_to_workbook_bytes(result, history)
    if workspace is not None and hasattr(workspace, "enrich_workbook"):
        result = workspace.enrich_workbook(ctx, result)
    checked = cockpit.verify_workbook(ctx.wb["bytes"], result, edits, app)
    workbook = load_workbook(io.BytesIO(result), read_only=True, data_only=False)
    try:
        sheet = workbook[app.STATUS_HISTORY_SHEET_NAME]
        headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
        if headers != app.STATUS_HISTORY_COLUMNS or sheet.sheet_state != "hidden":
            raise ValueError("The exported status history schema is invalid.")
        if sum(1 for _ in sheet.iter_rows(min_row=2)) != len(history):
            raise ValueError("The exported status history row count is invalid.")
    finally:
        workbook.close()
    agenda = app.build_agenda_dataframe(ctx.deck)
    return {"fingerprint": fingerprint(ctx, state), "workbook": result,
            "filename": f"{Path(ctx.workbook_name).stem}_meeting_{meeting_date.isoformat()}.xlsx",
            "history": app.status_history_to_csv_bytes(history),
            "updates": app.export_overrides_csv(ctx.raw_deck, ctx.deck),
            "agenda": cockpit.safe_csv(agenda), "verification": checked}


def render_exports(ctx, app, workspace=None):
    st.subheader("Verified workbook exports")
    st.caption("Exports include the full loaded portfolio, including deals outside the current filters. Saved edits and meeting records are checked before download.")
    current = fingerprint(ctx, st.session_state)
    bundle = st.session_state.get("prepared_export")
    if bundle and bundle.get("fingerprint") != current:
        bundle = None
        st.session_state.prepared_export = None
        st.info("Meeting data changed. Prepare a fresh export to include the latest work.")
    if st.button("Prepare and verify export", type="primary", icon=":material/fact_check:"):
        try:
            with st.status("Preparing meeting workbook…", expanded=True) as status:
                st.write("Applying saved edits and recording this meeting")
                bundle = prepare_export(ctx, app, st.session_state, workspace)
                st.session_state.prepared_export = bundle
                status.update(label=f"Verified {bundle['verification']['rows']:,} portfolio rows and {bundle['verification']['edits']:,} edited deals", state="complete", expanded=False)
        except Exception as exc:
            bundle = None
            st.session_state.prepared_export = None
            app.LOG.error("event=export_verification_failed type=%s", type(exc).__name__)
            st.error("The export could not be verified. Your saved edits remain available; correct the workbook issue and prepare it again.")
    if not bundle:
        return
    st.success("Workbook checked successfully. Download it to carry status history, follow-ups and meeting records into the next meeting.")
    cols = st.columns(2)
    cols[0].download_button("Download updated workbook", data=bundle["workbook"], file_name=bundle["filename"], mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch", on_click="ignore")
    cols[1].download_button("Download changed statuses/commentary CSV", data=bundle["updates"], file_name="weekly_portfolio_meeting_updates.csv", mime="text/csv", width="stretch", on_click="ignore")
    cols[0].download_button("Download status history CSV backup", data=bundle["history"], file_name="weekly_portfolio_status_history.csv", mime="text/csv", width="stretch", on_click="ignore")
    cols[1].download_button("Download full meeting agenda CSV", data=bundle["agenda"], file_name="weekly_portfolio_agenda.csv", mime="text/csv", width="stretch", on_click="ignore")
